from __future__ import annotations

from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
import openpyxl
import matplotlib.pyplot as plt
import pandas as pd
import pytest

from workbench.artifacts import read_json
from workbench.app import app
from workbench.config import WorkbenchConfig
from workbench.ingestion import ingest_files
from workbench.orchestrator import run_workflow
from workbench.projects import create_project, create_run
from workbench import visualization


def _run(tmp_path: Path, *, labels: dict[str, object] | None = None) -> Path:
    source = tmp_path / "labels.csv"
    pd.DataFrame(
        {
            "outcome": [float(index) + 0.2 * (index % 3) for index in range(48)],
            "treatment": [float(index) for index in range(48)],
        }
    ).to_csv(source, index=False)
    project = create_project(tmp_path, "labels")
    result = run_workflow(
        project.root,
        [source],
        mode="auto",
        y="outcome",
        x=["treatment"],
        model_type="ols",
        labels=labels,
    )
    assert result["status"] == "completed", result
    return project.root / "runs" / result["run_id"]


def test_declared_labels_reach_run_report_and_explicit_xlsx_export(tmp_path: Path) -> None:
    run_root = _run(
        tmp_path,
        labels={
            "variable_labels": {"outcome": "Outcome label", "treatment": "Treatment label"},
            "value_labels": {"treatment": {"0": "Control", "1": "Treated"}},
        },
    )

    inputs = read_json(run_root / "run_inputs.json")
    assert inputs["form"]["labels"]["variable_labels"]["outcome"] == "Outcome label"

    html = (run_root / "reports" / "report.html").read_text(encoding="utf-8")
    assert "Outcome label" in html
    assert "Treatment label" in html
    assert "Control" in html

    assert not (run_root / "exports" / "tables.xlsx").exists()

    response = TestClient(app).post(
        f"/runs/{run_root.name}/report/result-table-export",
        params={"project_root": str(run_root.parent.parent)},
        json={"sections": ["table_1", "regression_table"]},
    )
    assert response.status_code == 200, response.text
    workbook = openpyxl.load_workbook(BytesIO(response.content), read_only=True)
    table_1 = workbook["table_1"]
    rows = list(table_1.iter_rows(values_only=True))
    assert any("Outcome label" in row for row in rows)
    assert any("Treatment label" in row for row in rows)
    regression_table = workbook["regression_table"]
    regression_rows = list(regression_table.iter_rows(values_only=True))
    assert any("Treatment label" in row for row in regression_rows)

    figure_labels = read_json(run_root / "figures" / "figure_labels.json")
    assert figure_labels["scatter_plots"]["x_label"] == "Treatment label"
    assert figure_labels["scatter_plots"]["y_label"] == "Outcome label"


def test_undeclared_labels_use_explicit_column_name_fallback(tmp_path: Path) -> None:
    run_root = _run(tmp_path)
    html = (run_root / "reports" / "report.html").read_text(encoding="utf-8")
    assert "column_name_fallback" in html
    assert "Outcome label" not in html


def test_labels_reject_unknown_top_level_fields(tmp_path: Path) -> None:
    source = tmp_path / "invalid-labels.csv"
    pd.DataFrame({"y": [1.0, 2.0], "x": [1.0, 2.0]}).to_csv(source, index=False)
    project = create_project(tmp_path, "invalid-labels")
    with pytest.raises(ValueError, match="labels"):
        run_workflow(
            project.root,
            [source],
            mode="auto",
            y="y",
            x=["x"],
            labels={"unexpected": {}},
        )


def test_common_figures_use_variable_and_value_labels(monkeypatch, tmp_path: Path) -> None:
    frame = pd.DataFrame(
        {
            "outcome": [float(index) + 0.2 * (index % 3) for index in range(48)],
            "treatment": [float(index) for index in range(48)],
            "group": [index % 2 for index in range(48)],
            "period": pd.date_range("2024-01-01", periods=48, freq="D"),
        }
    )
    frame.attrs["variable_labels"] = {
        "outcome": "Outcome label",
        "treatment": "Treatment label",
        "group": "Group label",
        "period": "Period label",
    }
    frame.attrs["value_labels"] = {
        "group": {"0": "Control", "1": "Treated"},
    }
    captured: dict[str, list[dict[str, object]]] = {}

    def capture(fig, path, run_root, artifact_id, figures) -> None:
        panels: list[dict[str, object]] = []
        for axis in fig.axes:
            legend = axis.get_legend()
            panels.append(
                {
                    "title": axis.get_title(),
                    "xlabel": axis.get_xlabel(),
                    "ylabel": axis.get_ylabel(),
                    "xticks": [tick.get_text() for tick in axis.get_xticklabels()],
                    "yticks": [tick.get_text() for tick in axis.get_yticklabels()],
                    "legend": (
                        [text.get_text() for text in legend.get_texts()]
                        if legend is not None
                        else []
                    ),
                }
            )
        captured[artifact_id] = panels
        figures[artifact_id] = str(path)
        plt.close(fig)

    monkeypatch.setattr(visualization, "_save", capture)
    figures = visualization.create_figures(
        frame,
        tmp_path,
        numeric_columns=["outcome", "treatment", "group"],
        time_column="period",
        outcome_column="outcome",
        regressors=["treatment", "group"],
        model_type="ols",
    )

    assert set(
        (
            "histograms",
            "kde_plots",
            "boxplots",
            "correlation_heatmap",
            "category_counts",
            "scatter_plots",
            "group_boxplots",
            "time_trend",
        )
    ) <= set(figures)
    assert {panel["title"] for panel in captured["histograms"]} >= {
        "Outcome label",
        "Treatment label",
    }
    assert {panel["title"] for panel in captured["kde_plots"]} >= {
        "Outcome label",
        "Treatment label",
    }
    assert captured["boxplots"][0]["xticks"] == [
        "Outcome label",
        "Treatment label",
    ]
    assert captured["correlation_heatmap"][0]["xticks"] == [
        "Outcome label",
        "Treatment label",
    ]
    assert captured["correlation_heatmap"][0]["yticks"] == [
        "Outcome label",
        "Treatment label",
    ]
    assert captured["category_counts"][0]["title"] == "Group label"
    assert set(captured["category_counts"][0]["xticks"]) == {"Control", "Treated"}
    assert captured["scatter_plots"][0]["xlabel"] == "Treatment label"
    assert captured["scatter_plots"][0]["ylabel"] == "Outcome label"
    assert set(captured["group_boxplots"][0]["xticks"]) == {"Control", "Treated"}
    assert captured["group_boxplots"][0]["ylabel"] == "Outcome label"
    assert captured["time_trend"][0]["xlabel"] == "Period label"
    assert set(captured["time_trend"][0]["legend"]) == {
        "Outcome label",
        "Treatment label",
    }


@pytest.mark.parametrize("suffix", [".dta", ".sav"])
def test_external_label_formats_remain_explicitly_unsupported(
    tmp_path: Path, suffix: str
) -> None:
    source = tmp_path / f"external-labels{suffix}"
    source.write_bytes(b"not parsed by this release")
    project = create_project(tmp_path, f"unsupported{suffix[1:]}")
    run = create_run(project.root, mode="auto")

    with pytest.raises(ValueError, match="unsupported file type"):
        ingest_files([source], run.root, WorkbenchConfig())

    assert not (run.root / "raw_snapshot" / source.name).exists()
    assert read_json(run.root / "artifacts_index.json") == {
        "schema_version": 1,
        "artifacts": [],
    }
