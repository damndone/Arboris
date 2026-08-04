import io
import json
import time
from io import BytesIO
from pathlib import Path

import pandas as pd
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from workbench.agent.context_tools import (
    NodeOperationContextProvider,
    _bounded_statistics_evidence,
)
from workbench.agent.tools import ToolContext
from workbench.api import app
from workbench.artifacts import read_json


def _grouped_csv() -> bytes:
    group_values = {
        "a": [1.0, 2.0, 3.0, 4.0] * 3,
        "b": [2.0, 3.0, 4.0, 5.0] * 3,
        "c": [3.0, 4.0, 5.0, 6.0] * 3,
    }
    y = [value for values in group_values.values() for value in values]
    group = [name for name, values in group_values.items() for _ in values]
    frame = pd.DataFrame(
        {
            "y": y,
            "x": [float(index) for index in range(len(y))],
            "group": group,
        }
    )
    return frame.to_csv(index=False).encode()


def _wait_for_completed_run(
    client: TestClient, project_root: str, run_id: str
) -> dict:
    terminal = {"completed", "failed", "blocked", "partial", "interrupted"}
    last_detail: dict = {}
    for _ in range(240):
        response = client.get(
            f"/runs/{run_id}", params={"project_root": project_root}
        )
        assert response.status_code == 200, response.text
        last_detail = response.json()
        if last_detail.get("status") in terminal:
            return last_detail
        time.sleep(0.05)
    raise AssertionError(f"run did not reach a terminal state: {last_detail}")


def test_http_run_exposes_advanced_statistics_evidence_to_consumers(tmp_path):
    client = TestClient(app)
    project_root = client.post(
        "/projects", json={"parent": str(tmp_path), "name": "statistics-http"}
    ).json()["project_root"]

    response = client.post(
        "/runs",
        data={
            "project_root": project_root,
            "mode": "auto",
            "model_type": "ols",
            "y": "y",
            "x": "x,group",
            "labels": json.dumps({
                "variable_labels": {"y": "Outcome"},
                "value_labels": {"group": {"a": "Group A"}},
            }),
        },
        files={"file": ("grouped.csv", io.BytesIO(_grouped_csv()), "text/csv")},
    )

    assert response.status_code == 200, response.text
    run_id = response.json()["run_id"]
    detail = _wait_for_completed_run(client, project_root, run_id)
    assert detail["status"] == "completed"
    assert detail["table_1"]
    assert detail["labels"]["variable_labels"]["y"] == "Outcome"
    assert detail["labels"]["value_labels"]["group"]["a"] == "Group A"

    run_root = Path(project_root) / "runs" / run_id
    packet = read_json(run_root / "statistical_tests" / "evidence.json")
    assert packet["payload_schema"] == "workbench.statistics.evidence-packet"
    assert detail["statistical_evidence"]["payload_schema"] == packet["payload_schema"]
    assert packet["schema_version"] == 1
    assert len(packet["dataset_ref"]["dataset_sha256"]) == 64
    assert packet["lineage_parent"] == "cleaned_dataset"

    anova = next(
        row for row in packet["results"] if row["test_type"] == "anova_posthoc"
    )
    cohens = next(
        row
        for row in packet["results"]
        if row["test_type"] == "cohens_d"
        and row["test_id"] == "cohens_d:y:group:a:b"
    )
    assert anova["statistic"] == 8.8
    assert anova["p_value"] == 0.0008649609084088763
    assert cohens["statistic"] == -2.097617696340303
    assert cohens["p_value"] == 0.047648092397099114

    artifact_index = read_json(run_root / "artifacts_index.json")
    assert artifact_index["schema_version"] == 1
    evidence_record = next(
        record
        for record in artifact_index["artifacts"]
        if record["artifact_id"] == "statistical_tests_evidence"
    )
    assert evidence_record["path"] == "statistical_tests/evidence.json"
    assert evidence_record["artifact_type"] == "statistical_test"
    assert evidence_record["step"] == "statistical_tests"
    assert evidence_record["inputs"] == ["cleaned_dataset"]

    artifact_listing = client.get(
        f"/runs/{run_id}/artifacts", params={"project_root": project_root}
    )
    assert artifact_listing.status_code == 200, artifact_listing.text
    assert any(
        record["artifact_id"] == "statistical_tests_evidence"
        for group in artifact_listing.json()["groups"]
        for record in group["items"]
    )

    report_response = client.get(
        f"/runs/{run_id}/report", params={"project_root": project_root}
    )
    assert report_response.status_code == 200, report_response.text
    report_html = report_response.text
    assert 'id="statistical-evidence"' in report_html
    assert "anova_posthoc" in report_html
    assert "cohens_d" in report_html
    assert str(anova["statistic"]) in report_html
    assert str(anova["p_value_corrected"]) in report_html

    workbook_response = client.get(
        f"/runs/{run_id}/artifacts/tables_xlsx",
        params={"project_root": project_root},
    )
    assert workbook_response.status_code == 404, workbook_response.text

    explicit_export = client.post(
        f"/runs/{run_id}/report/result-table-export",
        params={"project_root": project_root},
        json={"sections": ["regression_table", "statistical_evidence"]},
    )
    assert explicit_export.status_code == 200, explicit_export.text
    workbook = load_workbook(BytesIO(explicit_export.content), read_only=True)
    assert "statistical_evidence" in workbook.sheetnames
    evidence_sheet = workbook["statistical_evidence"]
    evidence_rows = list(evidence_sheet.iter_rows(values_only=True))
    headers = list(evidence_rows[0])
    exported = [
        dict(zip(headers, row, strict=True))
        for row in evidence_rows[1:]
    ]
    exported_anova = next(
        row for row in exported if row["test_type"] == "anova_posthoc"
    )
    assert exported_anova["statistic"] == anova["statistic"]
    assert exported_anova["p_value"] == anova["p_value"]

    agent_projection = _bounded_statistics_evidence(run_root)
    assert agent_projection is not None
    assert agent_projection["payload_schema"] == packet["payload_schema"]
    assert agent_projection["schema_version"] == packet["schema_version"]
    assert agent_projection["dataset_ref"] == packet["dataset_ref"]
    assert agent_projection["lineage_parent"] == packet["lineage_parent"]
    assert agent_projection["result_count"] == len(packet["results"])
    # The bounded Agent projection can redact string labels after its public
    # string budget is consumed, so align rows to the producer packet while
    # requiring the exact numeric evidence to survive the projection.
    agent_anova = next(row for row in agent_projection["results"] if row.get("test_id") == anova["test_id"])
    agent_cohens = next(row for row in agent_projection["results"] if row.get("test_id") == cohens["test_id"])
    assert agent_anova["statistic"] == anova["statistic"]
    assert agent_anova["p_value"] == anova["p_value"]
    assert agent_cohens["statistic"] == cohens["statistic"]
    assert agent_cohens["p_value"] == cohens["p_value"]

    provider = NodeOperationContextProvider(Path(project_root))
    tool = next(
        definition
        for definition in provider.tool_definitions(
            chain_id="v186-statistics-agent-chain",
            session_id="v186-statistics-agent-session",
        )
        if definition.tool_id == "inspect_result_summary"
    )
    agent_result = tool.handler(
        {
            "request_id": "agent-v186-statistics",
            "owner_run_id": run_id,
            "op_node_id": "model:ols_1",
            "active_head_run_id": run_id,
        },
        ToolContext(session_id="v186-statistics-agent-session"),
    )
    agent_packet = agent_result["result_summary"]["statistical_evidence"]
    assert agent_packet["payload_schema"] == packet["payload_schema"]
    assert agent_packet["schema_version"] == packet["schema_version"]
    agent_anova = next(row for row in agent_packet["results"] if row.get("test_id") == anova["test_id"])
    agent_cohens = next(row for row in agent_packet["results"] if row.get("test_id") == cohens["test_id"])
    for field in ("statistic", "p_value", "p_value_corrected"):
        assert agent_anova[field] == anova[field]
        assert agent_cohens[field] == cohens[field]
