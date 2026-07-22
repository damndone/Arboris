from __future__ import annotations

import json
from pathlib import Path


def _write_artifact(run_root: Path, artifact_id: str, payload: object) -> None:
    path = run_root / "artifacts" / f"{artifact_id}.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({"payload": payload}), encoding="utf-8")
    index_path = run_root / "artifacts_index.json"
    index = json.loads(index_path.read_text(encoding="utf-8")) if index_path.exists() else {"artifacts": []}
    index["artifacts"].append({"artifact_id": artifact_id, "path": str(path.relative_to(run_root))})
    index_path.write_text(json.dumps(index), encoding="utf-8")


def test_arma_garch_deliverables_are_human_readable_and_nonempty(tmp_path: Path) -> None:
    """A time-series run must export its native evidence, never an empty coefficients sheet."""
    from workbench.engine.packs.arma_garch.deliverables import build_arma_garch_deliverables

    run_root = tmp_path / "run"
    run_root.mkdir()
    _write_artifact(run_root, "ts.analysis_contract", {
        "dataset_ref": "upload:VIXCLS.csv", "time_column": "observation_date",
        "value_column": "VIXCLS", "transform": "log_return_pct", "selection_mode": "manual",
        "estimation_strategy": "sequential", "innovation_distribution": "normal",
        "arma": {"p": 1, "q": 1, "constant_mode": "exclude"},
        "variance": {"model": "garch", "garch_p": 1, "garch_q": 1},
        "validation": {"validation_n": 20, "method": "expanding_window_one_step"},
    })
    _write_artifact(run_root, "ts.data_audit", {"time_index": {"original_row_count": 2610}, "data_quality": {"finite_value_count": 2541}})
    _write_artifact(run_root, "ts.arma_selection", {
        "final_selected_candidate_id": "arma-p1-q1-n",
        "final_selection_basis": "frozen_training_gates_aicc_bic_parsimony",
        "validation_data_role": "independent_evaluation_after_selection",
    })
    _write_artifact(run_root, "ts.volatility_selection", {
        "final_selected_candidate_id": "variance-garch-p1-q1-normal",
        "final_selection_basis": "frozen_training_gates_aicc_bic_parsimony",
    })
    _write_artifact(run_root, "ts.parameters", {
        "mean_candidate": {"ar.L1": 0.9208, "ma.L1": -0.9743, "sigma2": 62.223},
        "selected_variance_candidate": {"omega": 14.185, "alpha[1]": 0.1969, "beta[1]": 0.5784},
    })
    _write_artifact(run_root, "ts.forecast_metrics", {
        "validation_n": 20, "successful_forecast_n": 20, "mae": 6.4236,
        "rmse": 7.4409, "interval_coverage": 1.0, "average_interval_width": 30.1852,
    })
    _write_artifact(run_root, "ts.next_forecast", {
        "forecast_origin": {"time": "2025-12-12T00:00:00+00:00"},
        "conditional_mean": 0.6146, "conditional_variance": 48.4046,
        "conditional_volatility": 6.9573, "lower_bound": -13.0216, "upper_bound": 14.2507,
        "model_scale": "log_return_pct",
    })
    _write_artifact(run_root, "ts.final_diagnostics", {
        "adf": {"statistic": -20.99, "p_value": 0.0}, "warnings": [],
    })
    _write_artifact(run_root, "ts.arma_candidates", [{"candidate_id": "arma-p1-q1-n", "aic": 17711.38, "bic": 17728.90}])
    _write_artifact(run_root, "ts.volatility_candidates", [{"candidate_id": "variance-garch-p1-q1-normal", "aic": 17362.20, "bic": 17379.72}])
    _write_artifact(run_root, "ts.rolling_forecasts", [{"origin": 1, "actual": 2.0, "forecast": 1.5, "error": -0.5}])

    deliverables = build_arma_garch_deliverables(run_root)

    assert deliverables["report"]["title"] == "ARMA–GARCH Volatility Report"
    assert "ARMA(1,1)" in deliverables["report"]["model_label"]
    assert "GARCH(1,1)" in deliverables["report"]["model_label"]
    assert deliverables["report"]["metrics"]["rmse"] == 7.4409
    assert next(row for row in deliverables["tables"]["Overview"] if row["field"] == "Analysis observations")["value"] == 2541
    assert deliverables["report"]["limitations"]
    assert set(deliverables["tables"]) >= {
        "Overview", "ARMA candidates", "Volatility candidates", "Parameters",
        "Diagnostics", "Rolling forecasts", "Next forecast", "Acceptance",
    }
    assert all(rows for rows in deliverables["tables"].values())
    assert deliverables["tables"]["Parameters"][0]["component"] == "mean"


def test_arma_garch_report_and_exports_render_the_native_evidence(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    from workbench.engine.packs.arma_garch.deliverables import build_arma_garch_deliverables
    from workbench.exports import export_pdf, export_xlsx
    from workbench.reporting import render_html_report

    # Reuse the small, registered-artifact fixture above via the public builder
    # setup rather than creating a special report-only representation.
    run_root = tmp_path / "run"
    run_root.mkdir()
    _write_artifact(run_root, "ts.analysis_contract", {
        "dataset_ref": "upload:VIXCLS.csv", "time_column": "observation_date", "value_column": "VIXCLS",
        "transform": "log_return_pct", "arma": {"p": 1, "q": 1, "constant_mode": "exclude"},
        "variance": {"model": "garch", "garch_p": 1, "garch_q": 1},
    })
    _write_artifact(run_root, "ts.arma_selection", {"final_selected_candidate_id": "arma-p1-q1-n"})
    _write_artifact(run_root, "ts.volatility_selection", {"selected_candidate_id": "variance-garch-p1-q1-normal"})
    _write_artifact(run_root, "ts.parameters", {"mean_candidate": {"ar.L1": 0.9208}, "selected_variance_candidate": {"omega": 14.185}})
    _write_artifact(run_root, "ts.forecast_metrics", {"validation_n": 20, "rmse": 7.4409})
    _write_artifact(run_root, "ts.next_forecast", {"conditional_mean": 0.6146, "conditional_volatility": 6.9573})
    _write_artifact(run_root, "ts.final_diagnostics", {"adf": {"statistic": -20.99, "p_value": 0.0}})
    _write_artifact(run_root, "ts.arma_candidates", {"candidates": [{"candidate_id": "arma-p1-q1-n", "aic": 17711.38}]})
    _write_artifact(run_root, "ts.volatility_candidates", {"searches": [{"candidate_id": "variance-garch-p1-q1-normal", "aic": 17362.2}]})
    _write_artifact(run_root, "ts.rolling_forecasts", {"rows": [{"origin": 1, "actual": 2.0, "forecast": 1.5}]})

    deliverables = build_arma_garch_deliverables(run_root)
    html_path = render_html_report(deliverables["report"], run_root)
    pdf_path = export_pdf(deliverables["report"], run_root)
    xlsx_path = export_xlsx(deliverables["tables"], run_root)

    html = html_path.read_text(encoding="utf-8")
    assert 'id="time-series-overview"' in html
    assert "ARMA(1,1)" in html
    assert "7.4409" in html
    assert pdf_path.stat().st_size > 4_000
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    assert "Parameters" in workbook.sheetnames
    assert workbook["Parameters"].max_row > 1
