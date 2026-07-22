from __future__ import annotations

import json
import math
from pathlib import Path

import numpy as np
import pandas as pd

from workbench.engine.context import DataHandle, ModelingContext, RunEnv


class _Recorder:
    def __init__(self) -> None:
        self.nodes: list[dict[str, object]] = []
        self.edges: list[dict[str, object]] = []

    def record_stage(self, node_id, display_label, **kwargs) -> None:
        self.nodes.append({"node_id": node_id, "display_label": display_label, **kwargs})

    def record_edge(self, edge_id, source_id, target_id, op, **kwargs) -> None:
        self.edges.append(
            {
                "edge_id": edge_id,
                "source_id": source_id,
                "target_id": target_id,
                "op": op,
                **kwargs,
            }
        )


def _source(n: int = 130) -> pd.DataFrame:
    rng = np.random.default_rng(20260720)
    values = np.zeros(n)
    errors = np.zeros(n)
    variances = np.ones(n)
    for index in range(1, n):
        variances[index] = 0.1 + 0.12 * errors[index - 1] ** 2 + 0.80 * variances[index - 1]
        errors[index] = math.sqrt(variances[index]) * rng.normal()
        values[index] = 0.2 + 0.4 * values[index - 1] + errors[index]
    return pd.DataFrame(
        {"when": pd.bdate_range("2020-01-01", periods=n), "value": values}
    )


def _options() -> dict[str, object]:
    return {
        "dataset_ref": "dataset:pack-runtime:1",
        "time_column": "when",
        "value_column": "value",
        "time_index_semantics": "business_or_trading_observations",
        "transform": "level",
        "transform_confirmed": True,
        "selection_mode": "manual",
        "arma": {"p": 1, "q": 0, "constant_mode": "include"},
        "variance": {"model": "garch", "garch_p": 1, "garch_q": 1},
        "estimation_strategy": "joint",
        "innovation_distribution": "normal",
        "validation": {"validation_n": 4, "refit_every": 2},
        "random_seed": 20260720,
    }


def _context(source: pd.DataFrame) -> ModelingContext:
    cleaned = source.copy(deep=True)
    ctx = ModelingContext(
        data=DataHandle.of(cleaned, artifact_id="cleaned_dataset", provenance=("raw_input.csv",)),
        y_col="value",
        x_cols=[],
        y_type="continuous",
        requested_model_type="time_series.arma_garch",
    )
    ctx.artifacts.update(
        {
            "_model_options": _options(),
            "_frames": {"input.csv": source},
            "_upload_hash": "a" * 64,
            "_raw_inputs": ["raw_input.csv"],
        }
    )
    return ctx


def test_pack_declaration_exposes_one_generic_time_series_handler() -> None:
    from workbench.engine.capabilities import build_capabilities
    from workbench.engine.packs.loader import bootstrap_builtin_packs
    from workbench.engine.registry import MODEL_REGISTRY

    bootstrap_builtin_packs()

    handler = MODEL_REGISTRY["time_series.arma_garch"]
    assert handler.model_id == "arma_garch_1"
    assert handler.model_options_contract is not None
    capability = next(
        item
        for item in build_capabilities()["model_types"]
        if item["key"] == "time_series.arma_garch"
    )
    assert capability["group"] == "Time Series"
    assert {item["key"] for item in capability["params"]} >= {
        "model_type",
        "model_options",
    }


def test_runner_uses_raw_dataset_and_persists_traceable_artifacts_and_child(
    tmp_path: Path,
) -> None:
    from workbench.engine.packs.arma_garch.runner import fit_from_context

    source = _source()
    before = source.copy(deep=True)
    run_root = tmp_path / "run-1"
    run_root.mkdir()
    (run_root / "artifacts_index.json").write_text(
        json.dumps({"artifacts": []}), encoding="utf-8"
    )
    recorder = _Recorder()
    ctx = _context(source)
    env = RunEnv(run_root=run_root, run_id="run-1", recorder=recorder)

    model_id, result, fitted = fit_from_context(ctx, env)

    assert model_id == "arma_garch_1"
    assert fitted is None
    assert result["model_type"] == "time_series.arma_garch"
    assert result["estimation_strategy"] == "joint"
    assert result["joint_likelihood"] is True
    assert result["selection_repeated_during_validation"] is False
    assert result["validation"]["metrics"]["validation_n"] == 4
    assert result["validation"]["data_role"] == "independent_evaluation_after_selection"
    assert result["validation"]["candidate_selection_used_validation"] is False
    child = result["production_child"]
    assert child["result_role"] == "production_final_child"
    assert child["parent_result_id"] == "arma_garch_1"
    assert child["does_not_overwrite_parent"] is True
    assert child["next_forecast"]["fit_method"] == "full_sample_refit"
    refit = child["refit_summary"]
    assert refit["fit_scope"] == "all_available_observations"
    assert refit["nobs"] == len(source)
    assert refit["converged"] is True
    assert refit["finite_parameters"] is True
    assert refit["parameters"]
    assert refit["fit_statistics"]
    assert ctx.artifacts["_primary_model_input_ids"] == ["raw_input.csv"]
    assert ctx.artifacts["_primary_model_parent_node_id"] == "stage:ts-full-sample-child"
    pack_index = ctx.artifacts["_pack_node_index_entries"]
    assert set(pack_index) == {
        "stage:ts-analysis-view",
        "stage:ts-split",
        "stage:ts-mean-selection",
        "stage:ts-volatility-selection",
        "stage:ts-rolling-validation",
        "stage:ts-full-sample-child",
    }
    assert all(len(entry["node_hash"]) == 64 for entry in pack_index.values())
    pd.testing.assert_frame_equal(source, before)

    required = {
        "ts.data_audit",
        "ts.analysis_contract",
        "ts.analysis_view_manifest",
        "ts.transform_profile",
        "ts.train_validation_split",
        "ts.arma_candidates",
        "ts.arma_selection",
        "ts.arma_diagnostics",
        "ts.volatility_candidates",
        "ts.volatility_selection",
        "ts.final_model",
        "ts.parameters",
        "ts.conditional_series",
        "ts.final_diagnostics",
        "ts.rolling_forecasts",
        "ts.forecast_metrics",
        "ts.next_forecast",
        "ts.arma_vs_garch_comparison",
        "ts.report",
        "ts.artifact_manifest",
        "ts.chart.series_transform",
        "ts.chart.acf",
        "ts.chart.pacf",
        "ts.chart.residual_series",
        "ts.chart.residual_acf",
        "ts.chart.squared_residual_acf",
        "ts.chart.qq",
        "ts.chart.conditional_volatility",
        "ts.chart.rolling_interval",
        "ts.chart.quantile_exceptions",
        "ts.chart.model_comparison",
    }
    index = json.loads((run_root / "artifacts_index.json").read_text())
    records = {item["artifact_id"]: item for item in index["artifacts"]}
    assert required <= set(records)
    manifest = json.loads(
        (run_root / records["ts.artifact_manifest"]["path"]).read_text()
    )["payload"]
    assert manifest["status"] == "complete"
    assert manifest["complete"] is True
    for artifact_id in required - {"ts.artifact_manifest"}:
        payload = json.loads((run_root / records[artifact_id]["path"]).read_text())
        metadata = payload["metadata"]
        assert metadata["run_id"] == "run-1"
        assert metadata["dataset_hash"] == "a" * 64
        assert metadata["contract_hash"] == result["contract_hash"]
        assert metadata["analysis_view_hash"]
        assert metadata["node_id"]
        assert metadata["dependency_versions"]["arch"]
    expected_chart_nodes = {
        "ts.chart.series_transform": "stage:ts-analysis-view",
        "ts.chart.acf": "stage:ts-mean-selection",
        "ts.chart.pacf": "stage:ts-mean-selection",
        "ts.chart.residual_series": "stage:ts-mean-selection",
        "ts.chart.residual_acf": "stage:ts-mean-selection",
        "ts.chart.squared_residual_acf": "stage:ts-mean-selection",
        "ts.chart.qq": "stage:ts-volatility-selection",
        "ts.chart.conditional_volatility": "stage:ts-volatility-selection",
        "ts.chart.rolling_interval": "stage:ts-rolling-validation",
        "ts.chart.quantile_exceptions": "stage:ts-rolling-validation",
        "ts.chart.model_comparison": "stage:ts-rolling-validation",
    }
    for artifact_id, node_id in expected_chart_nodes.items():
        payload = json.loads((run_root / records[artifact_id]["path"]).read_text())
        assert payload["metadata"]["node_id"] == node_id
    json.dumps(result, allow_nan=False)

    node_ids = {item["node_id"] for item in recorder.nodes}
    assert {
        "stage:ts-analysis-view",
        "stage:ts-split",
        "stage:ts-mean-selection",
        "stage:ts-volatility-selection",
        "stage:ts-rolling-validation",
        "stage:ts-full-sample-child",
    } <= node_ids


def test_runner_blocks_when_generic_cleaning_hid_an_exact_duplicate(tmp_path: Path) -> None:
    from workbench.engine.packs.arma_garch.errors import ArmaGarchInputError
    from workbench.engine.packs.arma_garch.runner import fit_from_context

    source = _source()
    source = pd.concat([source.iloc[:1], source], ignore_index=True)
    cleaned = source.drop_duplicates().reset_index(drop=True)
    ctx = _context(source)
    ctx.data = DataHandle.of(
        cleaned, artifact_id="cleaned_dataset", provenance=("raw_input.csv",)
    )
    run_root = tmp_path / "run-duplicate"
    run_root.mkdir()
    (run_root / "artifacts_index.json").write_text(
        json.dumps({"artifacts": []}), encoding="utf-8"
    )

    try:
        fit_from_context(
            ctx,
            RunEnv(run_root=run_root, run_id="run-duplicate", recorder=_Recorder()),
        )
    except ArmaGarchInputError as error:
        assert error.code == "DUPLICATE_TIMESTAMP"
    else:
        raise AssertionError("the runner used the cleaned frame and hid a duplicate")
    index = json.loads((run_root / "artifacts_index.json").read_text())
    manifest_record = next(
        item for item in index["artifacts"] if item["artifact_id"] == "ts.artifact_manifest"
    )
    manifest = json.loads((run_root / manifest_record["path"]).read_text())["payload"]
    assert manifest["status"] == "blocked"
    assert manifest["terminal_code"] == "DUPLICATE_TIMESTAMP"
    assert manifest["complete"] is False


def test_public_workflow_completes_with_pack_graph_and_raw_lineage(
    tmp_path: Path,
) -> None:
    from workbench.graph_store import GraphStore
    from workbench.orchestrator import run_workflow
    from workbench.projects import create_project

    source_path = tmp_path / "series.csv"
    _source().to_csv(source_path, index=False)
    project = create_project(tmp_path, "arma-garch-project")

    outcome = run_workflow(
        project.root,
        [source_path],
        mode="auto",
        y="value",
        x=[],
        model_type="time_series.arma_garch",
        model_options=_options(),
    )

    assert outcome["status"] == "completed"
    run_root = project.root / "runs" / outcome["run_id"]
    result = json.loads((run_root / "model_results" / "arma_garch_1.json").read_text())
    assert result["model_type"] == "time_series.arma_garch"
    report_html = (run_root / "reports" / "report.html").read_text(encoding="utf-8")
    assert 'id="time-series-overview"' in report_html
    assert "ARMA(1,0)" in report_html
    assert "RMSE" in report_html
    from openpyxl import load_workbook
    workbook = load_workbook(run_root / "exports" / "tables.xlsx", read_only=True, data_only=True)
    assert {"Overview", "Parameters", "Acceptance", "Next forecast"}.issubset(workbook.sheetnames)
    assert workbook["Parameters"].max_row > 1
    index = json.loads((run_root / "artifacts_index.json").read_text())
    model_record = next(
        item for item in index["artifacts"] if item["artifact_id"] == "arma_garch_1"
    )
    assert model_record["inputs"] == ["raw_series.csv"]
    graph = GraphStore(project.root / "runs").read(outcome["run_id"])
    assert "stage:ts-full-sample-child" in graph.nodes
    model_edge = next(
        edge
        for edge in graph.edges.values()
        if edge.target_id == "model:arma_garch_1"
    )
    assert model_edge.source_id == "stage:ts-full-sample-child"


def test_public_workflow_preserves_duplicate_timestamp_error_code(
    tmp_path: Path,
) -> None:
    from workbench.orchestrator import run_workflow
    from workbench.projects import create_project

    source = _source()
    source = pd.concat([source.iloc[:1], source], ignore_index=True)
    source_path = tmp_path / "duplicate-series.csv"
    source.to_csv(source_path, index=False)
    project = create_project(tmp_path, "arma-garch-duplicate-project")

    outcome = run_workflow(
        project.root,
        [source_path],
        mode="auto",
        y="value",
        x=[],
        model_type="time_series.arma_garch",
        model_options=_options(),
    )

    assert outcome["status"] == "failed"
    errors = json.loads(
        (project.root / "runs" / outcome["run_id"] / "errors.json").read_text()
    )
    issue = next(item for item in errors["issues"] if item["code"] == "DUPLICATE_TIMESTAMP")
    assert issue["evidence"]["duplicate_timestamp_count"] == 2
    assert issue["evidence"]["impact"]
    run_root = project.root / "runs" / outcome["run_id"]
    tombstone = json.loads(
        (run_root / "model_results" / "arma_garch_1.tombstone.json").read_text()
    )
    assert tombstone["model_id"] == "arma_garch_1"
    assert tombstone["model_type"] == "time_series.arma_garch"
    assert tombstone["status"] == "blocked"
    assert tombstone["configured_not_fitted"] is True
    assert tombstone["terminal_code"] == "DUPLICATE_TIMESTAMP"
    from workbench.contracts.model.arma_garch import ArmaGarchAnalysisContract

    assert tombstone["contract_hash"] == ArmaGarchAnalysisContract.from_dict(
        tombstone["model_options"]
    ).contract_hash
    assert tombstone["model_options"]["missing_value_policy"] == "block"
    assert not ({"parameters", "converged", "forecast", "production_child"} & tombstone.keys())
    graph = __import__(
        "workbench.graph_store", fromlist=["GraphStore"]
    ).GraphStore(project.root / "runs").read(outcome["run_id"])
    failed_model = graph.nodes["model:arma_garch_1"]
    assert failed_model.trust.value == "blocker"
    assert failed_model.summary == "Configured, not fitted · DUPLICATE_TIMESTAMP"
    node_index = json.loads((run_root / "node_index.json").read_text())
    assert len(node_index["stage:raw"]["node_hash"]) == 64
    assert len(node_index["model:arma_garch_1"]["node_hash"]) == 64
    from workbench.lineage.node_write_validation import build_rerun_operation_context

    context = build_rerun_operation_context(
        project.root / "runs",
        request_id="failed-model-rerun",
        owner_run_id=outcome["run_id"],
        op_node_id="model:arma_garch_1",
        active_head_run_id=outcome["run_id"],
    )
    assert context.node_hash == node_index["model:arma_garch_1"]["node_hash"]


def test_runner_supports_sequential_student_t_without_composite_likelihood(
    tmp_path: Path,
) -> None:
    from workbench.engine.packs.arma_garch.runner import fit_from_context

    ctx = _context(_source())
    options = _options()
    options["estimation_strategy"] = "sequential"
    options["innovation_distribution"] = "student_t"
    ctx.artifacts["_model_options"] = options
    run_root = tmp_path / "sequential-student-t"
    run_root.mkdir()
    (run_root / "artifacts_index.json").write_text(
        json.dumps({"artifacts": []}), encoding="utf-8"
    )

    _, result, _ = fit_from_context(
        ctx,
        RunEnv(run_root=run_root, run_id="sequential-student-t", recorder=_Recorder()),
    )

    assert result["estimation_strategy"] == "sequential"
    assert result["joint_likelihood"] is False
    final_record = next(
        item
        for item in json.loads((run_root / "artifacts_index.json").read_text())["artifacts"]
        if item["artifact_id"] == "ts.final_model"
    )
    final_payload = json.loads((run_root / final_record["path"]).read_text())["payload"]
    validation_fit = final_payload["validation_fit"]
    assert validation_fit["innovation_distribution"] == "student_t"
    assert "mean_stage" in validation_fit
    assert "variance_stage" in validation_fit
    assert "composite_information_criterion" not in validation_fit


def test_runner_auto_mode_searches_a_bounded_grid_once(
    tmp_path: Path, monkeypatch,
) -> None:
    from workbench.engine.packs.arma_garch import runner

    rolling_calls = 0
    original_rolling = runner.run_rolling_validation

    def counted_rolling(*args, **kwargs):
        nonlocal rolling_calls
        rolling_calls += 1
        return original_rolling(*args, **kwargs)

    monkeypatch.setattr(runner, "run_rolling_validation", counted_rolling)

    ctx = _context(_source())
    options = _options()
    options.update(
        {
            "selection_mode": "auto",
            "arma": {
                "p": None,
                "q": None,
                "constant_mode": "exclude",
                "auto_max_p": 1,
                "auto_max_q": 1,
                "auto_max_total_order": 1,
            },
            "variance": {
                "model": "auto",
                "arch_p": None,
                "garch_p": None,
                "garch_q": None,
                "auto_arch_max_p": 2,
                "include_garch_1_1": True,
            },
            "estimation_strategy": "auto",
        }
    )
    ctx.artifacts["_model_options"] = options
    run_root = tmp_path / "bounded-auto"
    run_root.mkdir()
    (run_root / "artifacts_index.json").write_text(
        json.dumps({"artifacts": []}), encoding="utf-8"
    )

    _, result, _ = runner.fit_from_context(
        ctx,
        RunEnv(run_root=run_root, run_id="bounded-auto", recorder=_Recorder()),
    )

    assert result["selection_repeated_during_validation"] is False
    assert rolling_calls == 1
    assert result["selected_specification"]["mean_candidate_id"]
    assert result["selected_specification"]["variance_candidate_id"]
    candidate_record = next(
        item
        for item in json.loads((run_root / "artifacts_index.json").read_text())["artifacts"]
        if item["artifact_id"] == "ts.arma_candidates"
    )
    candidates = json.loads((run_root / candidate_record["path"]).read_text())["payload"]
    assert len(candidates["candidates"]) == 3
    assert candidates["selection_repeated_during_validation"] is False
    selection_record = next(
        item
        for item in json.loads((run_root / "artifacts_index.json").read_text())["artifacts"]
        if item["artifact_id"] == "ts.arma_selection"
    )
    selection = json.loads((run_root / selection_record["path"]).read_text())["payload"]
    assert selection["candidate_selection_used_validation"] is False
    assert selection["validation_data_role"] == "independent_evaluation_after_selection"


def test_runner_cooperatively_cancels_and_records_incomplete_manifest(
    tmp_path: Path,
) -> None:
    from workbench.engine.context import RunInterruptionRequested
    from workbench.engine.packs.arma_garch.runner import fit_from_context

    run_root = tmp_path / "cancelled"
    run_root.mkdir()
    (run_root / "artifacts_index.json").write_text(
        json.dumps({"artifacts": []}), encoding="utf-8"
    )
    ctx = _context(_source())
    env = RunEnv(
        run_root=run_root,
        run_id="cancelled",
        recorder=_Recorder(),
        stop_reason=lambda: "cancelled",
    )

    try:
        fit_from_context(ctx, env)
    except RunInterruptionRequested as error:
        assert error.reason == "cancelled"
    else:
        raise AssertionError("runner ignored the cooperative cancellation signal")

    index = json.loads((run_root / "artifacts_index.json").read_text())
    manifest_record = next(
        item for item in index["artifacts"] if item["artifact_id"] == "ts.artifact_manifest"
    )
    manifest = json.loads((run_root / manifest_record["path"]).read_text())["payload"]
    assert manifest["status"] == "interrupted"
    assert manifest["terminal_code"] == "WORKFLOW_CANCELLED"
    assert manifest["missing_required_artifact_ids"]
    assert {
        "ts.chart.series_transform",
        "ts.chart.acf",
        "ts.chart.pacf",
        "ts.chart.residual_series",
        "ts.chart.residual_acf",
        "ts.chart.squared_residual_acf",
        "ts.chart.qq",
        "ts.chart.conditional_volatility",
        "ts.chart.rolling_interval",
        "ts.chart.quantile_exceptions",
        "ts.chart.model_comparison",
    } <= set(manifest["missing_required_artifact_ids"])


def test_late_graph_failure_rewrites_complete_manifest_to_failed(
    tmp_path: Path, monkeypatch,
) -> None:
    from workbench.engine.packs.arma_garch import runner

    run_root = tmp_path / "late-graph-failure"
    run_root.mkdir()
    (run_root / "artifacts_index.json").write_text(
        json.dumps({"artifacts": []}), encoding="utf-8"
    )

    def fail_graph(*_args, **_kwargs):
        raise RuntimeError("injected graph write failure")

    monkeypatch.setattr(runner, "_record_graph", fail_graph)
    try:
        runner.fit_from_context(
            _context(_source()),
            RunEnv(run_root=run_root, run_id="late-graph-failure", recorder=_Recorder()),
        )
    except RuntimeError as error:
        assert str(error) == "injected graph write failure"
    else:
        raise AssertionError("fault injection did not reach the graph write")

    index = json.loads((run_root / "artifacts_index.json").read_text())
    manifest_records = [
        item for item in index["artifacts"] if item["artifact_id"] == "ts.artifact_manifest"
    ]
    assert len(manifest_records) == 1
    manifest = json.loads((run_root / manifest_records[0]["path"]).read_text())["payload"]
    assert manifest["status"] == "failed"
    assert manifest["terminal_code"] == "RuntimeError"
    assert manifest["complete"] is False
    assert manifest["persisted_artifact_ids"]
    assert "ts.artifact_manifest" not in manifest["persisted_artifact_ids"]
    envelope = json.loads((run_root / manifest_records[0]["path"]).read_text())
    assert envelope["metadata"]["node_id"] == "stage:ts-full-sample-child"
    assert envelope["metadata"]["contract_hash"]


def test_late_cancellation_rewrites_complete_manifest_to_interrupted(
    tmp_path: Path,
) -> None:
    from workbench.engine.context import RunInterruptionRequested
    from workbench.engine.packs.arma_garch.runner import fit_from_context

    run_root = tmp_path / "late-cancellation"
    run_root.mkdir()
    (run_root / "artifacts_index.json").write_text(
        json.dumps({"artifacts": []}), encoding="utf-8"
    )
    recorder = _Recorder()
    env = RunEnv(
        run_root=run_root,
        run_id="late-cancellation",
        recorder=recorder,
        stop_reason=lambda: "cancelled" if recorder.nodes else None,
    )
    try:
        fit_from_context(_context(_source()), env)
    except RunInterruptionRequested as error:
        assert error.reason == "cancelled"
    else:
        raise AssertionError("late cancellation was not observed")

    index = json.loads((run_root / "artifacts_index.json").read_text())
    record = next(
        item for item in index["artifacts"] if item["artifact_id"] == "ts.artifact_manifest"
    )
    manifest = json.loads((run_root / record["path"]).read_text())["payload"]
    assert manifest["status"] == "interrupted"
    assert manifest["terminal_code"] == "WORKFLOW_CANCELLED"
    assert manifest["complete"] is False
