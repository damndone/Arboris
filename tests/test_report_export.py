from __future__ import annotations

import io
import zipfile
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from workbench.app import app
from workbench.artifacts import write_json
from workbench.projects import create_project, create_run
from workbench.report_export import export_report


def _fixture(tmp_path: Path) -> tuple[Path, str]:
    project = create_project(tmp_path, "demo")
    run = create_run(project.root, mode="auto")
    (run.root / "figures").mkdir(parents=True, exist_ok=True)
    (run.root / "figures" / "coef_plot.png").write_bytes(b"\x89PNG fake")
    write_json(
        run.root / "artifacts_index.json",
        {
            "schema_version": 1,
            "artifacts": [
                {
                    "artifact_id": "coef_plot",
                    "path": "figures/coef_plot.png",
                    "artifact_type": "figure",
                    "step": "visualization",
                    "sha256": "stub",
                    "inputs": [],
                }
            ],
        },
    )
    return project.root, run.run_id


def test_report_export_embeds_figures_in_html_and_print_html(tmp_path: Path) -> None:
    project_root, run_id = _fixture(tmp_path)
    markdown = "# Results\n\nThe effect is **positive**.\n\n[[fig:coef_plot]]"
    figures = [{"artifact_id": "coef_plot", "chart_type": "Coefficient plot"}]

    run_root = project_root / "runs" / run_id
    html, media, filename = export_report(
        run_root, markdown=markdown, figures=figures, format="html"
    )
    text = html.decode("utf-8")
    assert media == "text/html"
    assert filename == "report.html"
    assert "data:image/png;base64" in text
    assert "[[fig:" not in text
    assert "Coefficient plot" in text

    printed, _, print_name = export_report(
        run_root, markdown=markdown, figures=figures, format="pdf-print"
    )
    assert print_name == "report-print.html"
    assert b"@media print" in printed


def test_report_export_produces_docx_and_latex_zip_without_new_dependencies(tmp_path: Path) -> None:
    project_root, run_id = _fixture(tmp_path)
    run_root = project_root / "runs" / run_id
    markdown = "# Results\n\nThe effect is **positive** [[fig:coef_plot]].\n\n- one\n- two"
    figures = [{"artifact_id": "coef_plot", "chart_type": "Coefficient plot"}]

    docx, media, _ = export_report(
        run_root, markdown=markdown, figures=figures, format="docx"
    )
    assert media.startswith("application/vnd.openxmlformats")
    with zipfile.ZipFile(io.BytesIO(docx)) as archive:
        names = set(archive.namelist())
        assert "word/document.xml" in names
        assert "word/media/image1.png" in names
        document = archive.read("word/document.xml")
        assert b"Results" in document
        assert b"<w:b/>" in document
        assert b"**positive**" not in document
        assert b"<w:drawing>" in document
        assert "• one".encode() in document

    tex_zip, media, filename = export_report(
        run_root, markdown=markdown, figures=figures, format="tex"
    )
    assert media == "application/zip"
    assert filename == "report.tex.zip"
    with zipfile.ZipFile(io.BytesIO(tex_zip)) as archive:
        assert "report.tex" in archive.namelist()
        assert "figures/coef_plot.png" in archive.namelist()
        tex = archive.read("report.tex").decode("utf-8")
        assert "includegraphics" in tex
        assert "[[fig:" not in tex


def test_report_export_route_rejects_unknown_figure_and_returns_file(tmp_path: Path) -> None:
    project_root, run_id = _fixture(tmp_path)
    client = TestClient(app)
    response = client.post(
        f"/runs/{run_id}/report/export",
        params={"project_root": str(project_root), "format": "html"},
        json={"markdown": "# Report\n\n[[fig:coef_plot]]", "figures": [{"artifact_id": "coef_plot"}]},
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/html")
    assert b"data:image/png;base64" in response.content

    missing = client.post(
        f"/runs/{run_id}/report/export",
        params={"project_root": str(project_root), "format": "html"},
        json={"markdown": "[[fig:missing]]", "figures": [{"artifact_id": "missing"}]},
    )
    assert missing.status_code == 422
    assert missing.json()["error"]["code"] == "REPORT_EXPORT_INVALID"


def test_persisted_report_export_is_bound_to_the_stored_revision(tmp_path: Path) -> None:
    project_root, run_id = _fixture(tmp_path)
    client = TestClient(app)
    saved = client.post(
        f"/runs/{run_id}/ai-reports",
        params={"project_root": str(project_root)},
        json={
            "id": "rpt_export1",
            "generatedAt": "2026-07-21T22:17:30Z",
            "instruction": "Write report",
            "text": "# Stored report\n\nThe result is [[c:c1]].\n\n[[fig:coef_plot]]",
            "scope": {"run_id": run_id, "node_count": 1, "node_keys": ["model:ols"]},
            "facts": [{"id": "c1", "value": 1.25}],
            "excluded_fact_ids": [],
            "figures": [{
                "artifact_id": "coef_plot",
                "chart_type": "Coefficient plot",
                "path": "figures/coef_plot.png",
                "source": {"artifact_id": "ols_1", "kind": "model"},
            }],
        },
    )
    assert saved.status_code == 200, saved.text

    exported = client.post(
        f"/runs/{run_id}/report/export",
        params={"project_root": str(project_root), "format": "html"},
        json={
            "report_id": "rpt_export1",
            "markdown": "# Stored report\n\nThe result is [[c:c1]].\n\n[[fig:coef_plot]]",
            "figures": [{"artifact_id": "coef_plot", "chart_type": "Coefficient plot"}],
        },
    )
    assert exported.status_code == 200, exported.text
    assert b"Stored report" in exported.content
    assert b"data:image/png;base64" in exported.content

    mismatch = client.post(
        f"/runs/{run_id}/report/export",
        params={"project_root": str(project_root), "format": "html"},
        json={
            "report_id": "rpt_export1",
            "markdown": "# Tampered report\n\nThe result is [[c:c1]].",
            "figures": [],
        },
    )
    assert mismatch.status_code == 409
    assert mismatch.json()["error"]["code"] == "AI_REPORT_SOURCE_MISMATCH"

    missing = client.post(
        f"/runs/{run_id}/report/export",
        params={"project_root": str(project_root), "format": "html"},
        json={
            "report_id": "rpt_missing1",
            "markdown": "# Stored report",
            "figures": [],
        },
    )
    assert missing.status_code == 404
    assert missing.json()["error"]["code"] == "AI_REPORT_NOT_FOUND"


def test_result_table_export_route_reads_authoritative_run_evidence(tmp_path: Path) -> None:
    project_root, run_id = _fixture(tmp_path)
    run_root = project_root / "runs" / run_id
    (run_root / "model_results").mkdir(parents=True, exist_ok=True)
    write_json(
        run_root / "model_results" / "ols_1.json",
        {
            "model_id": "ols_1",
            "model_type": "ols",
            "nobs": 12,
            "coefficients": {
                "x": {
                    "estimate": 1.25,
                    "std_error": 0.2,
                    "p_value": 0.04,
                    "source_id": "model_results.ols_1.coefficients.x",
                }
            },
        },
    )
    write_json(
        run_root / "diagnostic_summary.json",
        {
            "table_1": [{"column": "y", "count": 12}],
            "statistical_evidence": {
                "results": [{"test_type": "shapiro_wilk", "p_value": 0.3}],
            },
            "labels": {"variable_labels": {"x": "Treatment"}},
            "model_family_evidence": None,
        },
    )

    response = TestClient(app).post(
        f"/runs/{run_id}/report/result-table-export",
        params={"project_root": str(project_root)},
        json={"sections": ["regression_table", "statistical_evidence"]},
    )

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert {"regression_table", "statistical_evidence"} <= set(workbook.sheetnames)
    rows = list(workbook["regression_table"].iter_rows(values_only=True))
    assert "ols_1 estimate" in rows[0]
    assert any("model_results.ols_1.coefficients.x" in row for row in rows)
    assert not (run_root / "reports" / "report.pdf").exists()
    assert not (run_root / "exports" / "tables.xlsx").exists()

    rejected = TestClient(app).post(
        f"/runs/{run_id}/report/result-table-export",
        params={"project_root": str(project_root)},
        json={"sections": ["regression_table"], "rows": [{"estimate": 999}]},
    )
    assert rejected.status_code == 422
