from __future__ import annotations

import hashlib
import io
import json
import copy
import pickle
import atexit
from contextlib import contextmanager
from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4
from collections.abc import Iterator

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from workbench.canonical import sha256_canonical
from workbench.contracts.common.contract_ref import ContractRefV1
from workbench.contracts.common.foundation_status import TransactionStatusV1
from workbench.contracts.evidence.envelope import EvidenceEnvelopeV1, NotApplicableV1
from workbench.contracts.evidence.persistence import (
    EvidenceCommitBindingV1,
    EvidenceCommitRequestV1,
)
from workbench.evidence.consumer_adapters import (
    EvidenceProjectionError,
    PersistedEvidenceEnvelopeV1,
    default_projection_registry,
)
from workbench.evidence.contracts import ArtifactAdmissionV1
from workbench.api import app
from workbench.app import configure_report_evidence_owner
from workbench.http import runs_routes
from workbench.report_view_model import (
    build_indexed_report_evidence,
    build_persisted_evidence_report_view,
    build_report_view_model,
    build_report_evidence_bundle,
)
from workbench.services.results_service import (
    PersistedEvidenceResultsFacadeV1,
    ReportEvidenceRequestIdentityV1,
    ServerOwnedReportEvidencePersistenceV1,
    UnavailableReportEvidenceCapabilityV1,
)
from workbench.report_export import EvidenceProjectionV1, ReportExportError, ReportEvidenceWriterV1
from workbench import reporting
from workbench.reporting import ReportDocumentWriterV1, render_legacy_html_report


_ROOT = Path(__file__).parents[1]
_FIXTURE = _ROOT / "tests" / "fixtures" / "contracts" / "v189" / "evidence_envelope.json"
_LIVE_SERVER_EVIDENCE_CONTEXTS: list[object] = []
_CONSUMER_PROJECTION_EVIDENCE_DIGEST = "a80048ec382239a8e8050f1219322201591816828f1ccf8f12570f38a45c7361"


def _close_live_server_evidence_contexts() -> None:
    while _LIVE_SERVER_EVIDENCE_CONTEXTS:
        context = _LIVE_SERVER_EVIDENCE_CONTEXTS.pop()
        context.__exit__(None, None, None)  # type: ignore[attr-defined]


atexit.register(_close_live_server_evidence_contexts)


def _ref(kind: str, logical_id: str) -> ContractRefV1:
    return ContractRefV1(
        kind=kind,
        logical_id=logical_id,
        semantic_version="1.0",
        contract_digest=sha256_canonical(
            {"kind": kind, "logical_id": logical_id, "semantic_version": "1.0"}
        ),
    )


def _persisted_record() -> PersistedEvidenceEnvelopeV1:
    envelope = EvidenceEnvelopeV1.from_dict(json.loads(_FIXTURE.read_text(encoding="utf-8")))
    projector_ref = _ref("projection", "workbench.consumer_projection.results")
    request = EvidenceCommitRequestV1(
        artifact_admission_ref="artifact:admitted:result-1",
        projector_ref=projector_ref,
        envelope_digest=sha256_canonical(envelope.to_dict()),
        idempotency_key="evidence-commit:result-1",
        supersedes_ref=NotApplicableV1("no supersession"),
        invalidation_refs=(),
    )
    binding = EvidenceCommitBindingV1(
        commit_id="evidence-commit-1",
        request_digest=request.request_digest,
        artifact_admission_ref=request.artifact_admission_ref,
        projector_ref=projector_ref,
        envelope_digest=request.envelope_digest,
        idempotency_key=request.idempotency_key,
        transaction_status=TransactionStatusV1.COMMITTED,
        supersedes_ref=request.supersedes_ref,
        invalidation_refs=(),
        failure_receipt=NotApplicableV1("not failed"),
    )
    return PersistedEvidenceEnvelopeV1(
        commit_id=binding.commit_id,
        artifact_ref="artifact-1",
        envelope_ref="evidence-envelope:result-1",
        sensitivity="public",
        binding=binding,
        envelope=envelope,
    )


def _indexed_artifact(**updates: object) -> dict[str, object]:
    return {
        "artifact_id": "workflow_model_quadratic_stationary_point_1",
        "artifact_type": "post_estimation",
        "path": "artifacts/post_estimation/result.json",
        "sha256": "a" * 64,
        "step": "locate_stationary_point",
        **updates,
    }


def _canonical_indexed_artifact(
    record: ArtifactAdmissionV1,
    binding: EvidenceCommitBindingV1,
    *,
    artifact_id: str | None = None,
    run_id: str = "run-1",
    model_run_id: str = "run-1",
    project_ref: str = "project:1",
    session_ref: str = "session:1",
    scope_ref: str = "scope:project:1",
    raw_payload: bytes | None = None,
) -> tuple[dict[str, object], dict[str, object], str]:
    resolved_artifact_id = artifact_id or record.artifact_id
    payload: dict[str, object] = {
        "source": {"run_id": run_id, "model_run_id": model_run_id},
        "result": {"committed": True},
    }
    content = raw_payload or json.dumps(
        payload,
        separators=(",", ":"),
        sort_keys=True,
    ).encode("utf-8")
    index = _indexed_artifact(
        artifact_id=resolved_artifact_id,
        artifact_type=record.artifact_type,
        sha256=hashlib.sha256(content).hexdigest(),
        run_id=run_id,
        model_run_id=model_run_id,
        project_ref=project_ref,
        session_ref=session_ref,
        scope_ref=scope_ref,
        content_hash=record.content_hash,
        evidence_content_hash=record.content_hash,
        actor_ref="actor:1",
        confirmation_ref=f"confirmation:{run_id}",
        evidence_commit_ref=binding.commit_id,
        evidence_commit_binding=binding.to_dict(),
        evidence_envelope=record.envelope.to_dict(),
        evidence_envelope_ref=f"evidence-envelope:{record.envelope_digest}",
        evidence_artifact_ref=resolved_artifact_id,
        evidence_artifact_admission_ref=record.artifact_id,
        evidence_run_id=run_id,
        evidence_model_run_id=model_run_id,
        source_ref=record.source_ref,
        validator_ref=(
            record.validator_ref.to_dict()
            if hasattr(record.validator_ref, "to_dict")
            else record.validator_ref
        ),
        execution_ref=record.execution_ref,
        contract_digest=record.contract_digest,
    )
    return index, payload, hashlib.sha256(content).hexdigest()


@contextmanager
def _server_owned_evidence(
    tmp_path: Path,
    *,
    run_id: str = "run-1",
    model_run_id: str = "run-1",
    canonical_project_root: Path | None = None,
    keep_open: bool = False,
    envelope_payload: dict[str, object] | None = None,
    with_project_file_store: bool = False,
) -> Iterator[tuple[ArtifactAdmissionV1, EvidenceCommitBindingV1, ServerOwnedReportEvidencePersistenceV1]]:
    """Build the same C0 registry/store/receipt capability as the owner seam."""

    from tests.test_v189_evidence_envelope import _admit, _admit_store, _authorized_registry

    tmp_path.mkdir(parents=True, exist_ok=True)
    registry, host = _authorized_registry("scope:project:1")
    record = registry.admit_artifact(
        artifact_id="artifact-1",
        artifact_type="post_estimation",
        scope_ref="scope:project:1",
        content=(
            envelope_payload
            if envelope_payload is not None
            else json.loads(_FIXTURE.read_text(encoding="utf-8"))
        ),
        source_ref="source:dataset:1",
        validator_ref="validator:evidence-envelope:1",
        execution_ref=f"execution:{run_id}",
        contract_digest=_CONSUMER_PROJECTION_EVIDENCE_DIGEST,
    )
    store_context = _admit_store(tmp_path, registry)
    store = store_context.__enter__()
    if keep_open:
        _LIVE_SERVER_EVIDENCE_CONTEXTS.append(store_context)
    project_files = None
    try:
        if with_project_file_store:
            from workbench.services.report_evidence_owner_factory import HostBoundProjectFileStoreV1
            project_files = HostBoundProjectFileStoreV1.from_host(
                canonical_project_root, authority_handle=host.channel, scope_ref="scope:project:1")
        receipt = registry.persistence_receipt(record)
        binding = store.commit(record, receipt=receipt)
        owner_payload = {
            "source": {"run_id": run_id, "model_run_id": model_run_id},
            "result": {"committed": True},
        }
        owner_raw = json.dumps(
            owner_payload,
            separators=(",", ":"),
            sort_keys=True,
        ).encode("utf-8")
        capability = ServerOwnedReportEvidencePersistenceV1._from_host(
            store,
            registry,
            receipt,
            actor_ref="actor:1",
            project_ref="project:1",
            session_ref="session:1",
            scope_ref="scope:project:1",
            run_id=run_id,
            model_run_id=model_run_id,
            artifact_id=record.artifact_id,
            commit_id=binding.commit_id,
            confirmation_ref=f"confirmation:{run_id}",
            artifact_file_sha256=hashlib.sha256(owner_raw).hexdigest(),
            canonical_project_root=canonical_project_root,
            authority_handle=host.channel,
            project_file_store=project_files,
        )
        yield record, binding, capability
    finally:
        if project_files is not None:
            project_files.close()
        if not keep_open:
            store_context.__exit__(None, None, None)


def _report_projection_wire() -> dict[str, object]:
    return {
        "schemaVersion": 1,
        "projectionId": "report:evidence-commit-1",
        "consumerId": "report",
        "consumerVersion": "1.0",
        "actorRef": "actor:1",
        "projectRef": "project:1",
        "sessionRef": "session:1",
        "scopeRef": "scope:project:1",
        "runId": "run-1",
        "modelRunId": "run-1",
        "artifactRef": "artifact-1",
        "commitId": "evidence-commit-1",
        "confirmationRef": "confirmation:run-1",
        "envelopeRef": "evidence-envelope:1",
        "envelopeDigest": "a" * 64,
        "status": "available",
        "executionStatus": "completed",
        "availability": "available",
        "completenessStatus": "complete",
        "validationStatus": "met",
        "validationSummary": {"requirement_status": "met"},
        "facts": [{"factId": "fact-1", "statisticId": "estimate", "value": 1.25}],
        "sections": [],
        "provenance": {"artifactRef": "artifact-1", "commitId": "evidence-commit-1"},
    }


def _owner_wire(owner: ServerOwnedReportEvidencePersistenceV1, wire: dict[str, object]) -> dict[str, object]:
    del wire
    return dict(owner._canonical_report_wire())  # type: ignore[attr-defined]


def _owner_issued_projection(tmp_path: Path):
    with _server_owned_evidence(tmp_path / "owner-store", keep_open=True) as (record, _binding, owner):
        wire = _owner_wire(owner, _report_projection_wire())
        wire["envelopeDigest"] = record.envelope_digest
        return owner.issue_report_projection(wire)


@contextmanager
def _produced_report(run_root, *, model_results, diagnostic_summary, report_fields=None, variable_labels=None,
                     complete_collection=False):
    """Use real producer/host admission; never append display fields after issuance."""
    from types import SimpleNamespace
    from workbench.artifacts import initialize_artifact_index, register_artifact, write_json
    from workbench.report_view_model import build_regression_table, build_report_view_model
    from workbench.services.report_evidence_owner_factory import (
        ServerOwnedReportEvidenceOwnerFactoryV1, register_report_evidence_artifact,
    )
    from tests.test_v189_evidence_envelope import _admit_store, _authorized_registry

    run_root.mkdir(parents=True, exist_ok=True)
    initialize_artifact_index(run_root)
    for model_id, packet in model_results.items():
        path = run_root / "model_results" / f"{model_id}.json"
        write_json(path, packet)
        register_artifact(run_root, model_id, path, "model_result", "econometrics", [])
    registry, host = _authorized_registry("scope:project:report-fixture")
    store_root = run_root.parent.parent / f"evidence-store-{uuid4().hex}"
    store_root.mkdir()
    with _admit_store(store_root, registry) as store:
        factory = ServerOwnedReportEvidenceOwnerFactoryV1.from_host(store, registry,
            actor_ref="actor:report-fixture", project_ref="project:report-fixture",
            session_ref="session:report-fixture", scope_ref="scope:project:report-fixture",
            canonical_project_root=run_root.parent.parent, authority_handle=host.channel)
        try:
            primary = next(iter(model_results))
            envelope = factory.prepare_report_evidence_envelope(
                run_id=run_root.name, model_run_id=primary, upload_hash="a" * 64,
                primary_model_type=model_results[primary]["model_type"],
                regression_table=build_regression_table(list(model_results.items()), variable_labels=variable_labels),
                model_artifact_ids=list(model_results), diagnostic_summary=diagnostic_summary,
                report_fields=report_fields or {}, issue_dicts=[])
            receipt = factory.issue_production_receipt(run_id=run_root.name, model_run_id=primary, upload_hash="a" * 64)
            owner = factory.create_for_run(run_id=run_root.name, model_run_id=primary,
                production_receipt=receipt, envelope=envelope,
                artifact_file_sha256=factory.report_evidence_file_sha256(envelope))
            factory.write_report_evidence_file(owner=owner, production_receipt=receipt, envelope=envelope)
            register_report_evidence_artifact(owner, run_root=run_root,
                path=run_root / "report_evidence.json", inputs=list(model_results))
            if complete_collection:
                issue = getattr(owner, "issue_complete_report_collection", None)
                assert callable(issue), "owner must issue a dedicated complete local collection"
                projection = issue()
            else:
                projection = owner.issue_report_projection(owner._canonical_report_wire())
            view = build_report_view_model({}, run_root, evidence_projection=projection,
                owner=owner, artifact_admission=owner._canonical_report_artifact())
            writer = owner.issue_report_document_writer(view, projection, canonical_run_root=run_root)
            yield SimpleNamespace(view=view, writer=writer, owner=owner, projection=projection)
        finally:
            factory.close()


def _collection_model(count=1):
    return {"ols_1": {"model_id": "ols_1", "model_type": "ols", "model_label": "Collection OLS",
        "nobs": 30, "coefficients": {
            f"term_{i}": {"estimate": 1.25 + i, "std_error": 0.2, "p_value": 1e-24}
            for i in range(count)}}}


@pytest.mark.parametrize("replace_existing", [False, True])
def test_export_withdraws_index_under_persistent_fsync_eio(tmp_path, monkeypatch, replace_existing):
    import errno
    import os
    import stat

    from workbench.evidence.persistence import EvidencePersistenceError
    from workbench.reporting import render_html_report

    run = tmp_path / "project" / "runs" / "persistent-eio"
    with _produced_report(run, model_results=_collection_model(), diagnostic_summary={},
                          complete_collection=True) as case:
        if replace_existing:
            render_html_report(case.writer, run)
        index_path = run / "artifacts_index.json"
        original_index = json.loads(index_path.read_bytes())
        expected = {**original_index, "artifacts": [row for row in original_index["artifacts"]
            if row["artifact_id"] != "report_html" and row["path"] != "reports/report.html"]}
        original_fsync, original_open, original_replace = os.fsync, os.open, os.replace
        failed = False
        file_syncs_after_failure = []

        def persistent_eio(fd):
            nonlocal failed
            info = os.fstat(fd)
            if failed and stat.S_ISREG(info.st_mode):
                file_syncs_after_failure.append(info.st_ino)
            if not failed and (info.st_dev, info.st_ino) == (run.stat().st_dev, run.stat().st_ino):
                failed = any(row["artifact_id"] == "report_html"
                    for row in json.loads(index_path.read_bytes())["artifacts"])
            if failed:
                raise OSError(errno.EIO, "persistent export fsync failure")
            return original_fsync(fd)

        def fd_only_open(path, flags, *args, **kwargs):
            if flags & (os.O_WRONLY | os.O_RDWR | os.O_CREAT):
                assert isinstance(kwargs.get("dir_fd"), int), "export writes must remain FD-only"
                assert not Path(path).is_absolute()
            return original_open(path, flags, *args, **kwargs)

        def fd_only_replace(src, dst, *args, **kwargs):
            assert isinstance(kwargs.get("src_dir_fd"), int)
            assert isinstance(kwargs.get("dst_dir_fd"), int)
            assert not Path(src).is_absolute() and not Path(dst).is_absolute()
            return original_replace(src, dst, *args, **kwargs)

        with monkeypatch.context() as faults:
            faults.setattr(os, "fsync", persistent_eio)
            faults.setattr(os, "open", fd_only_open)
            faults.setattr(os, "replace", fd_only_replace)
            with pytest.raises(EvidencePersistenceError) as failure:
                render_html_report(case.writer, run)
        assert failure.value.code == "EVIDENCE_PERSISTENCE_WRITE_FAILED"
        assert failed, "inject only after the complete export index was renamed into place"
        assert json.loads(index_path.read_bytes()) == expected, "failed export must be withdrawn without losing other records"
        assert file_syncs_after_failure == [], "withdrawal must not fsync a new staged file"
        assert not list(run.rglob("*.tmp"))
        assert not list(run.glob("*.rollback"))


def test_export_withdraws_index_if_rollback_snapshot_cleanup_fails(tmp_path, monkeypatch):
    import errno
    import os

    from workbench.reporting import render_html_report

    run = tmp_path / "project" / "runs" / "rollback-cleanup"
    with _produced_report(run, model_results=_collection_model(), diagnostic_summary={},
                          complete_collection=True) as case:
        index_path = run / "artifacts_index.json"
        original_index = json.loads(index_path.read_bytes())
        original_unlink = os.unlink
        failed = False

        def fail_snapshot_cleanup(path, *args, **kwargs):
            nonlocal failed
            if str(path).endswith(".rollback") and not failed:
                assert isinstance(kwargs.get("dir_fd"), int)
                assert any(row["artifact_id"] == "report_html"
                    for row in json.loads(index_path.read_bytes())["artifacts"])
                failed = True
                raise OSError(errno.EIO, "rollback snapshot cleanup failure")
            return original_unlink(path, *args, **kwargs)

        with monkeypatch.context() as faults:
            faults.setattr(os, "unlink", fail_snapshot_cleanup)
            with pytest.raises(OSError, match="rollback snapshot cleanup failure"):
                render_html_report(case.writer, run)
        assert failed
        assert json.loads(index_path.read_bytes()) == original_index
        assert not list(run.glob("*.rollback"))


def test_local_collection_exports_full_oversized_content(tmp_path, monkeypatch):
    from workbench import report_export
    from workbench.evidence.consumer_adapters import ConsumerProjectionRegistryV1, _structured_row_count
    from workbench.canonical import canonical_json_v1
    from workbench.reporting import render_html_report
    from workbench.exports import export_xlsx, export_pdf
    cls = getattr(report_export, "LocalReportCollectionV1", None)
    assert cls is not None, "complete collections must not masquerade as EvidenceProjectionV1"
    run = tmp_path / "project" / "runs" / "r1"
    with _produced_report(run, model_results=_collection_model(100), diagnostic_summary={},
                          variable_labels={"term_99": "Final labelled term"},
                          complete_collection=True) as case:
        assert type(case.projection) is cls
        assert not isinstance(case.projection, report_export.EvidenceProjectionV1)
        wire = case.projection.wire
        assert wire["deliveryKind"] == "local_complete_collection"
        assert wire["readCompleteness"] == "COMPLETE" and wire["sectionDirectory"]
        assert len(case.view["regression_table"]["rows"]) == 100
        assert case.view["regression_table"]["rows"][-1]["label"] == "Final labelled term"
        assert all(section["status"] in {"available", "unavailable", "not_applicable", "blocked"}
                   for section in wire["resolvedSections"])
        source, _ = case.owner._canonical_report_source_and_artifact()
        single = default_projection_registry().project(source, consumer_id="report", destination="local")
        assert single.status.value == "unavailable", "fixture must exceed single delivery budget"
        assert single.audit.restriction_reason == "structured_projection_budget_exceeded"
        central = case.projection._collection
        assert _structured_row_count(central.facts, central.sections, central.result_dependencies) <= 1024
        original = json.loads((run / "report_evidence.json").read_text())
        assert canonical_json_v1([fact.to_dict() for fact in central.facts]) == canonical_json_v1(original["numeric_facts"])
        assert canonical_json_v1(central.result_dependencies) == canonical_json_v1(original["family_evidence"]["result_dependencies"])
        assert {s["sectionId"]: (s["status"], list(s["reasonCodes"])) for s in wire["resolvedSections"]} == {
            s["section_id"]: (s["status"], s["reason_codes"])
            for s in original["family_evidence"]["structured_results"]["sections"]}
        def no_default_all(*args, **kwargs):
            pytest.fail("complete collection rendering must not reproject default-all")
        monkeypatch.setattr(ConsumerProjectionRegistryV1, "project", no_default_all)
        view = build_report_view_model({}, run, evidence_projection=case.projection,
            owner=case.owner, artifact_admission=case.owner._canonical_report_artifact())
        assert view["status"] == "available" and view["facts"] == case.view["facts"]
        html = render_html_report(case.writer, run).read_text()
        assert "term_99" in html and "1e-24" in html and "Final labelled term" in html
        pdf = export_pdf(case.writer, run).read_bytes()
        assert b"Collection OLS" in pdf and b"1e-24" in pdf and b"Final labelled term" in pdf
        writer = case.owner.issue_report_evidence_writer(case.projection, sections=("regression_table",))
        workbook = load_workbook(export_xlsx(writer, run), data_only=False)
        try:
            rows = list(workbook.active.iter_rows(values_only=True))
            assert len(rows) == 101
            assert rows[-1][rows[0].index("label")] == "Final labelled term"
            cell = json.loads(rows[-1][rows[0].index("models")])["ols_1"]
            assert cell["estimate"] == 100.25 and cell["p_value"] == 1e-24
        finally:
            workbook.close()


def test_local_collection_preserves_locked_section_row_cap(tmp_path):
    from workbench.evidence.consumer_adapters import EvidenceProjectionError
    with pytest.raises(EvidenceProjectionError, match="RESULT_PAGE_ROW_BUDGET_EXCEEDED"):
        with _produced_report(tmp_path / "project" / "runs" / "over-row-cap",
                              model_results=_collection_model(180), diagnostic_summary={},
                              complete_collection=True):
            pytest.fail("over-budget section was issued as complete")


@pytest.mark.parametrize("attack", ["copy", "pickle", "forged", "wire", "directory", "incomplete", "source", "view"])
def test_local_collection_rejects_forgery_and_changes_before_export(tmp_path, attack):
    from workbench import report_export
    from workbench.reporting import render_html_report
    cls = getattr(report_export, "LocalReportCollectionV1", None)
    assert cls is not None, "dedicated owner-issued collection is required"
    run = tmp_path / "project" / "runs" / "r1"
    with _produced_report(run, model_results=_collection_model(), diagnostic_summary={},
                          complete_collection=True) as case:
        collection = case.projection
        if attack in {"copy", "pickle"}:
            with pytest.raises(TypeError):
                (copy.copy if attack == "copy" else pickle.dumps)(collection)
            return
        if attack == "forged":
            fake = object.__new__(cls)
            for key in ("wire", "_owner", "_collection", "_issuer_token"):
                object.__setattr__(fake, key, getattr(collection, key))
            with pytest.raises(TypeError):
                case.owner.issue_report_document_writer(case.view, fake, canonical_run_root=run)
            return
        if attack == "view":
            view = json.loads(json.dumps(case.writer.view_model))
            view["regression_table"]["rows"][0]["label"] = "unbound replacement"
            object.__setattr__(case.writer, "view_model", view)
        elif attack == "source":
            (run / "model_results" / "ols_1.json").write_bytes(b"tampered")
        else:
            wire = json.loads(json.dumps(collection.wire))
            if attack == "wire": wire["facts"][0]["value"] += 7
            elif attack == "directory": wire["sectionDirectory"] = []
            else: wire["readCompleteness"] = "PARTIAL"
            object.__setattr__(collection, "wire", wire)
        with pytest.raises((TypeError, ValueError)):
            render_html_report(case.writer, run)
        assert not (run / "reports" / "report.html").exists()


def test_local_collection_rejects_foreign_owner(tmp_path):
    with _produced_report(tmp_path / "project" / "runs" / "a",
                          model_results=_collection_model(), diagnostic_summary={},
                          complete_collection=True) as first:
        with _produced_report(tmp_path / "project" / "runs" / "b",
                              model_results=_collection_model(), diagnostic_summary={},
                              complete_collection=True) as second:
            with pytest.raises((TypeError, ValueError)):
                second.owner.issue_report_evidence_writer(first.projection)
            from workbench.report_export import LocalReportCollectionV1
            with pytest.raises((TypeError, ValueError)):
                LocalReportCollectionV1._from_server_collection(first.projection.wire,
                    owner=second.owner, collection=first.projection._collection)


def test_local_collection_admission_releases_unreferenced_capability():
    import gc
    import weakref
    from workbench.report_export import LocalReportCollectionV1, _REPORT_COLLECTION_ADMISSIONS

    # Registry ownership alone must not retain a potentially large report wire.
    instance = object.__new__(LocalReportCollectionV1)
    reference = weakref.ref(instance)
    before = len(_REPORT_COLLECTION_ADMISSIONS)
    _REPORT_COLLECTION_ADMISSIONS[instance] = (object(), object(), object())
    del instance
    gc.collect()
    assert reference() is None
    assert len(_REPORT_COLLECTION_ADMISSIONS) == before


@pytest.mark.parametrize("term", ["x", "<script>untrusted</script>"])
def test_actual_producer_owner_preserves_report_presentation(tmp_path, term):
    from workbench.reporting import render_html_report
    run_root = tmp_path / "project" / "runs" / "r1"
    packet = {"model_id": "ols_1", "model_type": "ols", "model_label": "Declared OLS",
              "nobs": 20, "coefficients": {term: {"estimate": 1.25, "std_error": 0.2, "p_value": 1e-20}}}
    summary = {"diagnostics": {"blockers": [], "warnings": [{"code": "W", "severity": "WARNING",
                    "message": "Declared warning <unsafe>"}], "cautions": [], "info": []},
               "narrative_contract": {"constraints": {"must_not_interpret_treatment_independently": True},
                                      "required_mentions": [], "forbidden_claims": []}}
    with _produced_report(run_root, model_results={"ols_1": packet}, diagnostic_summary=summary) as case:
        assert case.view["status"] == "available"
        assert case.view["title"] == "Declared OLS"
        assert case.view["warnings"][0]["text"] == "Declared warning <unsafe>"
        assert "treatment-like variable" in case.view["causal_caution"]
        path = render_html_report(case.writer, run_root)
        html = path.read_text()
        assert "Declared OLS" in html and "Declared warning &lt;unsafe&gt;" in html
        assert 'id="coefficient-interpretation"' in html
        assert "coefficient = 1.25; p-value = 1e-20" in html
        assert "<script>untrusted</script>" not in html
        if term.startswith("<"):
            assert "&lt;script&gt;untrusted&lt;/script&gt;" in html
        from workbench.exports import export_pdf, export_xlsx
        from openpyxl import load_workbook
        pdf_path = export_pdf(case.writer, run_root)
        assert b"Declared OLS" in pdf_path.read_bytes()
        table_writer = case.owner.issue_report_evidence_writer(case.projection, sections=("regression_table",))
        workbook_path = export_xlsx(table_writer, run_root)
        workbook = load_workbook(workbook_path, data_only=False)
        try:
            rows = list(workbook.active.iter_rows(values_only=True))
            model_column = rows[0].index("models")
            cell = json.loads(rows[1][model_column])["ols_1"]
            assert cell["estimate"] == 1.25 and cell["p_value"] == 1e-20
        finally:
            workbook.close()


def test_actual_producer_survey_text_preserves_counts_without_inventing_missing_design(tmp_path):
    run_root = tmp_path / "project" / "runs" / "r1"
    packet = {"model_id": "ols_1", "model_type": "ols", "nobs": 20,
              "coefficients": {"x": {"estimate": 1.25, "std_error": 0.2}},
              "survey_design": {"variance_method": "jackknife", "replicates_attempted": 12,
                                "replicates_succeeded": 10, "replicates_failed": 2,
                                "replicate_failure_reasons": {"singular_fit": 2}}}
    with _produced_report(run_root, model_results={"ols_1": packet}, diagnostic_summary={}) as case:
        text = "\n".join(case.view["presentation_facts"])
        assert "jackknife" in text
        assert "replicates attempted=12" in text and "succeeded=10" in text and "failed=2" in text
        assert "singular_fit=2" in text
        assert "0 strata" not in text and "confidence intervals use this" not in text
        from workbench.reporting import render_html_report
        html = render_html_report(case.writer, run_root).read_text()
        assert "replicates attempted=12" in html and "singular_fit=2" in html


def _owner_issued_document_writer(
    tmp_path: Path,
    run_root: Path,
):
    with _server_owned_evidence(
        tmp_path / f"owner-store-{uuid4().hex}",
        run_id=run_root.name,
        model_run_id=run_root.name,
        canonical_project_root=run_root.parent.parent,
        keep_open=True,
    ) as (record, _binding, owner):
        wire = _owner_wire(owner, _report_projection_wire())
        wire["envelopeDigest"] = record.envelope_digest
        projection = owner.issue_report_projection(wire)
        bound_report = build_report_view_model(
            {}, run_root, evidence_projection=projection, owner=owner,
            artifact_admission=owner._canonical_report_artifact(),
        )
        return owner.issue_report_document_writer(
            bound_report,
            projection,
            canonical_run_root=run_root,
        )


def _owner_issued_evidence_writer(
    tmp_path: Path,
    sections: tuple[str, ...],
    table_rows: dict[str, list[dict[str, object]]] | None = None,
):
    with _server_owned_evidence(tmp_path / "owner-store", keep_open=True) as (record, _binding, owner):
        wire = _owner_wire(owner, _report_projection_wire())
        wire["envelopeDigest"] = record.envelope_digest
        projection = owner.issue_report_projection(wire)
        return owner.issue_report_evidence_writer(
            projection,
            sections=sections,
            table_rows=table_rows,
        )


class _ObjectWithLoad:
    def __init__(self, value: object) -> None:
        self.value = value

    def load(self, _commit_id: str) -> object:
        return self.value


def test_legacy_facade_rejects_object_with_load_instead_of_minting_trust() -> None:
    record = _persisted_record()

    with pytest.raises(EvidenceProjectionError, match="EVIDENCE_PERSISTENCE_CAPABILITY_REQUIRED"):
        PersistedEvidenceResultsFacadeV1(
            _ObjectWithLoad(record.binding),
            _ObjectWithLoad(record),
        )


def test_copied_registry_receipt_is_typed_rejected_before_report_projection(tmp_path: Path) -> None:
    with _server_owned_evidence(tmp_path) as (record, binding, capability):
        with pytest.raises(TypeError, match="host-bound capabilities"):
            copy.copy(capability._receipt)  # type: ignore[attr-defined]

        indexed_artifact, payload, content_hash = _canonical_indexed_artifact(record, binding)
        view = build_indexed_report_evidence(
            indexed_artifact,
            payload,
            artifact_content_sha256=content_hash,
            authoritative_persistence=capability,
        )

    assert view["status"] == "available"


def test_report_accepts_owner_issued_degraded_typed_projection(tmp_path: Path) -> None:
    envelope = json.loads(_FIXTURE.read_text(encoding="utf-8"))
    envelope["completeness_status"] = "degraded"
    envelope["degradation_reasons"] = [
        {
            "code": "CATEGORICAL_AUTO_DUMMY_CODED",
            "message": "Categorical predictors were dummy coded automatically.",
            "source_refs": ["run:run-1"],
        }
    ]

    with _server_owned_evidence(
        tmp_path / "owner-store",
        envelope_payload=envelope,
    ) as (record, _binding, owner):
        wire = dict(owner._canonical_report_wire())  # type: ignore[attr-defined]
        assert wire["executionStatus"] == "completed"
        assert wire["availability"] == "available"
        assert wire["completenessStatus"] == "degraded"
        assert wire["facts"]
        wire["envelopeDigest"] = record.envelope_digest

        projection = owner.issue_report_projection(wire)

    assert projection.wire["status"] == "degraded"
    assert projection.wire["completenessStatus"] == "degraded"
    assert projection.wire["facts"]


def test_degraded_available_report_projection_mints_a_server_receipt() -> None:
    view = {
        "status": "degraded",
        "projectionDigest": "a" * 64,
        "provenance": {"scopeRef": "scope:project:1"},
        "actorRef": "actor:1",
        "projectRef": "project:1",
        "sessionRef": "session:1",
        "scopeRef": "scope:project:1",
        "runId": "run-1",
        "modelRunId": "run-1",
        "artifactRef": "artifact:1",
        "commitId": "commit:1",
        "confirmationRef": "confirmation:1",
    }

    receipt = runs_routes._report_evidence_server_receipt(view)

    assert receipt is not None
    assert receipt["kind"] == "evidence.report.server-receipt.v1"
    assert receipt["projectionDigest"] == view["projectionDigest"]


def test_indexed_post_estimation_without_committed_binding_is_unavailable_and_not_scalarized() -> None:
    view = build_indexed_report_evidence(
        _indexed_artifact(),
        {
            "source": {"run_id": "run-1", "model_run_id": "run-1"},
            "result": {
                "stationary_point": 211.59338521178753,
                "nested": {"secret_scalar": 99},
            },
        },
    )

    assert view["status"] == "unavailable"
    assert view["reasonCode"] == "EVIDENCE_COMMIT_BINDING_REQUIRED"
    assert view["facts"] == []
    assert "211.59338521178753" not in json.dumps(view, sort_keys=True)
    assert "secret_scalar" not in json.dumps(view, sort_keys=True)


def test_tampered_committed_envelope_is_blocked_without_facts(tmp_path: Path) -> None:
    with _server_owned_evidence(tmp_path) as (record, binding, capability):
        tampered_envelope = record.envelope.to_dict()
        tampered_envelope["numeric_facts"][0]["value"] = 999.0  # type: ignore[index]
        indexed_artifact, payload, content_hash = _canonical_indexed_artifact(record, binding)
        indexed_artifact["evidence_envelope"] = tampered_envelope

        view = build_indexed_report_evidence(
            indexed_artifact,
            payload,
            artifact_content_sha256=content_hash,
            authoritative_persistence=capability,
        )

    assert view["status"] in {"blocked", "unavailable"}
    assert view["reasonCode"] == "PROJECTION_DIGEST_MISMATCH"
    assert view["facts"] == []


def test_unbound_projection_is_unavailable_without_server_identity_and_admission() -> None:
    record = _persisted_record()

    with pytest.raises(EvidenceProjectionError, match="EVIDENCE_SERVER_SOURCE_REQUIRED"):
        default_projection_registry().project(
            record,
            consumer_id="report",
            destination="local",
        )


def test_report_owner_reloads_live_packet_after_cached_field_tamper(tmp_path: Path) -> None:
    with _server_owned_evidence(tmp_path / "owner-store") as (record, _binding, owner):
        object.__setattr__(owner, "_canonical_record", _persisted_record())
        object.__setattr__(owner, "_canonical_artifact", object())

        wire = _owner_wire(owner, _report_projection_wire())
        wire["envelopeDigest"] = record.envelope_digest
        projection = owner.issue_report_projection(wire)

    assert projection.wire["status"] == "available"
    assert projection.wire["facts"]
    assert projection.wire["artifactRef"] == record.artifact_id


def test_report_request_identity_cache_tamper_is_not_a_fact_source(tmp_path: Path) -> None:
    with _server_owned_evidence(tmp_path / "owner-store") as (_record, _binding, owner):
        tampered = ReportEvidenceRequestIdentityV1(
            actor_ref=owner.actor_ref,
            project_ref=owner.project_ref,
            session_ref=owner.session_ref,
            scope_ref=owner.scope_ref,
            run_id="run-forged",
            model_run_id=owner.model_run_id,
            confirmation_ref=owner.confirmation_ref,
        )
        object.__setattr__(owner, "_report_request_identity", tampered)

        identity = owner._require_report_request_identity(None)  # type: ignore[attr-defined]

    assert identity.run_id == "run-1"
    assert identity != tampered


def test_remediation_five_caller_mapping_cannot_build_or_render_legacy_report_before_io(
    tmp_path: Path,
    monkeypatch,
) -> None:
    forged_projection = _report_projection_wire()
    forged_view = build_report_view_model(
        {},
        tmp_path,
        evidence_projection=forged_projection,
    )
    assert forged_view["status"] in {"blocked", "unavailable"}
    assert forged_view["facts"] == []

    run_root = tmp_path / "run-1"
    run_root.mkdir()
    write_bytes = __import__("unittest.mock", fromlist=["Mock"]).Mock()
    monkeypatch.setattr(reporting, "_write_report_output", write_bytes)

    with pytest.raises(TypeError, match="owner|capability|projection|Report"):
        render_legacy_html_report(forged_projection, run_root)

    write_bytes.assert_not_called()


def test_remediation_five_object_new_html_writer_cannot_reach_io(
    tmp_path: Path,
    monkeypatch,
) -> None:
    run_root = tmp_path / "run-1"
    run_root.mkdir()
    forged_writer = object.__new__(ReportDocumentWriterV1)
    object.__setattr__(forged_writer, "view_model", {})
    object.__setattr__(forged_writer, "projection", object.__new__(EvidenceProjectionV1))
    object.__setattr__(forged_writer, "canonical_run_root", run_root)
    write_bytes = __import__("unittest.mock", fromlist=["Mock"]).Mock()
    monkeypatch.setattr(reporting, "_write_report_output", write_bytes)

    with pytest.raises(TypeError, match="owner|capability|projection|Report"):
        reporting.render_html_report(forged_writer, run_root)

    write_bytes.assert_not_called()


def test_remediation_five_server_admitted_owner_can_build_and_render_legacy_report(
    tmp_path: Path,
) -> None:
    project_root = tmp_path / "project"
    run_root = project_root / "runs" / "run-1"
    run_root.mkdir(parents=True)
    from workbench.artifacts import initialize_artifact_index

    initialize_artifact_index(run_root)
    with _server_owned_evidence(
        tmp_path / "owner-store",
        run_id="run-1",
        model_run_id="run-1",
        canonical_project_root=project_root,
        with_project_file_store=True,
    ) as (record, _binding, owner):
        wire = _owner_wire(owner, _report_projection_wire())
        wire["envelopeDigest"] = record.envelope_digest
        projection = owner.issue_report_projection(wire)
        view = build_report_view_model(
            {},
            run_root,
            evidence_projection=projection,
            owner=owner,
            artifact_admission=owner._canonical_report_artifact(),  # type: ignore[attr-defined]
        )
        assert view["status"] == "available"
        assert view["facts"]

        output = render_legacy_html_report(
            view,
            run_root,
            owner=owner,
            projection=projection,
        )

    assert output == run_root / "reports" / "report.html"
    assert output.is_file()


def test_remediation_four_caller_created_request_identity_cannot_admit_report_view(
    tmp_path: Path,
) -> None:
    with _server_owned_evidence(tmp_path / "owner-store") as (record, _binding, owner):
        projection = default_projection_registry().project(
            owner._canonical_record,  # type: ignore[attr-defined]
            consumer_id="report",
            destination="local",
        )
        caller_identity = ReportEvidenceRequestIdentityV1(
            actor_ref=owner.actor_ref,
            project_ref=owner.project_ref,
            session_ref=owner.session_ref,
            scope_ref=owner.scope_ref,
            run_id=owner.run_id,
            model_run_id=owner.model_run_id,
            confirmation_ref=owner.confirmation_ref,
        )

        caller_view = build_persisted_evidence_report_view(
            projection,
            owner=owner,
            request_identity=caller_identity,
            artifact_admission=owner._canonical_report_artifact(),  # type: ignore[attr-defined]
        )
        valid_view = build_persisted_evidence_report_view(
            projection,
            owner=owner,
            artifact_admission=owner._canonical_report_artifact(),  # type: ignore[attr-defined]
        )

    assert caller_view["status"] in {"blocked", "unavailable"}
    assert caller_view["facts"] == []
    assert caller_view["reasonCode"] == "EVIDENCE_OWNER_REQUEST_REQUIRED"
    assert valid_view["status"] == "available"
    assert valid_view["facts"]
    assert record.artifact_id == valid_view["artifactRef"]


def test_remediation_four_object_new_projection_is_typed_unavailable_before_report_facts(
    tmp_path: Path,
) -> None:
    from workbench.evidence.consumer_adapters import EvidenceProjectionV1

    with _server_owned_evidence(tmp_path / "owner-store") as (_record, _binding, owner):
        forged_projection = object.__new__(EvidenceProjectionV1)
        view = build_persisted_evidence_report_view(
            forged_projection,
            owner=owner,
            artifact_admission=owner._canonical_report_artifact(),  # type: ignore[attr-defined]
        )

    assert view["status"] in {"blocked", "unavailable"}
    assert view["facts"] == []
    assert view["reasonCode"] == "EVIDENCE_PROJECTION_INVALID"


def test_report_view_serializes_the_declared_renderer_identity(tmp_path: Path) -> None:
    with _server_owned_evidence(tmp_path / "renderer-identity") as (_record, _binding, owner):
        projection = default_projection_registry().project(
            owner._canonical_record,  # type: ignore[attr-defined]
            consumer_id="report",
            destination="local",
        )
        view = build_persisted_evidence_report_view(
            projection,
            owner=owner,
            request_identity=owner._report_request_identity,  # type: ignore[attr-defined]
            artifact_admission=owner._canonical_report_artifact(),  # type: ignore[attr-defined]
        )

    assert view["rendererKey"] == "workbench.consumer_projection.report"


def test_report_view_preserves_numeric_fact_validation_provenance(tmp_path: Path) -> None:
    with _server_owned_evidence(tmp_path / "fact-validation") as (_record, _binding, owner):
        projection = default_projection_registry().project(
            owner._canonical_record,  # type: ignore[attr-defined]
            consumer_id="report",
            destination="local",
        )
        view = build_persisted_evidence_report_view(
            projection,
            owner=owner,
            request_identity=owner._report_request_identity,  # type: ignore[attr-defined]
            artifact_admission=owner._canonical_report_artifact(),  # type: ignore[attr-defined]
        )

    first_fact = view["facts"][0]
    assert first_fact["validationProvenance"]["level"] == "same_library_regression"
    assert first_fact["validationProvenance"]["oracleChecks"][0]["status"] == "pass"


@pytest.mark.parametrize(
    ("outer_field", "nested_field", "outer_value"),
    (
        ("evidence_commit_binding", "evidence_commit_binding", None),
        ("evidence_envelope", "evidence_envelope", {}),
        ("evidence_envelope_ref", "envelope_ref", ""),
    ),
)
def test_remediation_four_present_invalid_outer_wire_never_falls_back_to_nested(
    tmp_path: Path,
    outer_field: str,
    nested_field: str,
    outer_value: object,
) -> None:
    with _server_owned_evidence(tmp_path / outer_field) as (record, binding, owner):
        indexed_artifact, payload, content_hash = _canonical_indexed_artifact(record, binding)
        indexed_artifact["evidence"] = {
            nested_field: (
                binding.to_dict()
                if nested_field == "evidence_commit_binding"
                else record.envelope.to_dict()
                if nested_field == "evidence_envelope"
                else f"evidence-envelope:{record.envelope_digest}"
            )
        }
        indexed_artifact[outer_field] = outer_value

        view = build_indexed_report_evidence(
            indexed_artifact,
            payload,
            artifact_content_sha256=content_hash,
            authoritative_persistence=owner,
        )

    assert view["status"] in {"blocked", "unavailable"}
    assert view["facts"] == []
    assert view["reasonCode"] == "EVIDENCE_WIRE_FIELD_INVALID"


def test_remediation_three_report_provenance_is_complete_and_canonical(tmp_path: Path) -> None:
    with _server_owned_evidence(tmp_path / "owner-store") as (record, binding, owner):
        indexed_artifact, payload, content_hash = _canonical_indexed_artifact(record, binding)
        view = build_indexed_report_evidence(
            indexed_artifact,
            payload,
            artifact_content_sha256=content_hash,
            authoritative_persistence=owner,
        )

    assert view["status"] == "available"
    assert set(
        (
            "artifactRef",
            "artifactType",
            "scopeRef",
            "contentHash",
            "sourceRef",
            "validatorRef",
            "executionRef",
            "contractDigest",
            "commitId",
            "envelopeRef",
            "envelopeDigest",
            "projectionRef",
            "resultIdentity",
            "lineageRefs",
        )
    ) <= set(view["provenance"])
    assert view["provenance"]["contentHash"] == record.content_hash
    assert view["provenance"]["sourceRef"] == record.source_ref
    assert view["provenance"]["executionRef"] == record.execution_ref
    assert view["provenance"]["contractDigest"] == record.contract_digest


@pytest.mark.parametrize(
    ("field", "reason_code"),
    (
        ("source_ref", "EVIDENCE_SOURCE_REF_REQUIRED"),
        ("validator_ref", "EVIDENCE_VALIDATOR_REF_REQUIRED"),
        ("execution_ref", "EVIDENCE_EXECUTION_REF_REQUIRED"),
        ("contract_digest", "EVIDENCE_CONTRACT_DIGEST_REQUIRED"),
    ),
)
def test_remediation_three_outer_index_missing_provenance_is_unavailable(
    tmp_path: Path,
    field: str,
    reason_code: str,
) -> None:
    with _server_owned_evidence(tmp_path / field) as (record, binding, owner):
        indexed_artifact, payload, content_hash = _canonical_indexed_artifact(record, binding)
        indexed_artifact.pop(field)
        view = build_indexed_report_evidence(
            indexed_artifact,
            payload,
            artifact_content_sha256=content_hash,
            authoritative_persistence=owner,
        )

    assert view["status"] in {"blocked", "unavailable"}
    assert view["reasonCode"] == reason_code
    assert view["facts"] == []


@pytest.mark.parametrize(
    ("field", "value", "reason_code"),
    (
        ("source_ref", "source:forged", "EVIDENCE_SOURCE_REF_MISMATCH"),
        ("execution_ref", "execution:foreign", "EVIDENCE_EXECUTION_REF_MISMATCH"),
        ("contract_digest", "f" * 64, "EVIDENCE_CONTRACT_DIGEST_MISMATCH"),
    ),
)
def test_remediation_three_outer_index_tampered_provenance_is_unavailable(
    tmp_path: Path,
    field: str,
    value: object,
    reason_code: str,
) -> None:
    with _server_owned_evidence(tmp_path / field) as (record, binding, owner):
        indexed_artifact, payload, content_hash = _canonical_indexed_artifact(record, binding)
        indexed_artifact[field] = value
        view = build_indexed_report_evidence(
            indexed_artifact,
            payload,
            artifact_content_sha256=content_hash,
            authoritative_persistence=owner,
        )

    assert view["status"] in {"blocked", "unavailable"}
    assert view["reasonCode"] == reason_code
    assert view["facts"] == []


def test_indexed_canonical_commit_is_the_only_path_that_admits_report_facts(tmp_path: Path) -> None:
    with _server_owned_evidence(tmp_path) as (record, binding, capability):
        indexed_artifact, payload, content_hash = _canonical_indexed_artifact(record, binding)
        view = build_indexed_report_evidence(
            indexed_artifact,
            payload,
            artifact_content_sha256=content_hash,
            authoritative_persistence=capability,
        )

        assert view["status"] == "available"
        assert view["envelopeDigest"] == record.envelope_digest
        assert [fact["value"] for fact in view["facts"]] == [2.0, 0.2]


def test_foreign_artifact_cannot_reuse_a_valid_committed_pair(tmp_path: Path) -> None:
    with _server_owned_evidence(tmp_path) as (record, binding, capability):
        indexed_artifact, payload, content_hash = _canonical_indexed_artifact(record, binding)
        indexed_artifact["artifact_id"] = "foreign-artifact"

        view = build_indexed_report_evidence(
            indexed_artifact,
            payload,
            artifact_content_sha256=content_hash,
            authoritative_persistence=capability,
        )

        assert view["status"] == "blocked"
        assert view["reasonCode"] == "EVIDENCE_ARTIFACT_ID_MISMATCH"
        assert view["facts"] == []


def test_registry_rejects_a_self_consistent_foreign_copy_of_a_committed_pair(tmp_path: Path) -> None:
    with _server_owned_evidence(tmp_path) as (record, binding, capability):
        indexed_artifact, payload, content_hash = _canonical_indexed_artifact(record, binding)
        foreign_payload = dict(payload)
        foreign_raw = json.dumps(
            foreign_payload,
            separators=(",", ":"),
            sort_keys=True,
        ).encode("utf-8")
        foreign = dict(indexed_artifact)
        foreign.update(
            {
                "artifact_id": "artifact-foreign",
                "evidence_artifact_ref": "artifact-foreign",
                "sha256": hashlib.sha256(foreign_raw).hexdigest(),
            }
        )

        view = build_indexed_report_evidence(
            foreign,
            foreign_payload,
            artifact_content_sha256=hashlib.sha256(foreign_raw).hexdigest(),
            authoritative_persistence=capability,
        )

    assert view["status"] in {"unavailable", "blocked"}
    assert view["facts"] == []
    assert view["reasonCode"] == "EVIDENCE_ADMISSION_NOT_FOUND"


def test_foreign_run_cannot_reuse_a_valid_committed_pair(tmp_path: Path) -> None:
    with _server_owned_evidence(tmp_path) as (record, binding, capability):
        indexed_artifact, payload, content_hash = _canonical_indexed_artifact(record, binding)
        indexed_artifact["run_id"] = "foreign-run"
        indexed_artifact["model_run_id"] = "foreign-run"
        indexed_artifact["evidence_run_id"] = "foreign-run"
        indexed_artifact["evidence_model_run_id"] = "foreign-run"

        view = build_indexed_report_evidence(
            indexed_artifact,
            payload,
            artifact_content_sha256=content_hash,
            authoritative_persistence=capability,
        )

        assert view["status"] == "blocked"
        assert view["reasonCode"] == "EVIDENCE_RUN_ID_MISMATCH"
        assert view["facts"] == []


def test_foreign_project_and_scope_with_the_same_run_are_blocked(tmp_path: Path) -> None:
    with _server_owned_evidence(tmp_path) as (record, binding, capability):
        indexed_artifact, payload, content_hash = _canonical_indexed_artifact(record, binding)
        indexed_artifact.update(
            {
                "project_ref": "project:foreign",
                "session_ref": "session:foreign",
                "scope_ref": "scope:foreign",
            }
        )

        view = build_indexed_report_evidence(
            indexed_artifact,
            payload,
            artifact_content_sha256=content_hash,
            authoritative_persistence=capability,
        )

    assert view["status"] in {"blocked", "unavailable"}
    assert view["reasonCode"] == "EVIDENCE_OWNER_CONTEXT_MISMATCH"
    assert view["facts"] == []


def test_foreign_content_hash_and_admission_ref_cannot_reuse_a_pair(tmp_path: Path) -> None:
    with _server_owned_evidence(tmp_path) as (record, binding, capability):
        indexed_artifact, payload, content_hash = _canonical_indexed_artifact(record, binding)
        indexed_artifact["content_hash"] = "f" * 64

        hash_view = build_indexed_report_evidence(
            indexed_artifact,
            payload,
            artifact_content_sha256=content_hash,
            authoritative_persistence=capability,
        )

        assert hash_view["status"] == "blocked"
        assert hash_view["reasonCode"] == "EVIDENCE_ARTIFACT_CONTENT_HASH_MISMATCH"
        assert hash_view["facts"] == []

        indexed_artifact, payload, content_hash = _canonical_indexed_artifact(record, binding)
        indexed_artifact["sha256"] = "f" * 64
        file_hash_view = build_indexed_report_evidence(
            indexed_artifact,
            payload,
            artifact_content_sha256=content_hash,
            authoritative_persistence=capability,
        )

        assert file_hash_view["status"] == "blocked"
        assert file_hash_view["reasonCode"] == "EVIDENCE_ARTIFACT_FILE_HASH_MISMATCH"
        assert file_hash_view["facts"] == []

        indexed_artifact, payload, content_hash = _canonical_indexed_artifact(record, binding)
        indexed_artifact["evidence_artifact_admission_ref"] = "foreign-admission"
        admission_view = build_indexed_report_evidence(
            indexed_artifact,
            payload,
            artifact_content_sha256=content_hash,
            authoritative_persistence=capability,
        )

        assert admission_view["status"] == "blocked"
        assert admission_view["reasonCode"] == "EVIDENCE_ARTIFACT_ADMISSION_MISMATCH"
        assert admission_view["facts"] == []


def test_multi_artifact_bundle_fails_closed_without_a_single_canonical_identity(tmp_path: Path) -> None:
    with _server_owned_evidence(tmp_path / "owner-store") as (_record, _binding, owner):
        projection = owner.issue_report_projection(owner._canonical_report_wire())
        first = build_persisted_evidence_report_view(
            projection,
            owner=owner,
            artifact_admission=owner._canonical_report_artifact(),  # type: ignore[attr-defined]
        )
    second = dict(first)
    second["projectionId"] = "report:evidence-commit-2"
    second["artifactRef"] = "artifact-2"

    bundle = build_report_evidence_bundle([first, second])

    assert bundle is not None
    assert bundle["status"] == "blocked"
    assert bundle["reasonCode"] == "REPORT_EVIDENCE_BUNDLE_IDENTITY_UNBOUND"
    assert bundle["commitId"] == "UNAVAILABLE"
    assert bundle["envelopeDigest"] == "UNAVAILABLE"
    assert bundle["facts"] == []


def test_runs_route_injects_the_server_owned_projection_capability(tmp_path: Path, monkeypatch) -> None:
    run_root = tmp_path / "run-1"
    artifact_path = run_root / "artifacts" / "post_estimation" / "result.json"
    artifact_path.parent.mkdir(parents=True)

    owner_root = tmp_path / "owner-store"
    owner_root.mkdir()
    with _server_owned_evidence(owner_root) as (record, binding, capability):
        indexed_artifact, payload, _content_hash = _canonical_indexed_artifact(record, binding)
        indexed_artifact["path"] = "artifacts/post_estimation/result.json"
        indexed_artifact.pop("evidence_commit_binding")
        indexed_artifact.pop("evidence_envelope")
        artifact_path.write_bytes(
            json.dumps(payload, separators=(",", ":"), sort_keys=True).encode("utf-8")
        )
        (run_root / "artifacts_index.json").write_text(
            json.dumps({"schema_version": 1, "artifacts": [indexed_artifact]}),
            encoding="utf-8",
        )
        monkeypatch.setattr(
            runs_routes,
            "current_report_evidence_persistence",
            lambda **kwargs: capability if kwargs.get("run_id") == "run-1" else None,
        )

        view = runs_routes._report_evidence_for_run(run_root, [])

    assert view["status"] == "available"
    assert view["commitId"] == binding.commit_id
    assert view["envelopeDigest"] == record.envelope_digest
    assert [fact["value"] for fact in view["facts"]] == [2.0, 0.2]


def test_runs_route_rejects_a_test_injected_fake_capability(tmp_path: Path, monkeypatch) -> None:
    run_root = tmp_path / "run-1"
    artifact_path = run_root / "artifacts" / "post_estimation" / "result.json"
    artifact_path.parent.mkdir(parents=True)

    owner_root = tmp_path / "owner-store"
    owner_root.mkdir()
    with _server_owned_evidence(owner_root) as (record, binding, _capability):
        indexed_artifact, payload, _content_hash = _canonical_indexed_artifact(record, binding)
        indexed_artifact["path"] = "artifacts/post_estimation/result.json"
        indexed_artifact.pop("evidence_commit_binding")
        indexed_artifact.pop("evidence_envelope")
        artifact_path.write_bytes(
            json.dumps(payload, separators=(",", ":"), sort_keys=True).encode("utf-8")
        )
        (run_root / "artifacts_index.json").write_text(
            json.dumps({"schema_version": 1, "artifacts": [indexed_artifact]}),
            encoding="utf-8",
        )
        monkeypatch.setattr(
            runs_routes,
            "current_report_evidence_persistence",
            lambda **_kwargs: object(),
        )

        view = runs_routes._report_evidence_for_run(run_root, [])

    assert view["status"] in {"blocked", "unavailable"}
    assert view["reasonCode"] == "EVIDENCE_PERSISTENCE_CAPABILITY_REQUIRED"
    assert view["facts"] == []


def test_http_run_detail_rejects_a_rehashed_index_for_tampered_payload(
    tmp_path: Path,
    monkeypatch,
) -> None:
    project_root = tmp_path / "project"
    run_root = project_root / "runs" / "run-1"
    artifact_path = run_root / "artifacts" / "post_estimation" / "result.json"
    artifact_path.parent.mkdir(parents=True)
    (run_root / "run_manifest.json").write_text(
        json.dumps({"run_id": "run-1", "status": "completed", "mode": "auto", "lineage": []}),
        encoding="utf-8",
    )

    owner_root = tmp_path / "owner-store"
    owner_root.mkdir()
    with _server_owned_evidence(owner_root, canonical_project_root=project_root) as (record, binding, capability):
        indexed_artifact, _payload, _content_hash = _canonical_indexed_artifact(record, binding)
        tampered_payload = {
            "source": {"run_id": "run-1", "model_run_id": "run-1"},
            "result": {"committed": False, "tampered_scalar": 999.0},
        }
        tampered_raw = json.dumps(
            tampered_payload,
            separators=(",", ":"),
            sort_keys=True,
        ).encode("utf-8")
        indexed_artifact["path"] = "artifacts/post_estimation/result.json"
        indexed_artifact["sha256"] = hashlib.sha256(tampered_raw).hexdigest()
        indexed_artifact.pop("evidence_commit_binding")
        indexed_artifact.pop("evidence_envelope")
        artifact_path.write_bytes(tampered_raw)
        (run_root / "artifacts_index.json").write_text(
            json.dumps({"schema_version": 1, "artifacts": [indexed_artifact]}),
            encoding="utf-8",
        )
        monkeypatch.setattr(
            runs_routes,
            "current_report_evidence_persistence",
            lambda **kwargs: capability if kwargs.get("run_id") == "run-1" else None,
        )

        response = TestClient(app).get(
            "/runs/run-1",
            params={"project_root": str(project_root)},
        )

    assert response.status_code == 200, response.text
    report_evidence = response.json()["report_evidence"]
    assert report_evidence["status"] in {"blocked", "unavailable"}
    assert report_evidence["reasonCode"] == "EVIDENCE_ARTIFACT_FILE_HASH_MISMATCH"
    assert report_evidence["facts"] == []
    assert "999.0" not in json.dumps(report_evidence, sort_keys=True)


def test_http_run_detail_rejects_foreign_actor_and_copied_confirmation(
    tmp_path: Path,
    monkeypatch,
) -> None:
    project_root = tmp_path / "project"
    run_root = project_root / "runs" / "run-1"
    artifact_path = run_root / "artifacts" / "post_estimation" / "result.json"
    artifact_path.parent.mkdir(parents=True)
    (run_root / "run_manifest.json").write_text(
        json.dumps({"run_id": "run-1", "status": "completed", "mode": "auto", "lineage": []}),
        encoding="utf-8",
    )
    owner_root = tmp_path / "owner-store"
    owner_root.mkdir()
    with _server_owned_evidence(owner_root, canonical_project_root=project_root) as (record, binding, capability):
        index, payload, _content_hash = _canonical_indexed_artifact(record, binding)
        index["path"] = "artifacts/post_estimation/result.json"
        index["actor_ref"] = "actor:foreign"
        index["confirmation_ref"] = "confirmation:copied"
        index.pop("evidence_commit_binding")
        index.pop("evidence_envelope")
        artifact_path.write_bytes(json.dumps(payload, separators=(",", ":"), sort_keys=True).encode())
        (run_root / "artifacts_index.json").write_text(
            json.dumps({"schema_version": 1, "artifacts": [index]}),
            encoding="utf-8",
        )
        monkeypatch.setattr(
            runs_routes,
            "current_report_evidence_persistence",
            lambda **_kwargs: capability,
        )
        response = TestClient(app).get(
            "/runs/run-1",
            params={"project_root": str(project_root)},
        )

    report_evidence = response.json()["report_evidence"]
    assert response.status_code == 200
    assert report_evidence["status"] == "blocked"
    assert report_evidence["reasonCode"] == "EVIDENCE_OWNER_CONTEXT_MISMATCH"
    assert report_evidence["facts"] == []


def test_run_detail_exposes_typed_unavailable_projection_for_indexed_result(tmp_path, monkeypatch) -> None:
    project_root = tmp_path / "project"
    run_root = project_root / "runs" / "run-1"
    artifact_path = run_root / "artifacts" / "post_estimation.json"
    artifact_path.parent.mkdir(parents=True)
    payload = {
        "source": {"run_id": "run-1", "model_run_id": "run-1"},
        "result": {"stationary_point": 211.5, "nested": {"secret_scalar": 99}},
    }
    raw_payload = json.dumps(payload, separators=(",", ":")).encode("utf-8")
    artifact_path.write_bytes(raw_payload)
    record = _indexed_artifact(
        path="artifacts/post_estimation.json",
        sha256=hashlib.sha256(raw_payload).hexdigest(),
    )
    run_root.mkdir(exist_ok=True)
    (run_root / "run_manifest.json").write_text(
        json.dumps({"run_id": "run-1", "status": "completed", "mode": "auto", "lineage": []}),
        encoding="utf-8",
    )
    (run_root / "artifacts_index.json").write_text(
        json.dumps({"schema_version": 1, "artifacts": [record]}),
        encoding="utf-8",
    )
    response = TestClient(app).get("/runs/run-1", params={"project_root": str(project_root)})

    assert response.status_code == 200, response.text
    report_evidence = response.json()["report_evidence"]
    assert report_evidence["status"] == "unavailable"
    assert report_evidence["reasonCode"] == "EVIDENCE_PERSISTENCE_CAPABILITY_REQUIRED"
    assert report_evidence["facts"] == []
    assert "211.5" not in json.dumps(report_evidence, sort_keys=True)
    assert "secret_scalar" not in json.dumps(report_evidence, sort_keys=True)


def test_default_app_lifespan_uses_the_server_owned_report_capability(tmp_path: Path) -> None:
    project_root = tmp_path / "project"
    run_root = project_root / "runs" / "run-1"
    artifact_path = run_root / "artifacts" / "post_estimation" / "result.json"
    artifact_path.parent.mkdir(parents=True)
    (run_root / "run_manifest.json").write_text(
        json.dumps({"run_id": "run-1", "status": "completed", "mode": "auto", "lineage": []}),
        encoding="utf-8",
    )
    owner_root = tmp_path / "owner-store"
    owner_root.mkdir()
    with _server_owned_evidence(owner_root, canonical_project_root=project_root) as (record, binding, capability):
        index, payload, _content_hash = _canonical_indexed_artifact(record, binding)
        index["path"] = "artifacts/post_estimation/result.json"
        index.pop("evidence_commit_binding")
        index.pop("evidence_envelope")
        artifact_path.write_bytes(json.dumps(payload, separators=(",", ":"), sort_keys=True).encode())
        (run_root / "artifacts_index.json").write_text(
            json.dumps({"schema_version": 1, "artifacts": [index]}),
            encoding="utf-8",
        )
        configure_report_evidence_owner(capability)
        try:
            with TestClient(app) as client:
                response = client.get(
                    "/runs/run-1",
                    params={"project_root": str(project_root)},
                )
        finally:
            configure_report_evidence_owner(None)

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["report_evidence"]["status"] == "available"
    assert body["report_evidence_receipt"]["kind"] == "evidence.report.server-receipt.v1"
    assert body["report_evidence_receipt"]["projectionDigest"] == body["report_evidence"]["projectionDigest"]
    assert body["report_evidence_receipt"]["provenance"] == body["report_evidence"]["provenance"]
    app.state.report_evidence_persistence = None


def test_default_app_lifespan_never_leaves_report_owner_as_none() -> None:
    with TestClient(app):
        owner = app.state.report_evidence_persistence

    assert owner is not None
    assert type(owner) in {
        ServerOwnedReportEvidencePersistenceV1,
        UnavailableReportEvidenceCapabilityV1,
    }


def test_round_three_default_lifespan_installs_a_real_c0_owner_without_manual_configuration(
    tmp_path: Path,
) -> None:
    """A host-injected owner crosses real HTTP without the legacy configure helper."""

    project_root = tmp_path / "project"
    run_root = project_root / "runs" / "run-1"
    artifact_path = run_root / "artifacts" / "post_estimation" / "result.json"
    artifact_path.parent.mkdir(parents=True)
    (run_root / "run_manifest.json").write_text(
        json.dumps({"run_id": "run-1", "status": "completed", "mode": "auto", "lineage": []}),
        encoding="utf-8",
    )
    with _server_owned_evidence(
        tmp_path / "owner-store",
        canonical_project_root=project_root,
    ) as (record, binding, capability):
        index, payload, _content_hash = _canonical_indexed_artifact(record, binding)
        index["path"] = "artifacts/post_estimation/result.json"
        index.pop("evidence_commit_binding")
        index.pop("evidence_envelope")
        artifact_path.write_bytes(json.dumps(payload, separators=(",", ":"), sort_keys=True).encode())
        (run_root / "artifacts_index.json").write_text(
            json.dumps({"schema_version": 1, "artifacts": [index]}),
            encoding="utf-8",
        )
        # This is the deployment host injection; no configure helper is used.
        app.state.report_evidence_persistence = capability
        try:
            with TestClient(app) as client:
                assert type(app.state.report_evidence_persistence) is ServerOwnedReportEvidencePersistenceV1
                response = client.get("/runs/run-1", params={"project_root": str(project_root)})
        finally:
            app.state.report_evidence_persistence = None

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["report_evidence"]["status"] == "available"
    assert body["report_evidence_receipt"]["kind"] == "evidence.report.server-receipt.v1"
    assert body["report_evidence_receipt"]["projectionDigest"] == body["report_evidence"]["projectionDigest"]


def test_round_three_public_projection_and_writer_seams_reject_forged_inputs() -> None:
    wire = {
        "schemaVersion": 1,
        "projectionId": "report:commit-1",
        "consumerId": "report",
        "consumerVersion": "1.0",
        "actorRef": "actor:1",
        "projectRef": "project:1",
        "sessionRef": "session:1",
        "scopeRef": "scope:project:1",
        "runId": "run-1",
        "modelRunId": "run-1",
        "artifactRef": "artifact-1",
        "commitId": "commit-1",
        "confirmationRef": "confirmation:run-1",
        "envelopeDigest": "a" * 64,
        "status": "available",
        "validationStatus": "met",
        "facts": [{"factId": "fact-1", "statisticId": "estimate", "value": 1.25}],
        "provenance": {"artifactRef": "artifact-1", "commitId": "commit-1"},
    }

    with pytest.raises(TypeError, match="server-owned|capability|factory"):
        EvidenceProjectionV1(wire)  # type: ignore[arg-type]
    with pytest.raises(TypeError, match="server-owned|capability|factory"):
        EvidenceProjectionV1.from_wire(wire)
    with pytest.raises(TypeError, match="server-owned|capability|factory"):
        ReportEvidenceWriterV1.from_projection(object())  # type: ignore[arg-type]
    with pytest.raises(TypeError, match="server-owned|capability|factory"):
        ReportDocumentWriterV1.from_view_model(
            {"title": "forged"}, object(), canonical_run_root=Path("/tmp/forged")  # type: ignore[arg-type]
        )


def test_round_three_writer_rejects_raw_tables_and_raw_report_payloads() -> None:
    with pytest.raises(TypeError, match="server-owned|capability|factory"):
        ReportEvidenceWriterV1.from_projection(  # type: ignore[arg-type]
            object(),
            sections=("facts",),
            tables={"facts": ({"value": 99},)},
        )
    with pytest.raises(TypeError, match="server-owned|capability|factory"):
        ReportDocumentWriterV1.from_view_model(  # type: ignore[arg-type]
            {"facts": [{"value": 99}]},
            object(),
            canonical_run_root=Path("/tmp/forged"),
        )


def test_round_three_owner_factory_rejects_forged_commit_digest_and_identity(
    tmp_path: Path,
) -> None:
    with _server_owned_evidence(tmp_path / "owner-store") as (record, _binding, owner):
        valid_wire = _owner_wire(owner, _report_projection_wire())
        valid_wire["envelopeDigest"] = record.envelope_digest
        for key, forged in (
            ("commitId", "commit:forged"),
            ("envelopeDigest", "b" * 64),
            ("projectRef", "project:forged"),
        ):
            candidate = dict(valid_wire)
            candidate[key] = forged
            with pytest.raises(EvidenceProjectionError, match="OWNER_CONTEXT_MISMATCH|owner|projection"):
                owner.issue_report_projection(candidate)


def test_round_four_owner_rejects_forged_facts_and_projection_digest(
    tmp_path: Path,
) -> None:
    with _server_owned_evidence(tmp_path / "owner-store") as (record, _binding, owner):
        valid_wire = _owner_wire(owner, _report_projection_wire())
        valid_wire["envelopeDigest"] = record.envelope_digest
        for mutation in (
            {"facts": [{"factId": "fact-forged", "statisticId": "estimate", "value": 999.0}]},
            {"projectionDigest": "f" * 64},
        ):
            candidate = dict(valid_wire)
            candidate.update(mutation)
            with pytest.raises(EvidenceProjectionError, match="PROJECTION|FACT|DIGEST|CANONICAL|canonical|owner"):
                owner.issue_report_projection(candidate)


def test_round_four_projection_wire_is_deeply_immutable(tmp_path: Path) -> None:
    projection = _owner_issued_projection(tmp_path)

    with pytest.raises(TypeError):
        projection.wire["commitId"] = "forged"  # type: ignore[index]
    with pytest.raises(TypeError):
        projection.wire["facts"][0]["value"] = 999.0  # type: ignore[index]


def test_round_four_writer_factory_rejects_raw_tables_and_unbound_view_model(
    tmp_path: Path,
) -> None:
    with _server_owned_evidence(
        tmp_path / "owner-store",
        canonical_project_root=tmp_path / "canonical-project",
    ) as (record, _binding, owner):
        valid_wire = _owner_wire(owner, _report_projection_wire())
        valid_wire["envelopeDigest"] = record.envelope_digest
        projection = owner.issue_report_projection(valid_wire)
        with pytest.raises(TypeError, match="raw|typed|server-owned|capability"):
            owner.issue_report_evidence_writer(
                projection,
                sections=("facts",),
                table_rows={"facts": [{"value": 999.0}]},
            )
        with pytest.raises(TypeError, match="bound|typed|server-owned|projection"):
            owner.issue_report_document_writer(
                {"title": "forged"},
                projection,
                canonical_run_root=tmp_path / "canonical-project" / "runs" / "run-1",
            )


def test_round_four_document_writer_rejects_a_foreign_run_root(tmp_path: Path) -> None:
    canonical_root = tmp_path / "canonical-project"
    with _server_owned_evidence(
        tmp_path / "owner-store",
        canonical_project_root=canonical_root,
    ) as (record, _binding, owner):
        valid_wire = _owner_wire(owner, _report_projection_wire())
        valid_wire["envelopeDigest"] = record.envelope_digest
        projection = owner.issue_report_projection(valid_wire)
        with pytest.raises(TypeError, match="root|owner|canonical"):
            owner.issue_report_document_writer(
                {"title": "forged"},
                projection,
                canonical_run_root=tmp_path / "foreign-project" / "runs" / "run-1",
            )


def test_round_four_index_cannot_synchronize_forged_facts_or_digest(
    tmp_path: Path,
) -> None:
    with _server_owned_evidence(tmp_path) as (record, binding, capability):
        indexed_artifact, payload, content_hash = _canonical_indexed_artifact(record, binding)
        indexed_artifact["facts"] = [{"factId": "fact-forged", "statisticId": "estimate", "value": 999.0}]
        indexed_artifact["projectionDigest"] = "f" * 64
        payload["facts"] = [{"factId": "fact-forged", "statisticId": "estimate", "value": 999.0}]
        payload["projectionDigest"] = "f" * 64

        view = build_indexed_report_evidence(
            indexed_artifact,
            payload,
            artifact_content_sha256=content_hash,
            authoritative_persistence=capability,
        )

    assert view["status"] == "blocked"
    assert view["reasonCode"] in {"REPORT_EVIDENCE_FACTS_MISMATCH", "REPORT_EVIDENCE_PROJECTION_DIGEST_MISMATCH"}
    assert view["facts"] == []


def test_round_five_imported_factory_token_and_forged_owner_fail_before_projection_io() -> None:
    """A self-consistent caller wire cannot mint an exportable projection."""

    import workbench.report_export as report_export_module

    forged_owner = {"owner": "caller-created"}
    wire = _report_projection_wire()
    wire["projectionDigest"] = report_export_module.report_projection_digest(wire)

    with pytest.raises(TypeError, match="owner|capability|server-owned"):
        EvidenceProjectionV1._from_server_wire(wire, owner=forged_owner)

    imported_token = getattr(report_export_module, "_REPORT_FACTORY_TOKEN", object())
    with pytest.raises(TypeError, match="owner|capability|server-owned"):
        EvidenceProjectionV1(
            wire,
            _factory_token=imported_token,
            _owner=forged_owner,
        )


def test_round_five_forged_and_serialized_projection_objects_fail_before_writer_io() -> None:
    """A caller-shaped projection object cannot reach the XLSX writer."""

    import workbench.report_export as report_export_module

    forged_owner = {"owner": "caller-created"}
    wire = _report_projection_wire()
    wire["projectionDigest"] = report_export_module.report_projection_digest(wire)
    forged_projection = object.__new__(EvidenceProjectionV1)
    object.__setattr__(forged_projection, "wire", wire)
    object.__setattr__(forged_projection, "_owner", forged_owner)
    object.__setattr__(forged_projection, "_authority_marker", object())
    try:
        serialized_projection = pickle.loads(pickle.dumps(forged_projection))
    except TypeError:
        # A server-issued projection is also allowed to reject serialization
        # outright; the original forged object must still fail at the writer.
        serialized_projection = None

    for candidate in (forged_projection, serialized_projection):
        if candidate is None:
            continue
        with pytest.raises(TypeError, match="owner|capability|server-owned"):
            ReportEvidenceWriterV1._from_server_projection(
                candidate,
                owner=forged_owner,
                sections=("facts",),
            )


def test_round_five_real_owner_rejects_a_self_consistent_forged_wire_before_writer_io(
    tmp_path: Path,
) -> None:
    with _server_owned_evidence(tmp_path / "owner-store") as (record, _binding, owner):
        wire = _owner_wire(owner, _report_projection_wire())
        wire["envelopeDigest"] = record.envelope_digest
        wire["facts"] = [{"factId": "fact-forged", "statisticId": "estimate", "value": 999.0}]
        import workbench.report_export as report_export_module

        wire["projectionDigest"] = report_export_module.report_projection_digest(wire)

        with pytest.raises(ReportExportError, match="OWNER_WIRE"):
            EvidenceProjectionV1._from_server_wire(wire, owner=owner)


def test_round_five_unowned_xlsx_render_rejects_before_workbook_io(monkeypatch) -> None:
    import workbench.exports as exports_module
    import workbench.report_export as report_export_module

    forged_owner = {"owner": "caller-created"}
    wire = _report_projection_wire()
    wire["projectionDigest"] = report_export_module.report_projection_digest(wire)
    forged_projection = object.__new__(EvidenceProjectionV1)
    object.__setattr__(forged_projection, "wire", wire)
    object.__setattr__(forged_projection, "_owner", forged_owner)
    object.__setattr__(forged_projection, "_authority_marker", object())
    forged_writer = object.__new__(ReportEvidenceWriterV1)
    object.__setattr__(forged_writer, "projection", forged_projection)
    object.__setattr__(forged_writer, "sections", ("facts",))
    object.__setattr__(forged_writer, "tables_override", None)
    monkeypatch.setattr(
        exports_module,
        "Workbook",
        lambda: pytest.fail("unowned writer reached XLSX workbook IO"),
    )

    with pytest.raises(TypeError, match="owner|capability|server-owned"):
        exports_module.render_typed_xlsx(forged_writer)


def test_remediation_three_imported_factory_roots_and_untrusted_owner_fail_before_io(
    tmp_path: Path,
    monkeypatch,
) -> None:
    """Importing a legacy token or selecting a caller-owned owner cannot reach persistence IO."""

    import workbench.report_export as report_export_module
    import workbench.services.results_service as results_service_module

    assert not hasattr(report_export_module, "_REPORT_FACTORY_TOKEN")
    assert not hasattr(results_service_module, "_REPORT_OWNER_FACTORY_TOKEN")

    with _server_owned_evidence(tmp_path / "owner-store") as (_record, _binding, owner):
        monkeypatch.setattr(
            type(owner._evidence_store),
            "load",
            lambda *_args, **_kwargs: pytest.fail("untrusted owner construction reached persistence IO"),
        )
        with pytest.raises((TypeError, EvidenceProjectionError), match="host|authority|capability"):
            ServerOwnedReportEvidencePersistenceV1._from_host(
                owner._evidence_store,
                owner._registry,
                owner._receipt,
                actor_ref=owner.actor_ref,
                project_ref=owner.project_ref,
                session_ref=owner.session_ref,
                scope_ref=owner.scope_ref,
                run_id=owner.run_id,
                model_run_id=owner.model_run_id,
                artifact_id=owner.artifact_id,
                commit_id=owner.commit_id,
                confirmation_ref=owner.confirmation_ref,
                artifact_file_sha256=owner.artifact_file_sha256,
                canonical_project_root=owner.canonical_project_root,
            )


@contextmanager
def _server_host_issued_report_evidence(tmp_path: Path):
    """Build a Report owner only from the real test server's admitted endpoint."""

    from tests.test_v189_evidence_envelope import _admit, _admit_store, _authorized_registry

    tmp_path.mkdir(parents=True, exist_ok=True)
    registry, host = _authorized_registry("scope:project:1")
    record = _admit(
        registry,
        json.loads(_FIXTURE.read_text(encoding="utf-8")),
        artifact_id="artifact-1",
        scope_ref="scope:project:1",
        execution_ref="execution:run-1",
        contract_digest=_CONSUMER_PROJECTION_EVIDENCE_DIGEST,
    )
    with _admit_store(tmp_path, registry) as store:
        receipt = registry.persistence_receipt(record)
        binding = store.commit(record, receipt=receipt)
        owner = ServerOwnedReportEvidencePersistenceV1._from_host(
            store,
            registry,
            receipt,
            authority_handle=host.channel,
            actor_ref="actor:1",
            project_ref="project:1",
            session_ref="session:1",
            scope_ref="scope:project:1",
            run_id="run-1",
            model_run_id="run-1",
            artifact_id=record.artifact_id,
            commit_id=binding.commit_id,
            confirmation_ref="confirmation:run-1",
            artifact_file_sha256=hashlib.sha256(b"server-host-artifact").hexdigest(),
            canonical_project_root=tmp_path / "project",
        )
        yield record, binding, owner, host


def test_remediation_three_controlled_server_host_harness_issues_report_owner(tmp_path: Path) -> None:
    with _server_host_issued_report_evidence(tmp_path / "owner-store") as (_record, _binding, owner, host):
        projection = owner.issue_report_projection(owner._canonical_report_wire())
        assert projection.wire["status"] == "available"
        assert any(key[0] == "evidence.report.owner" for key in host.issued)


def test_remediation_three_forged_authority_and_object_new_owner_fail_before_io(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    with _server_host_issued_report_evidence(tmp_path / "owner-store") as (
        _record,
        _binding,
        owner,
        _host,
    ):
        owner_wire = owner._canonical_report_wire()
        store = owner._evidence_store
        registry = owner._registry
        receipt = owner._receipt
        monkeypatch.setattr(
            type(store),
            "load",
            lambda *_args, **_kwargs: pytest.fail("forged owner reached persistence IO"),
        )
        with pytest.raises((TypeError, EvidenceProjectionError), match="owner|authority|capability"):
            type(owner)._from_host(
                store,
                registry,
                receipt,
                actor_ref=owner.actor_ref,
                project_ref=owner.project_ref,
                session_ref=owner.session_ref,
                scope_ref=owner.scope_ref,
                run_id=owner.run_id,
                model_run_id=owner.model_run_id,
                artifact_id=owner.artifact_id,
                commit_id=owner.commit_id,
                confirmation_ref=owner.confirmation_ref,
                artifact_file_sha256=owner.artifact_file_sha256,
                canonical_project_root=owner.canonical_project_root,
                authority_handle=object(),
            )

        forged_owner = object.__new__(type(owner))
        for slot in type(owner).__slots__:
            if slot == "__weakref__":
                continue
            object.__setattr__(forged_owner, slot, getattr(owner, slot))
        with pytest.raises((TypeError, EvidenceProjectionError), match="owner|authority|capability"):
            forged_owner.issue_report_projection(owner_wire)


def test_remediation_three_configuration_rejects_object_new_owner(
    tmp_path: Path,
) -> None:
    from workbench.services.results_service import configure_report_evidence_persistence

    with _server_host_issued_report_evidence(tmp_path / "owner-store") as (
        _record,
        _binding,
        owner,
        _host,
    ):
        forged_owner = object.__new__(type(owner))
        for slot in type(owner).__slots__:
            if slot == "__weakref__":
                continue
            object.__setattr__(forged_owner, slot, getattr(owner, slot))

        with pytest.raises((TypeError, EvidenceProjectionError), match="owner|authority|capability"):
            configure_report_evidence_persistence(forged_owner)


def test_copied_valid_report_pair_in_foreign_project_root_is_typed_rejected(
    tmp_path: Path,
    monkeypatch,
) -> None:
    canonical_root = tmp_path / "canonical-project"
    foreign_root = tmp_path / "foreign-project"
    run_root = foreign_root / "runs" / "run-1"
    artifact_path = run_root / "artifacts" / "post_estimation" / "result.json"
    artifact_path.parent.mkdir(parents=True)
    (run_root / "run_manifest.json").write_text(
        json.dumps({"run_id": "run-1", "status": "completed", "mode": "auto", "lineage": []}),
        encoding="utf-8",
    )

    owner_root = tmp_path / "owner-store"
    owner_root.mkdir()
    with _server_owned_evidence(owner_root) as (record, binding, capability):
        index, payload, _content_hash = _canonical_indexed_artifact(record, binding)
        index["path"] = "artifacts/post_estimation/result.json"
        index.pop("evidence_commit_binding")
        index.pop("evidence_envelope")
        raw_payload = json.dumps(payload, separators=(",", ":"), sort_keys=True).encode()
        artifact_path.write_bytes(raw_payload)
        (run_root / "artifacts_index.json").write_text(
            json.dumps({"schema_version": 1, "artifacts": [index]}),
            encoding="utf-8",
        )
        monkeypatch.setattr(runs_routes, "current_report_evidence_persistence", lambda **_: capability)

        response = TestClient(app).get(
            "/runs/run-1",
            params={"project_root": str(foreign_root)},
        )

    assert canonical_root != foreign_root
    assert response.status_code == 200, response.text
    evidence = response.json()["report_evidence"]
    assert evidence["status"] in {"unavailable", "blocked"}
    assert evidence["reasonCode"] in {
        "EVIDENCE_PROJECT_ROOT_MISMATCH",
        "EVIDENCE_SCOPE_ROOT_MISMATCH",
        "EVIDENCE_OWNER_CONTEXT_MISMATCH",
        "EVIDENCE_PERSISTENCE_CAPABILITY_REQUIRED",
    }
    assert evidence["facts"] == []


def test_http_run_detail_never_returns_raw_model_results(tmp_path: Path) -> None:
    project_root = tmp_path / "project"
    run_root = project_root / "runs" / "run-1"
    (run_root / "model_results").mkdir(parents=True)
    (run_root / "run_manifest.json").write_text(
        json.dumps({"run_id": "run-1", "status": "completed", "mode": "auto", "lineage": []}),
        encoding="utf-8",
    )
    (run_root / "model_results" / "ols_1.json").write_text(
        json.dumps({"model_id": "ols_1", "coefficients": {"leaked": {"estimate": 999.0}}}),
        encoding="utf-8",
    )

    response = TestClient(app).get("/runs/run-1", params={"project_root": str(project_root)})

    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["model_results"] == []
    assert payload.get("model_family_evidence") in (None, {})
    assert "999.0" not in json.dumps(payload, sort_keys=True)


def test_result_table_export_missing_evidence_is_a_typed_unavailable_contract(tmp_path: Path) -> None:
    project_root = tmp_path / "project"
    run_root = project_root / "runs" / "run-1"
    run_root.mkdir(parents=True)
    (run_root / "run_manifest.json").write_text(
        json.dumps({"run_id": "run-1", "status": "completed", "mode": "auto", "lineage": []}),
        encoding="utf-8",
    )

    response = TestClient(app).post(
        "/runs/run-1/report/result-table-export",
        params={"project_root": str(project_root)},
        json={"sections": ["regression_table"]},
    )

    assert response.status_code != 500
    assert response.headers["content-type"].startswith("application/json")
    body = response.json()
    projection = body.get("error", {}).get("details", {}).get("projection")
    assert projection["status"] == "unavailable"
    assert projection["facts"] == []


def test_result_table_export_valid_commit_downloads_only_typed_facts(tmp_path: Path, monkeypatch) -> None:
    project_root = tmp_path / "project"
    run_root = project_root / "runs" / "run-1"
    run_root.mkdir(parents=True)
    (run_root / "run_manifest.json").write_text(
        json.dumps({"run_id": "run-1", "status": "completed", "mode": "auto", "lineage": []}),
        encoding="utf-8",
    )
    packet = {"model_id": "ols_1", "model_type": "ols", "nobs": 30,
              "coefficients": {"x": {"estimate": 2.0, "std_error": 0.2, "p_value": 1e-24}}}
    with _produced_report(run_root, model_results={"ols_1": packet}, diagnostic_summary={}) as case:
        monkeypatch.setattr(
            runs_routes,
            "current_report_evidence_persistence",
            lambda **_kwargs: case.owner,
        )
        expected = case.owner.issue_report_evidence_writer(
            case.projection, sections=("regression_table",),
        ).tables()["regression_table"]
        response = TestClient(app).post(
            "/runs/run-1/report/result-table-export",
            params={"project_root": str(project_root)},
            json={"sections": ["regression_table"]},
        )

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
    values = list(workbook["regression_table"].values)
    workbook.close()
    assert len(values) == 2
    row = dict(zip(values[0], values[1]))
    assert row["section_type"] == "regression_table" and row["status"] == "available"
    assert row["term"] == "x"
    cell = json.loads(row["models"])["ols_1"]
    assert cell["estimate"] == 2.0 and cell["p_value"] == 1e-24
    for key, value in expected[0].items():
        observed = json.loads(row[key]) if isinstance(value, (dict, list, tuple)) else row[key]
        assert sha256_canonical(observed) == sha256_canonical(value), key


def test_result_table_export_download_failure_is_typed_blocked(tmp_path: Path, monkeypatch) -> None:
    project_root = tmp_path / "project"
    run_root = project_root / "runs" / "run-1"
    artifact_path = run_root / "artifacts" / "post_estimation" / "result.json"
    artifact_path.parent.mkdir(parents=True)
    (run_root / "run_manifest.json").write_text(
        json.dumps({"run_id": "run-1", "status": "completed", "mode": "auto", "lineage": []}),
        encoding="utf-8",
    )
    owner_root = tmp_path / "owner-store"
    owner_root.mkdir()
    with _server_owned_evidence(owner_root, canonical_project_root=project_root) as (record, binding, capability):
        index, payload, _content_hash = _canonical_indexed_artifact(record, binding)
        index["path"] = "artifacts/post_estimation/result.json"
        index.pop("evidence_commit_binding")
        index.pop("evidence_envelope")
        artifact_path.write_bytes(json.dumps(payload, separators=(",", ":"), sort_keys=True).encode())
        (run_root / "artifacts_index.json").write_text(
            json.dumps({"schema_version": 1, "artifacts": [index]}),
            encoding="utf-8",
        )
        monkeypatch.setattr(
            runs_routes,
            "current_report_evidence_persistence",
            lambda **_kwargs: capability,
        )
        monkeypatch.setattr(
            runs_routes,
            "render_typed_xlsx",
            lambda _tables: (_ for _ in ()).throw(OSError("download failed")),
        )
        response = TestClient(app).post(
            "/runs/run-1/report/result-table-export",
            params={"project_root": str(project_root)},
            json={"sections": ["regression_table"]},
        )

    assert response.status_code == 409
    assert response.headers["content-type"].startswith("application/json")
    error = response.json()["error"]
    assert error["code"] == "RESULT_TABLE_EXPORT_FAILED"
    assert error["details"]["projection"]["status"] == "blocked"
    assert error["details"]["projection"]["facts"] == []


def test_remediation_six_missing_legacy_admission_fails_before_any_presentation_io(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """include_legacy_presentation cannot make an unadmitted summary readable."""

    import workbench.report_view_model as report_view_model_module

    def fail_before_io(*_args: object, **_kwargs: object) -> object:
        pytest.fail("legacy presentation reached a filesystem/template seam before admission")

    monkeypatch.setattr(report_view_model_module, "_build_legacy_presentation", fail_before_io)
    monkeypatch.setattr(report_view_model_module, "render_template", fail_before_io)
    monkeypatch.setattr(report_view_model_module, "read_json", fail_before_io)

    view_model = build_report_view_model(
        {
            "diagnostics": {
                "blockers": [{"message": "caller-shaped"}],
                "warnings": [],
                "cautions": [],
                "info": [],
            },
            "coefficients_summary": {"rows": []},
        },
        tmp_path,
        include_legacy_presentation=True,
    )

    assert view_model["status"] in {"blocked", "unavailable"}
    assert view_model["facts"] == []
    assert view_model["typedFacts"] == []


def test_remediation_six_report_evidence_admits_owner_before_index_or_hash(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The evidence entrance must admit the C0 owner before artifact reads."""

    project_root = tmp_path / "project"
    run_root = project_root / "runs" / "run-1"
    artifact_path = run_root / "artifacts" / "post_estimation" / "result.json"
    artifact_path.parent.mkdir(parents=True)

    with _server_owned_evidence(
        tmp_path / "owner-store",
        canonical_project_root=project_root,
    ) as (record, binding, owner):
        indexed_artifact, payload, content_hash = _canonical_indexed_artifact(record, binding)
        indexed_artifact["path"] = "artifacts/post_estimation/result.json"
        indexed_artifact.pop("evidence_commit_binding")
        indexed_artifact.pop("evidence_envelope")
        artifact_path.write_bytes(json.dumps(payload, separators=(",", ":"), sort_keys=True).encode())

        events: list[str] = []

        def admit(**_kwargs: object) -> object:
            events.append("owner-admission")
            return owner

        monkeypatch.setattr(runs_routes, "current_report_evidence_persistence", admit)
        monkeypatch.setattr(
            runs_routes,
            "_read_artifact_records",
            lambda _root: (events.append("artifact-index-read") or [indexed_artifact]),
        )
        monkeypatch.setattr(
            runs_routes,
            "_resolve_artifact_path",
            lambda _root, _record: (events.append("artifact-path-resolve") or artifact_path),
        )
        monkeypatch.setattr(
            runs_routes,
            "sha256_file",
            lambda _path: (events.append("artifact-hash") or content_hash),
        )

        view = runs_routes._report_evidence_for_run(
            run_root,
            [],
            project_root=project_root,
        )

    assert view["status"] == "available"
    assert events[0] == "owner-admission"
    assert events.index("owner-admission") < events.index("artifact-index-read")
    assert events.index("owner-admission") < events.index("artifact-hash")


def test_remediation_six_report_export_admits_owner_before_index_or_hash(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The export owner resolver cannot inspect the artifact before admission."""

    project_root = tmp_path / "project"
    run_root = project_root / "runs" / "run-1"
    artifact_path = run_root / "artifacts" / "post_estimation" / "result.json"
    artifact_path.parent.mkdir(parents=True)

    with _server_owned_evidence(
        tmp_path / "owner-store",
        canonical_project_root=project_root,
    ) as (record, binding, owner):
        indexed_artifact, payload, content_hash = _canonical_indexed_artifact(record, binding)
        indexed_artifact["path"] = "artifacts/post_estimation/result.json"
        artifact_path.write_bytes(json.dumps(payload, separators=(",", ":"), sort_keys=True).encode())
        projection = dict(owner._canonical_report_wire())  # type: ignore[attr-defined]
        events: list[str] = []

        def admit(**_kwargs: object) -> object:
            events.append("owner-admission")
            return owner

        monkeypatch.setattr(runs_routes, "current_report_evidence_persistence", admit)
        monkeypatch.setattr(
            runs_routes,
            "_read_artifact_records",
            lambda _root: (events.append("artifact-index-read") or [indexed_artifact]),
        )
        monkeypatch.setattr(
            runs_routes,
            "_resolve_artifact_path",
            lambda _root, _record: (events.append("artifact-path-resolve") or artifact_path),
        )
        monkeypatch.setattr(
            runs_routes,
            "sha256_file",
            lambda _path: (events.append("artifact-hash") or content_hash),
        )

        resolved_owner = runs_routes._report_owner_for_export(
            run_root,
            projection,
            None,
            project_root=project_root,
        )

    assert resolved_owner is owner
    assert events[0] == "owner-admission"
    assert events.index("owner-admission") < events.index("artifact-index-read")
    assert events.index("owner-admission") < events.index("artifact-hash")


def test_remediation_six_real_owner_wildcard_admission_is_server_bound(
    tmp_path: Path,
) -> None:
    """The pre-index admission uses the live global C0 owner, not index fields."""

    from workbench.services.results_service import (
        configure_report_evidence_persistence,
        current_report_evidence_persistence,
    )

    project_root = tmp_path / "project"
    run_root = project_root / "runs" / "run-1"
    with _server_owned_evidence(
        tmp_path / "owner-store",
        canonical_project_root=project_root,
    ) as (_record, _binding, owner):
        configure_report_evidence_persistence(owner)
        try:
            identity = owner._require_report_request_identity(None)  # type: ignore[attr-defined]
            admitted = runs_routes._admit_report_owner_for_run(
                run_root,
                identity,
                project_root,
            )
            direct = current_report_evidence_persistence(
                actor_ref=owner.actor_ref,
                run_id=owner.run_id,
                project_ref=owner.project_ref,
                session_ref=owner.session_ref,
                scope_ref=owner.scope_ref,
                model_run_id=owner.model_run_id,
                artifact_id=None,
                commit_id=None,
                confirmation_ref=owner.confirmation_ref,
                artifact_file_sha256=None,
                project_root=project_root,
            )
        finally:
            configure_report_evidence_persistence(None)

    assert admitted is owner
    assert direct is owner


def test_remediation_six_exact_identity_rejects_copied_valid_projection_and_writers(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Copied valid fields do not constitute a server-issued projection/writer."""

    def clone(source: object) -> object:
        cloned = object.__new__(type(source))
        for slot in type(source).__slots__:  # type: ignore[attr-defined]
            if slot == "__weakref__":
                continue
            try:
                value = getattr(source, slot)
            except AttributeError:
                continue
            object.__setattr__(cloned, slot, value)
        return cloned

    attacks_reached = []
    project_root = tmp_path / "project"
    run_root = project_root / "runs" / "run-1"
    with _server_owned_evidence(
        tmp_path / "owner-store",
        canonical_project_root=project_root,
        keep_open=True,
    ) as (record, _binding, owner):
        projection = owner.issue_report_projection(owner._canonical_report_wire())  # type: ignore[attr-defined]
        evidence_writer = owner.issue_report_evidence_writer(projection, sections=("facts",))
        projection._assert_server_owned()
        evidence_writer._assert_server_owned()

        monkeypatch.setattr(
            type(owner),
            "_canonical_report_wire",
            lambda _owner: pytest.fail("copied projection reached canonical evidence IO"),
        )
        with pytest.raises(TypeError, match="issuer|identity|owner|projection"):
            attacks_reached.append("projection")
            clone(projection)._assert_server_owned()  # type: ignore[attr-defined]

        monkeypatch.setattr(
            EvidenceProjectionV1,
            "_assert_server_owned",
            lambda _projection: pytest.fail("copied evidence writer reached projection/facts IO"),
        )
        with pytest.raises(TypeError, match="issuer|identity|writer|owner"):
            attacks_reached.append("evidence_writer")
            clone(evidence_writer)._assert_server_owned()  # type: ignore[attr-defined]

    monkeypatch.undo()
    document_writer = _owner_issued_document_writer(
        tmp_path / "document-owner",
        run_root,
    )
    document_writer._assert_server_owned()
    cloned_document_writer = clone(document_writer)
    monkeypatch.setattr(
        reporting,
        "Environment",
        lambda *_args, **_kwargs: pytest.fail("copied document writer reached template IO"),
    )
    with pytest.raises(TypeError, match="issuer|identity|writer|owner"):
        attacks_reached.append("document_writer")
        reporting.render_html_report(cloned_document_writer, run_root)  # type: ignore[arg-type]
    assert attacks_reached == ["projection", "evidence_writer", "document_writer"]


def test_remediation_six_report_stage_adapts_exact_owner_projection_and_fails_closed_without_one(
    tmp_path: Path,
) -> None:
    """ReportStage receives only the exact owner/projection pair from context."""

    from workbench.engine.stages.report import _report_evidence_admission

    assert _report_evidence_admission(SimpleNamespace(artifacts={})) is None

    project_root = tmp_path / "project"
    run_root = project_root / "runs" / "run-1"
    with _server_owned_evidence(
        tmp_path / "owner-store",
        run_id="run-1",
        model_run_id="run-1",
        canonical_project_root=project_root,
        keep_open=True,
    ) as (record, _binding, owner):
        wire = dict(owner._canonical_report_wire())  # type: ignore[attr-defined]
        projection = owner.issue_report_projection(wire)
        context = SimpleNamespace(
            artifacts={
                "_report_evidence_owner": owner,
                "_report_evidence_projection": projection,
            }
        )

        admitted = _report_evidence_admission(context)

    assert admitted is not None
    assert admitted[0] is owner
    assert admitted[1] is projection


def test_report_owner_admission_rejects_a_locally_fabricated_confirmation(
    tmp_path: Path,
) -> None:
    from workbench.artifacts import initialize_artifact_index
    from workbench.engine.stages.report import _admit_report_evidence_from_factory
    from workbench.services.report_evidence_owner_factory import (
        ServerOwnedReportEvidenceOwnerFactoryV1,
    )

    project_root = tmp_path / "project"
    run_root = project_root / "runs" / "run-1"
    run_root.mkdir(parents=True)
    initialize_artifact_index(run_root)
    regression_table = {
        "models": [{"id": "ols_1", "nobs": 1}],
        "rows": [{"term": "x", "models": {"ols_1": {"estimate": 1.0}}}],
    }

    with _server_owned_evidence(
        tmp_path / "owner-store",
        run_id="run-1",
        model_run_id="ols_1",
        canonical_project_root=project_root,
        keep_open=True,
    ) as (_record, _binding, owner):
        factory = ServerOwnedReportEvidenceOwnerFactoryV1.from_host(
            owner._evidence_store,  # type: ignore[attr-defined]
            owner._registry,  # type: ignore[attr-defined]
            actor_ref=owner.actor_ref,
            project_ref=owner.project_ref,
            session_ref=owner.session_ref,
            scope_ref=owner.scope_ref,
            canonical_project_root=project_root,
            authority_handle=owner._authority_handle,  # type: ignore[attr-defined]
        )
        context = SimpleNamespace(
            artifacts={
                "_report_evidence_owner_factory": factory,
                "_upload_hash": "a" * 64,
                "_report_confirmation_ref": "confirmation:run-1",
            }
        )

        with pytest.raises(EvidenceProjectionError, match="confirmation"):
            _admit_report_evidence_from_factory(
                context,
                run_root=run_root,
                run_id="run-1",
                model_run_id="ols_1",
                primary_model_type="ols",
                regression_table=regression_table,
                issue_dicts=(),
            )


def _arma_report_stage_context(
    run_root: Path,
    *,
    owner: ServerOwnedReportEvidencePersistenceV1 | None = None,
    projection: EvidenceProjectionV1 | None = None,
) -> SimpleNamespace:
    artifacts: dict[str, object] = {
        "_cleaned": object(),
        "_categorical_vars": [],
        "_model_results": [("arma_garch_1", {"model_type": "time_series.arma_garch"})],
        "_normalized_y": "value",
        "_normalized_x": [],
        "_poisson_x": [],
        "_profile": {"row_count": 1, "column_count": 1},
        "_routing": {"kind": "time_series"},
        "_dropped_vars": [],
        "_reliability_info": {},
        "_effective_exposure_col": None,
        "_imputation_summary": {},
        "_diagnostic_artifacts": {},
        "_statistical_tests": {"evidence": None},
        "_statistical_test_summaries": [],
        "_claims": [],
        "_issue_dicts": [],
        "_coercion_actions": [],
        "_data_detected_y_type": "continuous",
        "_input_files": [],
        "_mode": "auto",
        "_y": "value",
        "_x": [],
        "_model_type": "time_series.arma_garch",
        "_started_at": "2026-08-20T00:00:00+00:00",
    }
    if owner is not None:
        artifacts["_report_evidence_owner"] = owner
    if projection is not None:
        artifacts["_report_evidence_projection"] = projection
    return SimpleNamespace(
        artifacts=artifacts,
        exposure_col=None,
        y_type="continuous",
        roles={},
        terminal_status=None,
        primary_type=None,
    )


def _patch_report_stage_side_effects(monkeypatch: pytest.MonkeyPatch) -> dict[str, list[object]]:
    from workbench import artifacts as artifact_module
    from workbench import diagnostic_summary as diagnostic_summary_module
    from workbench import orchestrator as orchestrator_module
    from workbench import report_view_model as report_view_model_module
    from workbench import reporting as reporting_module
    from workbench.engine.stages import report as report_module

    events: dict[str, list[object]] = {"view_models": [], "rendered": []}
    monkeypatch.setattr(orchestrator_module, "_build_descriptive_stats", lambda *_a, **_k: [])
    monkeypatch.setattr(orchestrator_module, "_build_variable_importance", lambda *_a, **_k: [])
    monkeypatch.setattr(orchestrator_module, "_build_model_routing_summary", lambda *_a, **_k: {})
    monkeypatch.setattr(orchestrator_module, "_lineage", lambda *_a, **_k: [])
    monkeypatch.setattr(orchestrator_module, "_safe_flush_recorder", lambda *_a, **_k: None)
    monkeypatch.setattr(orchestrator_module, "_write_manifest", lambda *_a, **_k: None)
    monkeypatch.setattr(orchestrator_module, "build_diagnostic_summary", lambda **_k: {})
    monkeypatch.setattr(report_view_model_module, "build_regression_table", lambda *_a, **_k: [])
    monkeypatch.setattr(
        report_view_model_module,
        "build_report_view_model",
        lambda *_a, **_k: events["view_models"].append({
            "status": "unavailable",
            "facts": [],
            "reasonCode": "EVIDENCE_OWNER_CAPABILITY_REQUIRED",
        }) or events["view_models"][-1],
    )
    monkeypatch.setattr(diagnostic_summary_module, "finalize_diagnostic_summary", lambda value, **_k: value)
    monkeypatch.setattr(artifact_module, "write_json", lambda *_a, **_k: None)
    monkeypatch.setattr(artifact_module, "register_artifact", lambda *_a, **_k: None)
    monkeypatch.setattr(report_module, "_refresh_artifact_checksum", lambda *_a, **_k: None)
    monkeypatch.setattr(
        reporting_module,
        "render_legacy_html_report",
        lambda view_model, _run_root, *, owner, projection: events["rendered"].append(
            (view_model, owner, projection)
        ),
    )
    return events


def _run_report_stage_for_probe(ctx: SimpleNamespace, run_root: Path, monkeypatch: pytest.MonkeyPatch):
    from workbench.engine.context import RunEnv
    from workbench.engine.stages.report import ReportStage

    events = _patch_report_stage_side_effects(monkeypatch)
    env = RunEnv(run_root=run_root, run_id=run_root.name, recorder=SimpleNamespace())
    return ReportStage().run(ctx, env), events


def test_orchestrator_report_admission_injects_exact_pair_before_real_report_stage(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from workbench import orchestrator
    from workbench.app import app
    from workbench.lineage import incremental

    project_root = tmp_path / "project"
    run_root = project_root / "runs" / "run-1"
    run_root.mkdir(parents=True)
    source = tmp_path / "source.csv"
    source.write_text("value\n1\n", encoding="utf-8")
    config = SimpleNamespace(
        random_seed=1,
        imputation_method="",
        imputation_m=1,
        imputation_max_iter=1,
        max_missing_rate=0.4,
        run_timeout_s=10.0,
    )
    captured: dict[str, object] = {}

    def capture_pipeline(pipeline, ctx, _env, **_kwargs):
        captured["pipeline"] = pipeline
        captured["ctx"] = ctx
        ctx.terminal_status = "completed"
        return ctx

    monkeypatch.setattr(incremental, "run_pipeline_traced", capture_pipeline)
    with _server_owned_evidence(
        tmp_path / "owner-store",
        run_id="run-1",
        model_run_id="run-1",
        canonical_project_root=project_root,
        keep_open=True,
    ) as (_record, _binding, owner):
        configure_report_evidence_owner(owner)
        try:
            orchestrator._run_workflow(
                run_root,
                "run-1",
                [source],
                "auto",
                "value",
                [],
                config,
                "2026-08-20T00:00:00+00:00",
                model_type="ols",
            )
        finally:
            configure_report_evidence_owner(None)

    ctx = captured["ctx"]
    assert any(type(stage).__name__ == "ReportStage" for stage in captured["pipeline"])
    assert ctx.artifacts["_report_evidence_owner"] is owner
    projection = ctx.artifacts["_report_evidence_projection"]
    assert type(projection) is EvidenceProjectionV1
    assert projection._owner is owner
    assert app.state.report_evidence_persistence is None


def test_arma_garch_report_admission_missing_is_typed_and_does_not_build_raw_deliverable(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from workbench.engine.packs.arma_garch import deliverables

    run_root = tmp_path / "project" / "runs" / "run-1"
    run_root.mkdir(parents=True)
    ctx = _arma_report_stage_context(run_root)
    calls: list[Path] = []

    def raw_builder(path: Path) -> dict[str, object]:
        calls.append(path)
        raise AssertionError("ARMA raw deliverable builder reached before Report admission")

    monkeypatch.setattr(deliverables, "build_arma_garch_deliverables", raw_builder)
    result, events = _run_report_stage_for_probe(ctx, run_root, monkeypatch)

    assert result is ctx
    assert calls == []
    assert events["view_models"]
    assert events["view_models"][0]["status"] in {"unavailable", "blocked"}
    assert events["rendered"] == []
    assert not (run_root / "reports" / "report.html").exists()


def test_arma_garch_report_admission_valid_real_workflow_renders_owner_document(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from tests.models.arma_garch.test_pack_runtime import _options, _source
    from workbench import orchestrator
    from fastapi import FastAPI
    from tests.test_v189_evidence_envelope import _admit_store, _authorized_registry
    from workbench.app import configure_report_evidence_owner_factory
    from workbench.artifacts import initialize_artifact_index, read_json, sha256_file
    from workbench.config import WorkbenchConfig
    from workbench.engine.packs.loader import bootstrap_builtin_packs
    from workbench.lineage import incremental
    from workbench.model_options import bind_new_model_options
    from workbench.report_export import LocalReportCollectionV1
    from workbench.services.report_evidence_owner_factory import ServerOwnedReportEvidenceOwnerFactoryV1

    project_root = tmp_path / "project"
    run_root = project_root / "runs" / "run-1"
    run_root.mkdir(parents=True)
    (run_root / "raw_snapshot").mkdir()
    initialize_artifact_index(run_root)
    source = tmp_path / "series.csv"
    _source().to_csv(source, index=False)
    bootstrap_builtin_packs()
    bound_options = bind_new_model_options("time_series.arma_garch", _options())
    assert bound_options.binding is not None
    config = WorkbenchConfig(
        random_seed=20260720,
        imputation_method="",
        imputation_m=1,
        imputation_max_iter=1,
        max_missing_rate=0.4,
        run_timeout_s=180.0,
    )
    captured: dict[str, object] = {}
    original_run_pipeline_traced = incremental.run_pipeline_traced

    def run_pipeline_traced_with_capture(pipeline, ctx, env, **kwargs):
        captured["before_ctx"] = ctx
        result = original_run_pipeline_traced(pipeline, ctx, env, **kwargs)
        captured["after_ctx"] = result
        return result

    monkeypatch.setattr(
        incremental,
        "run_pipeline_traced",
        run_pipeline_traced_with_capture,
    )
    registry, host = _authorized_registry("scope:project:arma-report")
    store_root = tmp_path / "owner-store"
    store_root.mkdir()
    execution_app = FastAPI()
    with _admit_store(store_root, registry) as store:
        factory = ServerOwnedReportEvidenceOwnerFactoryV1.from_host(
            store, registry, actor_ref="actor:arma-report", project_ref="project:arma-report",
            session_ref="session:arma-report", scope_ref="scope:project:arma-report",
            canonical_project_root=project_root, authority_handle=host.channel,
        )
        configure_report_evidence_owner_factory(factory, app_instance=execution_app)
        try:
            result = orchestrator._run_workflow(
                run_root,
                "run-1",
                [source],
                "auto",
                "value",
                [],
                config,
                "2026-08-20T00:00:00+00:00",
                model_type="time_series.arma_garch",
                model_options=_options(),
                model_options_binding=bound_options.binding.to_dict(),
                app_instance=execution_app,
            )
            assert result == {"run_id": "run-1", "status": "completed"}
            ctx = captured["after_ctx"]
            owner = ctx.artifacts["_report_evidence_owner"]
            projection = ctx.artifacts["_report_evidence_projection"]
            assert type(projection) is LocalReportCollectionV1
            assert projection._owner is owner
            projection._assert_server_owned()
            artifact = owner._canonical_report_artifact()
            assert projection.wire["artifactRef"] == artifact.artifact_id
            assert projection.wire["commitId"] == owner._commit_id
            assert projection.wire["envelopeDigest"] == artifact.envelope_digest
            assert projection.wire["readCompleteness"] == "COMPLETE"
            assert owner.accepts_canonical_run_root(run_root)
            runtime = read_json(run_root / "model_results" / "arma_garch_1.json")
            assert projection.wire["metrics"]["mae"] == runtime["validation"]["metrics"]["mae"]
            assert projection.wire["metrics"]["validation_n"] == runtime["validation"]["metrics"]["validation_n"]
            assert projection.wire["parameters"], "specialized report must contain real admitted estimates"
            expected_charts = sorted(key for key in runtime["artifact_refs"] if key.startswith("ts.chart."))
            assert expected_charts
            assert list(projection.wire["chart_artifacts"]) == expected_charts
            expected_forecast = runtime["production_child"]["next_forecast"]["conditional_mean"]
            assert next(row["value"] for row in projection.wire["next_forecast"]
                        if row["field"] == "conditional_mean") == expected_forecast
        finally:
            configure_report_evidence_owner_factory(None, app_instance=execution_app)
            factory.close()
    assert execution_app.state.report_evidence_owner_factory is None

    report_path = run_root / "reports" / "report.html"
    assert report_path.is_file()
    html = report_path.read_text(encoding="utf-8")
    assert 'id="time-series-overview"' in html
    assert "ARMA(1,0)" in html
    assert str(runtime["validation"]["metrics"]["mae"]) in html
    assert str(expected_forecast) in html
    assert all(artifact_id in html for artifact_id in expected_charts)
    index = read_json(run_root / "artifacts_index.json")
    report_record = next(
        item for item in index["artifacts"] if item["artifact_id"] == "report_html"
    )
    assert report_record["path"] == "reports/report.html"
    assert report_record["artifact_type"] == "report"
    assert report_record["sha256"] == sha256_file(report_path)


def test_html_report_endpoint_never_returns_an_unindexed_report_file(tmp_path: Path) -> None:
    project_root = tmp_path / "project"
    run_root = project_root / "runs" / "run-1"
    (run_root / "reports").mkdir(parents=True)
    (run_root / "run_manifest.json").write_text(
        json.dumps({"run_id": "run-1", "status": "completed", "mode": "auto", "lineage": []}),
        encoding="utf-8",
    )
    (run_root / "reports" / "report.html").write_text("<html>raw-secret-999</html>", encoding="utf-8")

    response = TestClient(app).get("/runs/run-1/report", params={"project_root": str(project_root)})

    assert response.status_code != 200 or not response.headers.get("content-type", "").startswith("text/html")
    assert "raw-secret-999" not in response.text
    if response.headers.get("content-type", "").startswith("application/json"):
        projection = response.json().get("error", {}).get("details", {}).get("projection")
        assert projection["status"] in {"unavailable", "blocked"}


def test_report_view_model_does_not_recompute_raw_summary_values(tmp_path: Path) -> None:
    view_builder = __import__("workbench.report_view_model", fromlist=["build_report_view_model"]).build_report_view_model
    view_model = view_builder(
        {"model_identity": {"model_label": "raw-secret-123"}, "coefficients_summary": {"rows": []}},
        tmp_path,
    )

    assert view_model["status"] == "unavailable"
    assert view_model["facts"] == []
    assert "raw-secret-123" not in json.dumps(view_model, sort_keys=True)

    malformed_projection = {
        "status": "available",
        "facts": [{"factId": "raw-fact", "statisticId": "estimate", "value": "raw-secret-456"}],
    }
    blocked = view_builder({}, tmp_path, evidence_projection=malformed_projection)
    assert blocked["status"] in {"blocked", "unavailable"}
    assert blocked["facts"] == []
    assert "raw-secret-456" not in json.dumps(blocked, sort_keys=True)


def test_report_template_does_not_render_unsupported_raw_sections() -> None:
    from jinja2 import Environment, FileSystemLoader

    environment = Environment(
        loader=FileSystemLoader(Path(__file__).parents[1] / "backend" / "workbench" / "templates"),
        autoescape=True,
    )
    html = environment.get_template("report.html.j2").render(
        view_model={
            "title": "Unavailable report",
            "status": "unavailable",
            "reasonCode": "EVIDENCE_PROJECTION_UNAVAILABLE",
            "facts": [],
            "typedFacts": [],
            "critical_errors": [],
            "warnings": [],
            "cautions": [],
            "system_notes": [],
            "coefficient_interpretations": [],
            "causal_caution": "",
            "regression_table": {
                "models": [{"id": "raw-model", "nobs": 999}],
                "rows": [{"term": "raw-secret", "models": {"raw-model": {"estimate": 999.0}}}],
            },
            "descriptive_stats": [{"column": "raw-column", "mean": 998.0}],
            "statistical_tests": {"y_related": [{"interpretation": "raw-stat-997"}]},
            "statistical_evidence": {"results": [{"statistic": 996.0}]},
            "model_family_evidence": {"raw_family_value": 995.0},
            "model_diagnostics": {"raw-diagnostic": {"durbin_watson": 994.0}},
            "exploration": None,
        }
    )

    assert "EVIDENCE_PROJECTION_UNAVAILABLE" in html
    assert all(value not in html for value in ("999.0", "998.0", "raw-stat-997", "996.0", "995.0", "994.0"))


def test_owner_bound_report_view_and_html_ignore_legacy_raw_sections(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """The ReportStage owner path receives only safe presentation inputs."""

    from jinja2 import Environment, FileSystemLoader
    from workbench import orchestrator as orchestrator_module
    from workbench import report_view_model as report_view_model_module
    from workbench.engine.context import RunEnv
    from workbench.engine.stages.report import ReportStage

    project_root = tmp_path / "project"
    run_root = project_root / "runs" / "run-1"
    run_root.mkdir(parents=True)
    raw_values = (
        "raw-model-family-995",
        "raw-diagnostic-994",
        "raw-statistical-996",
        "raw-regression-999",
    )
    raw_regression = {
        "models": [{"id": "run-1", "nobs": 999}],
        "rows": [{"term": "raw-regression-999", "models": {"run-1": {"estimate": 999.0}}}],
    }
    raw_summary = {
        "model_identity": {"model_label": "raw-model-family-995"},
        "model_family_evidence": {"value": "raw-model-family-995"},
        "diagnostics": {"warnings": [{"message": "raw-diagnostic-994"}]},
        "statistical_evidence": {"results": [{"statistic": "raw-statistical-996"}]},
    }
    with _server_owned_evidence(
        tmp_path / "owner-store",
        run_id="run-1",
        model_run_id="run-1",
        canonical_project_root=project_root,
        keep_open=True,
    ) as (record, _binding, owner):
        projection = owner.issue_report_projection(owner._canonical_report_wire())  # type: ignore[attr-defined]
        ctx = _arma_report_stage_context(
            run_root,
            owner=owner,
            projection=projection,
        )
        ctx.artifacts.update(
            {
                "_model_results": [("run-1", {"model_type": "ols"})],
                "_routing": {"kind": "regression"},
                "_model_type": "ols",
                "_normalized_x": ["x"],
                "_diagnostic_artifacts": raw_summary["diagnostics"],
                "_statistical_tests": raw_summary["statistical_evidence"],
                "_statistical_test_summaries": [
                    {"interpretation": "raw-statistical-996"}
                ],
            }
        )
        events = _patch_report_stage_side_effects(monkeypatch)
        monkeypatch.setattr(
            orchestrator_module,
            "_build_descriptive_stats",
            lambda *_a, **_k: [{"column": "raw-column", "mean": 998.0}],
        )
        monkeypatch.setattr(
            orchestrator_module,
            "build_diagnostic_summary",
            lambda **_k: raw_summary,
        )
        monkeypatch.setattr(
            report_view_model_module,
            "build_regression_table",
            lambda *_a, **_k: raw_regression,
        )
        real_builder = build_report_view_model
        captured: dict[str, object] = {}

        def capture_builder(summary: dict[str, object], root: Path, **kwargs: object) -> dict[str, object]:
            captured["summary"] = summary
            captured["kwargs"] = kwargs
            view = real_builder(summary, root, **kwargs)  # type: ignore[arg-type]
            captured["view"] = view
            return view

        monkeypatch.setattr(
            report_view_model_module,
            "build_report_view_model",
            capture_builder,
        )
        result = ReportStage().run(
            ctx,
            RunEnv(run_root=run_root, run_id="run-1", recorder=SimpleNamespace()),
        )

    environment = Environment(
        loader=FileSystemLoader(Path(__file__).parents[1] / "backend" / "workbench" / "templates"),
        autoescape=True,
    )
    view = captured["view"]
    assert result.terminal_status == "completed"
    assert captured["summary"] == {}
    safe_kwargs = captured["kwargs"]
    assert isinstance(safe_kwargs, dict)
    assert safe_kwargs["descriptive_stats"] == []
    assert safe_kwargs["statistical_tests"] == {}
    assert safe_kwargs["statistical_evidence"] == {}
    assert safe_kwargs["regression_table"] == {}
    assert safe_kwargs["include_legacy_presentation"] is False
    assert events["rendered"]
    html = environment.get_template("report.html.j2").render(view_model=view)
    serialized_view = json.dumps(view, sort_keys=True)
    assert view["status"] == "available"  # type: ignore[index]
    assert view["typedFacts"]  # type: ignore[index]
    assert all(value not in serialized_view for value in raw_values)
    assert all(value not in html for value in raw_values)
