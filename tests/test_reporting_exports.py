from pathlib import Path

import pytest

import workbench.orchestrator  # noqa: F401 - initialize the stage import seam first
from workbench.artifacts import read_json
from workbench.exports import export_pdf, export_xlsx
from workbench.narrative import build_claims
from workbench.projects import create_project, create_run
from workbench.reporting import render_html_report
from workbench.report_view_model import (
    build_regression_table,
    regression_table_export_rows,
)
from tests.test_v189_evidence_report_boundary import _produced_report


@pytest.mark.parametrize("format_name", ["html", "pdf", "xlsx"])
@pytest.mark.parametrize("fault", ["raw_path", "interrupted", "index_interrupted"])
def test_ordinary_export_fd_only_publication(tmp_path, monkeypatch, format_name, fault):
    import os
    import stat

    run = tmp_path / "project" / "runs" / "fd-export"
    with _produced_report(run, model_results=_ols_source(), diagnostic_summary={},
                          complete_collection=True) as case:
        writer = (case.owner.issue_report_evidence_writer(case.projection, sections=("regression_table",))
                  if format_name == "xlsx" else case.writer)
        export = {"html": render_html_report, "pdf": export_pdf, "xlsx": export_xlsx}[format_name]
        if fault == "index_interrupted":
            # A failed replacement must not leave a stale complete index row.
            export(writer, run)
        original_open, original_os_open, original_fsync = Path.open, os.open, os.fsync

        def guarded_open(path, mode="r", *args, **kwargs):
            if any(flag in mode for flag in ("w", "a", "+")):
                pytest.fail("export reopened a raw pathname for persistence")
            return original_open(path, mode, *args, **kwargs)

        def guarded_os_open(path, flags, *args, **kwargs):
            # openpyxl may stage renderer bytes in system temporary files;
            # only project persistence requires the admitted host descriptor.
            target = Path(path)
            if (flags & (os.O_WRONLY | os.O_RDWR | os.O_CREAT)
                    and (not target.is_absolute() or target.is_relative_to(run.parent.parent))):
                assert isinstance(kwargs.get("dir_fd"), int), "write requires an admitted directory FD"
            return original_os_open(path, flags, *args, **kwargs)

        def interrupted_fsync(fd):
            if stat.S_ISREG(os.fstat(fd).st_mode):
                publishing = any(
                    candidate.stat().st_ino == os.fstat(fd).st_ino
                    and any(row["artifact_id"] in {"report_html", "report_pdf", "tables_xlsx"}
                            for row in read_json(candidate)["artifacts"])
                    for candidate in run.glob(".artifacts_index.json.*.tmp")
                )
                if fault == "interrupted" or publishing:
                    raise KeyboardInterrupt("injected export interruption")
            return original_fsync(fd)

        if fault == "raw_path":
            monkeypatch.setattr(Path, "open", guarded_open)
            monkeypatch.setattr(os, "open", guarded_os_open)
            path = export(writer, run)
            assert path.read_bytes()
            import hashlib
            rows = read_json(run / "artifacts_index.json")["artifacts"]
            row = next(row for row in rows if row["path"] == path.relative_to(run).as_posix())
            assert row["sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
        else:
            monkeypatch.setattr(os, "fsync", interrupted_fsync)
            with pytest.raises(KeyboardInterrupt, match="injected export interruption"):
                export(writer, run)
            assert not list(run.rglob("*.tmp"))
            assert not any(row["artifact_id"] in {"report_html", "report_pdf", "tables_xlsx"}
                           for row in read_json(run / "artifacts_index.json")["artifacts"])


def test_ordinary_export_rejects_symlinks_traversal_and_foreign_run(tmp_path):
    from workbench.evidence.persistence import EvidencePersistenceError

    run = tmp_path / "project" / "runs" / "safe-export"
    outside = tmp_path / "outside"
    outside.mkdir()
    sentinel = outside / "sentinel"
    sentinel.write_bytes(b"do not replace")
    with _produced_report(run, model_results=_ols_source(), diagnostic_summary={},
                          complete_collection=True) as case:
        table_writer = case.owner.issue_report_evidence_writer(case.projection, sections=("regression_table",))
        for export, writer, directory in (
            (render_html_report, case.writer, "reports"),
            (export_pdf, case.writer, "reports"),
            (export_xlsx, table_writer, "exports"),
        ):
            with pytest.raises((ValueError, EvidencePersistenceError)):
                export(writer, run, filename="../escaped")
            with pytest.raises(ValueError, match="RUN_ROOT_MISMATCH"):
                export(writer, outside)
            parent = run / directory
            parent.mkdir(exist_ok=True)
            target = parent / "linked"
            target.symlink_to(sentinel)
            with pytest.raises((ValueError, EvidencePersistenceError)):
                export(writer, run, filename="linked")
            target.unlink()
            parent.rmdir()
            parent.symlink_to(outside, target_is_directory=True)
            with pytest.raises((ValueError, EvidencePersistenceError)):
                export(writer, run, filename="sentinel")
            parent.unlink()
        assert sentinel.read_bytes() == b"do not replace"
        assert list(outside.iterdir()) == [sentinel]
        assert not (run / "escaped").exists()
        assert not any(row["artifact_id"] in {"report_html", "report_pdf", "tables_xlsx"}
                       for row in read_json(run / "artifacts_index.json")["artifacts"])


def test_ordinary_export_withdraws_index_if_commit_sync_is_interrupted(tmp_path, monkeypatch):
    import os

    run = tmp_path / "project" / "runs" / "commit-interruption"
    with _produced_report(run, model_results=_ols_source(), diagnostic_summary={},
                          complete_collection=True) as case:
        original_fsync = os.fsync
        interrupted = False

        def interrupt_committed_index(fd):
            nonlocal interrupted
            if (not interrupted and os.fstat(fd).st_ino == run.stat().st_ino
                    and any(row["artifact_id"] == "report_html"
                            for row in read_json(run / "artifacts_index.json")["artifacts"])):
                interrupted = True
                raise KeyboardInterrupt("index directory sync interrupted")
            return original_fsync(fd)

        monkeypatch.setattr(os, "fsync", interrupt_committed_index)
        with pytest.raises(KeyboardInterrupt, match="index directory sync interrupted"):
            render_html_report(case.writer, run)
        assert interrupted
        assert not any(row["artifact_id"] == "report_html"
                       for row in read_json(run / "artifacts_index.json")["artifacts"])
        assert not list(run.rglob("*.tmp"))


@pytest.mark.parametrize("text", ["统计", "🧪"])
@pytest.mark.parametrize("location", [
    "title", "label", "section_id", "reason", "content_key",
    "legacy_title", "arma_title", "arma_value",
])
def test_pdf_font_rejects_unsupported_unicode_before_save(text, location, monkeypatch):
    from reportlab.pdfgen.canvas import Canvas
    from workbench.exports import _render_pdf_bytes
    from workbench.report_export import ReportExportError

    section = {"sectionId": "tests:a", "sectionType": "statistical_tests",
               "status": "available", "reasonCodes": [], "content": {"label": "ASCII", "p_value": 1.25e-24}}
    report = {"title": "Report", "resolvedSections": [section]}
    if location in {"legacy_title", "arma_title", "arma_value"}:
        report = {"title": text}
        if location != "legacy_title":
            report["report_kind"] = "arma_garch"
        if location == "arma_value":
            report.update(title="ASCII", overview=[{"field": "label", "value": "A" * 80 + text}])
    elif location == "title":
        report["title"] = text
    elif location == "label":
        section["content"]["label"] = text
    elif location == "section_id":
        section["sectionId"] = text
    elif location == "reason":
        section.update(status="unavailable", reasonCodes=[text], content=None)
    else:
        section["content"][text] = "ASCII"
    assert text in str(report), "the unsupported text must reach the renderer"
    saved = []
    save = Canvas.save

    def record_save(pdf):
        saved.append(True)
        return save(pdf)

    monkeypatch.setattr(Canvas, "save", record_save)
    with pytest.raises(ReportExportError, match="PDF_FONT_GLYPH_UNSUPPORTED"):
        _render_pdf_bytes(report, deterministic=True)
    assert not saved, "unsupported text must fail before PDF serialization"


def test_pdf_font_checks_actual_glyph_coverage(monkeypatch):
    from reportlab.pdfbase.pdfmetrics import getFont
    from workbench.exports import _render_pdf_bytes
    from workbench.report_export import ReportExportError

    glyphs = getFont("Helvetica-Bold").face.glyphWidths
    assert "Z" in glyphs
    monkeypatch.delitem(glyphs, "Z")
    assert "Z" not in glyphs
    with pytest.raises(ReportExportError, match="PDF_FONT_GLYPH_UNSUPPORTED"):
        _render_pdf_bytes({"title": "Z"}, deterministic=True)


def test_arma_pdf_preserves_supported_en_dash_in_drawn_text(monkeypatch):
    from reportlab.pdfbase.pdfmetrics import getFont
    from reportlab.pdfgen.canvas import Canvas
    from workbench.exports import _render_pdf_bytes
    from workbench.report_export import ReportExportError

    text = "A–B"
    for name in ("Helvetica", "Helvetica-Bold"):
        font = getFont(name)
        encoded = "–".encode(font.encName)
        assert encoded.decode(font.encName) == "–"
        assert font.encoding.vector[encoded[0]] == "endash"
        assert "endash" in font.face.glyphWidths
    report = {"report_kind": "arma_garch", "title": "ARMA report", "model_label": text,
              "overview": [{"field": text, "value": text}]}
    drawn = []
    draw_string = Canvas.drawString

    def record_draw(pdf, x, y, value, *args, **kwargs):
        drawn.append((value, pdf._fontname))
        return draw_string(pdf, x, y, value, *args, **kwargs)

    monkeypatch.setattr(Canvas, "drawString", record_draw)
    pdf = _render_pdf_bytes(report, deterministic=True)
    assert pdf.startswith(b"%PDF") and drawn
    assert [value for value, _ in drawn if value in (text, "A-B")] == [text] * 3
    assert (text, "Helvetica-Bold") in drawn and (text, "Helvetica") in drawn
    # Supported punctuation is preserved; missing glyphs still cannot fall back.
    with pytest.raises(ReportExportError, match="PDF_FONT_GLYPH_UNSUPPORTED"):
        _render_pdf_bytes({**report, "model_label": "A🧪B"}, deterministic=True)


def test_pdf_font_preserves_ascii_and_exact_tiny_p():
    from workbench.exports import _render_pdf_bytes

    report = {"title": "ASCII report", "resolvedSections": [{
        "sectionId": "tests:a", "sectionType": "statistical_tests", "status": "available",
        "reasonCodes": [], "content": {"label": "Final ASCII row", "p_value": 1.25e-24}}]}
    pdf = _render_pdf_bytes(report, deterministic=True)
    assert pdf.startswith(b"%PDF")
    assert b"Final ASCII row" in pdf and b"1.25e-24" in pdf


@pytest.mark.parametrize("text", ["统计", "🧪"])
def test_pdf_font_owner_export_rejects_without_output(tmp_path, text):
    from workbench.report_export import ReportExportError

    run = tmp_path / "project" / "runs" / "unsupported-font"
    with _produced_report(run, model_results=_ols_source(), diagnostic_summary={},
                          variable_labels={"x": text}, complete_collection=True) as case:
        assert case.writer.view_model["regression_table"]["rows"][0]["label"] == text
        with pytest.raises(ReportExportError, match="PDF_FONT_GLYPH_UNSUPPORTED"):
            export_pdf(case.writer, run)
        assert not (run / "reports" / "report.pdf").exists()
        assert not list(run.rglob("*.tmp"))
        assert all(row["artifact_id"] != "report_pdf"
                   for row in read_json(run / "artifacts_index.json")["artifacts"])


def test_structured_export_tables_are_section_specific_and_lossless():
    from workbench.report_export import _structured_section_rows
    section = {"sectionId": "statistical-tests", "sectionType": "statistical_tests",
               "status": "available", "reasonCodes": [],
               "content": {"results": [{"p_value": 1.25e-24, "label": "<untrusted>"}],
                           "warnings": ["declared warning"]}}
    rows = _structured_section_rows(section)
    assert rows[0]["p_value"] == 1.25e-24
    assert rows[0]["label"] == "<untrusted>"
    assert rows[0]["section_id"] == "statistical-tests"
    assert rows[0]["section_metadata"] == {"warnings": ["declared warning"]}
    blocked = {**section, "status": "unavailable", "content": None, "reasonCodes": ["MISSING"]}
    assert _structured_section_rows(blocked) == [{"section_id": "statistical-tests",
        "section_type": "statistical_tests", "status": "unavailable", "reason_codes": ["MISSING"]}]


def test_structured_html_and_pdf_render_all_partitions_without_raw_results():
    from workbench.exports import _render_pdf_bytes
    from workbench.report_view_model import _empty_legacy_presentation
    from jinja2 import Environment, FileSystemLoader
    report = _empty_legacy_presentation()
    report["resolvedSections"] = [{"sectionId": "tests:a", "sectionType": "statistical_tests",
        "status": "available", "reasonCodes": [],
        "content": {"results": [{"p_value": 1.25e-24, "label": "<script>bad</script>"}]}},
        {"sectionId": "tests:b", "sectionType": "statistical_tests", "status": "unavailable",
         "reasonCodes": ["SOURCE_MISSING"], "content": None}]
    import workbench.reporting as module
    template = Environment(loader=FileSystemLoader(Path(module.__file__).parent / "templates"),
                           autoescape=True).get_template("report.html.j2")
    html = template.render(view_model=report)
    assert "1.25e-24" in html and "SOURCE_MISSING" in html
    assert "&lt;script&gt;bad&lt;/script&gt;" in html and "<script>bad</script>" not in html
    pdf = _render_pdf_bytes(report, deterministic=True)
    assert b"1.25e-24" in pdf and b"SOURCE_MISSING" in pdf


@pytest.mark.parametrize("template_name", ["report.html.j2", "arma_garch_report.html.j2"])
def test_structured_html_appendix_preserves_complete_escaped_content(template_name):
    from jinja2 import Environment, FileSystemLoader
    from markupsafe import escape
    import workbench.reporting as module

    hostile = '<img src=x onerror="alert(1)">统计🧪 & 尾'
    reason = '<script>missing</script>来源不可用'
    report = {
        "title": "Template-only report", "metrics": {}, "chart_artifacts": [],
        "readCompleteness": "COMPLETE",
        "resolvedSections": [
            {"sectionId": "nested:统计", "sectionType": "statistical_tests",
             "status": "available", "reasonCodes": [],
             "content": {"nested": [{"p_value": 1.25e-24},
                                     {hostile: {"last_field": hostile}}]}},
            {"sectionId": "missing:last", "sectionType": "statistical_tests",
             "status": "unavailable", "reasonCodes": [reason], "content": None},
        ],
    }
    template = Environment(
        loader=FileSystemLoader(Path(module.__file__).parent / "templates"), autoescape=True,
    ).get_template(template_name)
    html = template.render(view_model=report)
    assert html.count('id="structured-evidence"') == 1
    appendix = html.split('id="structured-evidence"', 1)[1].split("</section>", 1)[0]
    assert "nested:统计" in appendix and "1.25e-24" in appendix
    assert "last_field" in appendix and appendix.count(str(escape(hostile))) == 2
    assert hostile not in html and reason not in html
    assert "missing:last" in appendix and "unavailable" in appendix
    assert str(escape(reason)) in appendix
    assert "Read completeness: COMPLETE" in appendix
    assert "Reading status only; not statistical validation or run completeness." in appendix
    if template_name == "arma_garch_report.html.j2":
        assert 'id="forecast-evaluation"' in html and 'id="selected-parameters"' in html


def test_generic_regression_mapping_retains_shared_section_value():
    from jinja2 import Environment, FileSystemLoader
    from markupsafe import escape
    import workbench.reporting as module

    label = "<b>回归🧪</b>"
    report = {"title": "Regression", "typed_sections": {
        "regression_table": {"rows": {label: {"p_value": 1.25e-24}}},
    }}
    html = Environment(
        loader=FileSystemLoader(Path(module.__file__).parent / "templates"), autoescape=True,
    ).get_template("report.html.j2").render(view_model=report)
    assert 'id="regression-table"' in html
    assert str(escape(label)) in html and label not in html and "1.25e-24" in html
    assert 'id="structured-evidence"' not in html


def test_exported_labels_and_headers_cannot_become_excel_formulas():
    import io
    from openpyxl import load_workbook
    from workbench.exports import _render_tables_xlsx
    formula = '=HYPERLINK("https://invalid.example", "label")'
    raw = _render_tables_xlsx({"section": [{"=1+1": formula, "p_value": 1.25e-24}]})
    workbook = load_workbook(io.BytesIO(raw), data_only=False)
    assert workbook.active["A1"].value == "=1+1"
    assert workbook.active["A1"].data_type == "s"
    assert workbook.active["A2"].value == formula
    assert workbook.active["A2"].data_type == "s"
    assert workbook.active["B2"].value == 1.25e-24
    workbook.close()


@pytest.mark.parametrize("cell_kind", ["nested", "text", "header"])
def test_xlsx_rejects_oversized_cells_before_assignment(cell_kind, monkeypatch):
    from openpyxl.worksheet.worksheet import Worksheet
    from workbench.exports import _render_tables_xlsx
    from workbench.report_export import ReportExportError, _structured_section_rows

    if cell_kind == "nested":
        # Every leaf fits the section text limit, but the flattened cell does not.
        section = {"sectionId": "tests:a", "sectionType": "statistical_tests",
                   "status": "available", "reasonCodes": [],
                   "content": {"results": [{"labels": ["W" * 4096] * 8}]}}
        rows = _structured_section_rows(section)
    elif cell_kind == "text":
        rows = [{"label": "W" * 32768}]
    else:
        rows = [{"W" * 32768: "label"}]
    append = Worksheet.append

    def checked_append(sheet, values):
        assert all(not isinstance(value, str) or len(value) <= 32767 for value in values), (
            "oversized cell reached openpyxl, which silently truncates it"
        )
        return append(sheet, values)

    monkeypatch.setattr(Worksheet, "append", checked_append)
    with pytest.raises(ReportExportError, match="XLSX_CELL_TEXT_LIMIT_EXCEEDED"):
        _render_tables_xlsx({"section": rows})


def test_xlsx_preserves_exact_cell_limit_and_nested_values():
    import io
    import json
    from openpyxl import load_workbook
    from workbench.exports import _render_tables_xlsx

    nested = ["W" * 4096] * 7 + ["W" * 4070]
    serialized = json.dumps(nested, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    assert len(serialized) == 32767
    text = "=" + "W" * 32766
    raw = _render_tables_xlsx({"section": [{"label": text, "nested": nested}]})
    workbook = load_workbook(io.BytesIO(raw), data_only=False)
    assert workbook.active["A2"].value == text
    assert workbook.active["A2"].data_type == "s"
    assert workbook.active["B2"].value == serialized
    assert json.loads(workbook.active["B2"].value) == nested
    workbook.close()


@pytest.mark.parametrize("missing_status", ["unavailable", "not_applicable"])
def test_regression_missing_rows_render_status_and_reason(missing_status):
    from jinja2 import Environment, FileSystemLoader
    import workbench.reporting as module
    from workbench.report_view_model import _empty_legacy_presentation

    template = Environment(loader=FileSystemLoader(Path(module.__file__).parent / "templates"),
                           autoescape=True).get_template("report.html.j2")
    report = _empty_legacy_presentation()
    report["typed_sections"] = {"regression_table": {
        "models": [{"id": "m", "label": "Admitted model"}],
        "rows": {"status": missing_status, "reason": "ROWS_NOT_DECLARED <untrusted>"},
        "notes": ["Retain this accepted note"],
    }}
    html = template.render(view_model=report)
    assert 'id="regression-table"' in html
    assert missing_status in html and "ROWS_NOT_DECLARED &lt;untrusted&gt;" in html
    assert "Admitted model" in html and "Retain this accepted note" in html
    assert "ROWS_NOT_DECLARED <untrusted>" not in html


@pytest.mark.parametrize("location", ["label", "section_id", "title"])
def test_structured_pdf_wraps_wide_text_to_font_metrics(location, monkeypatch):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen.canvas import Canvas
    from workbench.exports import _render_pdf_bytes

    wide = "W" * 80
    report = {"title": wide if location == "title" else "Report",
              "resolvedSections": [{"sectionId": wide if location == "section_id" else "tests:a",
                  "sectionType": "statistical_tests", "status": "available", "reasonCodes": [],
                  "content": {"label": wide if location == "label" else "ordinary label"}}]}
    drawn = []
    draw_string = Canvas.drawString

    def measured_draw(pdf, x, y, text, *args, **kwargs):
        drawn.append((x, y, text, pdf._fontname, pdf._fontsize))
        return draw_string(pdf, x, y, text, *args, **kwargs)

    monkeypatch.setattr(Canvas, "drawString", measured_draw)
    pdf = _render_pdf_bytes(report, deterministic=True)
    assert pdf.startswith(b"%PDF")
    assert drawn and sum(text.count("W") for _, _, text, _, _ in drawn) == 80
    for x, y, text, font, size in drawn:
        assert x + stringWidth(text, font, size) <= letter[0] - 72 + 1e-6, text
        assert 72 <= y <= letter[1] - 72


def test_structured_pdf_wrap_preserves_whitespace_and_paginates(monkeypatch):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen.canvas import Canvas
    from workbench.exports import _render_pdf_bytes

    label = ("W" * 80 + "  next-label\n") * 30
    report = {"resolvedSections": [{"sectionId": "tests:a", "status": "available",
                                   "content": {"label": label}}]}
    drawn = []
    draw_string = Canvas.drawString

    def measured_draw(pdf, x, y, text, *args, **kwargs):
        drawn.append((x, y, text, pdf._fontname, pdf._fontsize, pdf.getPageNumber()))
        return draw_string(pdf, x, y, text, *args, **kwargs)

    monkeypatch.setattr(Canvas, "drawString", measured_draw)
    _render_pdf_bytes(report, deterministic=True)
    body = [text[2:] for x, _, text, _, _, _ in drawn if x == 90]
    assert "".join(body) == "label: " + label.replace("\n", "")
    assert max(page for *_, page in drawn) >= 2
    for x, y, text, font, size, _ in drawn:
        assert x + stringWidth(text, font, size) <= letter[0] - 72 + 1e-6
        assert 72 <= y <= letter[1] - 72


def test_report_template_does_not_assume_interval_level_or_missingness():
    from jinja2 import Environment, FileSystemLoader
    import workbench.reporting as module
    from workbench.report_view_model import _empty_legacy_presentation
    template = Environment(loader=FileSystemLoader(Path(module.__file__).parent / "templates"),
                           autoescape=True).get_template("report.html.j2")
    report = _empty_legacy_presentation()
    report["typed_sections"] = {"regression_table": {"models": [{"id": "m", "label": "Model"}],
        "rows": [{"term": "x", "models": {"m": {"estimate": 1.0,
            "confidence_interval": [0.1, 2.0], "confidence_level": 0.9,
            "confidence_interval_method": "percentile_bootstrap"}}}]}}
    html = template.render(view_model=report)
    assert "95% CI" not in html
    assert "90%" in html and "percentile_bootstrap" in html
    for missing_rate in (None, {"status": "unavailable", "reason": "not measured"}):
        row = {"column": "x", "mean": 1.0}
        if missing_rate is not None:
            row["missing_rate"] = missing_rate
        report["typed_sections"]["descriptive_stats"] = [row]
        html = template.render(view_model=report)
        assert 'id="table-1"' in html and "Not declared" in html


def _ols_source(*, model_id="ols_1", label="OLS", nobs=20):
    """Explicit source for renderer fixtures, never recovered from display tables."""
    return {model_id: {
        "model_id": model_id, "model_type": "ols", "model_label": label, "nobs": nobs,
        "coefficients": {"x": {"estimate": 2.0, "std_error": 0.2, "p_value": 0.125,
            "source_id": f"model_results.{model_id}.coefficients.x"}},
    }}


def _xlsx(tables: dict, run_root: Path):
    from tests.test_v189_evidence_report_boundary import _owner_issued_evidence_writer

    writer = _owner_issued_evidence_writer(run_root.parent, tuple(tables), tables)
    return export_xlsx(writer, run_root)


def test_claims_have_source_ids():
    model_result = {
        "model_id": "regression_1",
        "coefficients": {
            "x": {
                "estimate": 2.0,
                "p_value": 0.01,
                "source_id": "model_results.regression_1.coefficients.x",
            }
        },
    }
    claims = build_claims([model_result], warnings=[])
    assert claims[0]["source_id"] == "model_results.regression_1.coefficients.x"


def test_claims_include_magnitude_significance_and_r_squared():
    model_result = {
        "model_id": "ols_1",
        "r_squared": 0.834,
        "coefficients": {
            "x": {
                "estimate": 2.5,
                "p_value": 0.04,
                "source_id": "model_results.ols_1.coefficients.x",
            }
        },
    }

    claims = build_claims([model_result], warnings=[])

    assert "one-unit increase in x" in claims[0]["claim"]
    assert "2.5000" in claims[0]["claim"]
    assert "5% level" in claims[0]["claim"]
    assert claims[1]["claim"] == (
        "Model ols_1 (R²) explains 83.4% of dependent-variable variation."
    )


def test_warning_claims_bind_to_severity_aware_sources():
    claims = build_claims([], warnings=[
        {"severity": "BLOCKER", "message": "Bad data"},
        {"severity": "WARNING", "message": "High correlation"},
        {"severity": "INFO", "message": "Dummy-coded categorical"},
        "Legacy issue",
    ])

    assert claims == [
        {"claim": "Bad data", "source_id": "errors.json", "confidence": 1.0},
        {"claim": "High correlation", "source_id": "warnings.json", "confidence": 1.0},
        {"claim": "Dummy-coded categorical", "source_id": "diagnostics.warnings", "confidence": 1.0},
        {"claim": "Legacy issue", "source_id": "errors.json", "confidence": 1.0},
    ]


def test_claims_skip_unavailable_estimates():
    model_result = {
        "model_id": "regression_1",
        "coefficients": {
            "x": {
                "estimate": None,
                "source_id": "model_results.regression_1.coefficients.x",
            },
            "z": {
                "estimate": "nan",
                "source_id": "model_results.regression_1.coefficients.z",
            },
            "w": {
                "estimate": float("inf"),
                "source_id": "model_results.regression_1.coefficients.w",
            },
        },
    }

    assert build_claims([model_result], warnings=[]) == []


def test_regression_table_projects_multiple_models_side_by_side_with_configurable_stars():
    table = build_regression_table(
        [
            (
                "ols_1",
                {
                    "model_id": "ols_1",
                    "model_label": "Baseline",
                    "nobs": 100,
                    "r_squared": 0.81,
                    "coefficients": {
                        "x": {
                            "estimate": 1.25,
                            "std_error": 0.2,
                            "confidence_interval": [0.8, 1.7],
                            "p_value": 0.15,
                            "source_id": "model_results.ols_1.coefficients.x",
                        },
                        "z": {
                            "estimate": -0.4,
                            "std_error": 0.1,
                            "p_value": 0.009,
                            "source_id": "model_results.ols_1.coefficients.z",
                        },
                    },
                },
            ),
            (
                "ols_2",
                {
                    "model_id": "ols_2",
                    "model_label": "With control",
                    "nobs": 95,
                    "r_squared": 0.84,
                    "coefficients": {
                        "x": {
                            "estimate": 1.1,
                            "std_error": 0.18,
                            "p_value": 0.04,
                            "source_id": "model_results.ols_2.coefficients.x",
                        },
                        "control": {
                            "estimate": 0.7,
                            "std_error": 0.3,
                            "p_value": 0.25,
                            "source_id": "model_results.ols_2.coefficients.control",
                        },
                    },
                },
            ),
        ],
        variable_labels={"x": "Treatment label", "z": "Baseline control"},
        significance_levels={"***": 0.01, "**": 0.05, "*": 0.20},
    )

    assert table["payload_schema"] == "workbench.regression-table"
    assert [(model["id"], model["label"]) for model in table["models"]] == [
        ("ols_1", "Baseline"),
        ("ols_2", "With control"),
    ]
    rows = {row["term"]: row for row in table["rows"]}
    assert set(rows) == {"x", "z", "control"}
    assert rows["x"]["label"] == "Treatment label"
    assert rows["x"]["models"]["ols_1"] == {
        "estimate": 1.25,
        "std_error": 0.2,
        "confidence_interval": [0.8, 1.7],
        "p_value": 0.15,
        "significance": "*",
        "source_id": "model_results.ols_1.coefficients.x",
    }
    assert rows["x"]["models"]["ols_2"]["significance"] == "**"
    assert rows["z"]["models"]["ols_1"]["significance"] == "***"
    assert rows["z"]["models"]["ols_2"] is None
    assert rows["control"]["models"]["ols_1"] is None
    assert rows["control"]["models"]["ols_2"]["source_id"] == (
        "model_results.ols_2.coefficients.control"
    )


def test_regression_table_packet_reaches_html_pdf_and_xlsx_exports(tmp_path: Path):
    project = create_project(tmp_path, "regression-table")
    run = create_run(project.root, mode="auto")
    source = {"ols_1": {
        "model_id": "ols_1", "model_type": "ols", "model_label": "Baseline",
        "coefficients": {"x": {
            "estimate": 1.25, "std_error": 0.2, "confidence_interval": [0.8, 1.7],
            "p_value": 0.04, "source_id": "model_results.ols_1.coefficients.x",
        }},
    }}
    packet = build_regression_table(
        list(source.items()),
        variable_labels={"x": "Treatment label"},
    )
    report = {
        "title": "Regression table report",
        "facts": [],
        "claims": [],
        "warnings": [],
        "regression_table": packet,
    }

    # Labels are producer inputs before sealing, not a post-issuance append.
    with _produced_report(run.root, model_results=source, diagnostic_summary={},
                          report_fields=report, variable_labels={"x": "Treatment label"}) as case:
        html_path = render_html_report(case.writer, run.root)
        pdf_path = export_pdf(case.writer, run.root)
        table_writer = case.owner.issue_report_evidence_writer(case.projection, sections=("regression_table",))
        xlsx_path = export_xlsx(table_writer, run.root)
        from openpyxl import load_workbook
        import json
        workbook = load_workbook(xlsx_path, data_only=False)
        try:
            rows = list(workbook.active.iter_rows(values_only=True))
            cell = json.loads(rows[1][rows[0].index("models")])["ols_1"]
            assert cell["estimate"] == source["ols_1"]["coefficients"]["x"]["estimate"]
            assert cell["p_value"] == source["ols_1"]["coefficients"]["x"]["p_value"]
        finally:
            workbook.close()
    from workbench.engine.stages.report import _xlsx_export_rows

    with pytest.raises(TypeError, match="caller-owned raw tables"):
        _xlsx(
            {"regression_table": _xlsx_export_rows(regression_table_export_rows(packet))},
            run.root,
        )

    html = html_path.read_text(encoding="utf-8")
    assert "Treatment label" in html
    assert "model_results.ols_1.coefficients.x" in html
    assert b"Regression table" in pdf_path.read_bytes()
    assert b"model_results.ols_1.coefficients.x" in pdf_path.read_bytes()


def test_survey_provenance_reaches_html_and_xlsx_without_recomputation(tmp_path: Path):
    """Report consumers must retain typed survey provenance and failure status."""
    project = create_project(tmp_path, "survey-provenance")
    run = create_run(project.root, mode="auto")
    result = {
        "model_id": "ols_1",
        "model_type": "ols",
        "model_label": "Survey OLS",
        "coefficients": {
            "x": {
                "estimate": 1.25,
                "std_error": 0.2,
                "ci_lower": 0.8,
                "ci_upper": 1.7,
                "p_value": 0.04,
                "variance_source": "survey_design",
                "source_id": "model_results.ols_1.coefficients.x",
                "warnings": ["SURVEY_REPLICATE_FAILURES", "LONELY_PSU_ADJUSTED"],
            }
        },
        "survey_design": {
            "variance_method": "replicate",
            "variance_source": "survey_design",
            "strata_column": "stratum",
            "psu_column": "psu",
            "degf": 4,
            "replicates_attempted": 12,
            "replicates_succeeded": 10,
            "replicates_failed": 2,
            "replicate_failure_reasons": {"ValueError": 2},
        },
    }
    packet = build_regression_table([("ols_1", result)])
    with _produced_report(run.root, model_results={"ols_1": result}, diagnostic_summary={}) as case:
        html = render_html_report(case.writer, run.root).read_text()
        table_writer = case.owner.issue_report_evidence_writer(case.projection, sections=("regression_table",))
        xlsx_path = export_xlsx(table_writer, run.root)
        from openpyxl import load_workbook
        import json
        workbook = load_workbook(xlsx_path, data_only=False)
        try:
            rows = list(workbook.active.iter_rows(values_only=True))
            cell = json.loads(rows[1][rows[0].index("models")])["ols_1"]
            for field in ("estimate", "std_error", "p_value", "variance_source", "source_id", "warnings"):
                assert cell[field] == result["coefficients"]["x"][field]
        finally:
            workbook.close()
    assert "survey_design" in html
    assert "SURVEY_REPLICATE_FAILURES" in html
    assert "replicates_attempted=12" in html
    assert "replicate_failure_reasons=ValueError:2" in html

    from workbench.engine.stages.report import _xlsx_export_rows

    xlsx_rows = _xlsx_export_rows(regression_table_export_rows(packet))
    assert "ols_1 variance_source" in xlsx_rows[0]
    assert "ols_1 warnings" in xlsx_rows[0]
    flattened = " ".join(str(value) for value in xlsx_rows[0].values())
    assert "survey_design" in flattened
    assert "SURVEY_REPLICATE_FAILURES" in flattened
    assert "replicates_attempted=12" in flattened
    assert "ValueError:2" in flattened


def test_pdf_and_xlsx_keep_variable_and_value_labels(tmp_path: Path):
    project = create_project(tmp_path, "labelled-exports")
    run = create_run(project.root, mode="auto")
    table_1 = [
        {
            "column": "group",
            "label": "Treatment group",
            "label_source": "declared",
            "value_labels": {"0": "Control", "1": "Treated"},
            "dtype": "int64",
            "count": 4,
            "missing": 0,
            "unique_count": 2,
            "mean": None,
            "std": None,
            "min": None,
            "max": None,
        }
    ]
    report = {
        "title": "Labelled export report",
        "facts": [],
        "claims": [],
        "descriptive_stats": table_1,
        "warnings": [],
    }

    with _produced_report(run.root, model_results=_ols_source(nobs=4), diagnostic_summary={},
                          report_fields=report) as case:
        pdf_path = export_pdf(case.writer, run.root)
        table_writer = case.owner.issue_report_evidence_writer(case.projection, sections=("descriptive_statistics",))
        xlsx_path = export_xlsx(table_writer, run.root)
        from openpyxl import load_workbook
        import json
        workbook = load_workbook(xlsx_path, data_only=False)
        try:
            rows = list(workbook.active.iter_rows(values_only=True))
            assert rows[1][rows[0].index("label")] == "Treatment group"
            assert json.loads(rows[1][rows[0].index("value_labels")]) == {"0": "Control", "1": "Treated"}
        finally:
            workbook.close()
    with pytest.raises(TypeError, match="caller-owned raw tables"):
        _xlsx({"table_1": table_1}, run.root)

    pdf_bytes = pdf_path.read_bytes()
    assert b"Treatment group" in pdf_bytes
    assert b"Control" in pdf_bytes
    assert b"Treated" in pdf_bytes


def test_render_and_export_reports(tmp_path: Path):
    project = create_project(tmp_path, "demo")
    run = create_run(project.root, mode="auto")
    report = {
        "title": "Demo Report",
        "facts": ["n=5"],
        "claims": [
            {
                "claim": "x is positive",
                "source_id": "model_results.regression_1.coefficients.x",
                "confidence": 0.9,
            }
        ],
        "warnings": [],
    }
    with _produced_report(run.root,
                          model_results=_ols_source(model_id="regression_1", label="Demo Report", nobs=5),
                          diagnostic_summary={}, report_fields=report) as case:
        html_path = render_html_report(case.writer, run.root)
        pdf_path = export_pdf(case.writer, run.root)
    with pytest.raises(TypeError, match="caller-owned raw tables"):
        _xlsx({"coefficients": [{"term": "x", "estimate": 2.0}]}, run.root)
    assert html_path.exists()
    assert pdf_path.exists()
    assert b"model_results.regression_1.coefficients.x" in pdf_path.read_bytes()
    artifacts = read_json(run.root / "artifacts_index.json")["artifacts"]
    artifact_types = {
        artifact["artifact_id"]: artifact["artifact_type"] for artifact in artifacts
    }
    artifact_inputs = {
        artifact["artifact_id"]: artifact["inputs"] for artifact in artifacts
    }
    assert artifact_types["report_html"] == "report"
    assert artifact_types["report_pdf"] == "report"
    assert artifact_inputs["report_pdf"] == []


def test_lmm_export_rows_are_excel_scalar_safe() -> None:
    """Versioned LMM coefficients include CI arrays that XLSX cannot store raw."""

    from workbench.engine.stages.report import _xlsx_export_rows

    rows = _xlsx_export_rows([
        {
            "model_id": "linear_mixed_effects_1",
            "term": "group_time_interaction",
            "estimate": 0.94,
            "confidence_interval": [0.88, 1.00],
        }
    ])

    assert rows == [
        {
            "model_id": "linear_mixed_effects_1",
            "term": "group_time_interaction",
            "estimate": 0.94,
            "confidence_interval": "[0.88,1.0]",
        }
    ]


def test_render_html_report_escapes_untrusted_content(tmp_path: Path):
    project = create_project(tmp_path, "demo")
    run = create_run(project.root, mode="auto")

    message = "<b>unsafe</b>"
    summary = {"diagnostics": {
        "blockers": [], "warnings": [{"severity": "WARNING", "code": "EXPORT_ESCAPE",
            "message": message}], "cautions": [], "info": [],
    }}
    with _produced_report(run.root,
                          model_results=_ols_source(model_id="regression_1", label="<script>alert(1)</script>"),
                          diagnostic_summary=summary) as case:
        diagnostics = next(section for section in case.view["resolvedSections"]
                           if section["sectionType"] == "diagnostics")
        assert diagnostics["status"] == "available"
        assert diagnostics["content"]["warnings"][0]["message"] == message
        html_path = render_html_report(case.writer, run.root)

    html = html_path.read_text(encoding="utf-8")
    assert "<script>" not in html
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html
    assert "&lt;b&gt;unsafe&lt;/b&gt;" in html
    assert "<b>unsafe</b>" not in html


def test_report_renders_statistical_tests_section(tmp_path: Path):
    project = create_project(tmp_path, "demo")
    run = create_run(project.root, mode="auto")
    report = {
        "title": "Demo Report",
        "facts": [],
        "claims": [],
        "variable_importance": [
            {
                "variable": "x",
                "correlation": 0.98,
                "best_p_value": 0.001,
                "test_type": "correlation",
            }
        ],
        "statistical_evidence": {
            "results": [
                {
                    "label": "Pearson correlation: y vs x",
                    "statistic": 0.98,
                    "p_value": 0.001,
                    "interpretation": "Statistic 0.9800; p = 0.001.",
                    "source_id": "statistical_tests.correlations.y.x",
                }
            ],
        },
        "warnings": [],
    }

    with _produced_report(run.root, model_results=_ols_source(label="Demo Report"),
                          diagnostic_summary={}, report_fields=report) as case:
        html_path = render_html_report(case.writer, run.root)
        pdf_path = export_pdf(case.writer, run.root)

    html = html_path.read_text(encoding="utf-8")
    assert "<h2>Statistical Tests</h2>" in html
    assert "Pearson correlation: y vs x" in html
    assert b"Statistical tests" in pdf_path.read_bytes()


def test_report_renders_typed_statistical_evidence_and_preserves_table_fields(tmp_path: Path):
    project = create_project(tmp_path, "evidence-report")
    run = create_run(project.root, mode="auto")
    report = {
        "title": "Evidence Report",
        "facts": [],
        "claims": [],
        "descriptive_stats": [{"column": "y", "dtype": "float64", "count": 9, "missing": 0, "missing_rate": 0.0, "mean": 2.0, "std": 0.8}],
        "statistical_tests": {"y_related": [], "other": [], "other_truncated": 0},
        "statistical_evidence": {
            "payload_schema": "workbench.statistics.evidence-packet",
            "schema_version": 1,
            "correction_scope": "advanced_evidence_family",
            "results": [
                {
                    "test_type": "anova_posthoc",
                    "test_id": "anova_posthoc:y:region",
                    "nobs": 9,
                    "statistic": 12.0,
                    "p_value": 0.01,
                    "p_value_corrected": 0.02,
                    "effect_size": {"effect_size_name": "eta_squared", "value": 0.7},
                    "assumptions": ["independent observations"],
                    "warnings": [],
                }
            ],
        },
        "warnings": [],
    }

    with _produced_report(run.root, model_results=_ols_source(label="Evidence Report", nobs=9),
                          diagnostic_summary={}, report_fields=report) as case:
        html_path = render_html_report(case.writer, run.root)
    html = html_path.read_text(encoding="utf-8")

    assert "Statistical Evidence" in html
    assert "anova_posthoc" in html
    assert "eta_squared" in html
    assert "advanced_evidence_family" in html


def test_report_formats_tiny_variable_importance_p_values(tmp_path: Path):
    project = create_project(tmp_path, "demo")
    run = create_run(project.root, mode="auto")
    tiny_p = 1.25e-24
    report = {"variable_importance": [{
        "variable": "x", "correlation": 0.8,
        "best_p_value": tiny_p, "test_type": "correlation",
    }]}
    with _produced_report(run.root, model_results=_ols_source(label="Demo Report"),
                          diagnostic_summary={}, report_fields=report) as case:
        importance = next(section for section in case.view["resolvedSections"]
                          if section["sectionType"] == "variable_importance")
        assert importance["status"] == "available"
        assert importance["content"]["rows"][0]["best_p_value"] == tiny_p
        assert importance["content"]["rows"][0]["best_p_value"] != 0
        facts = [fact for fact in case.projection.wire["facts"]
                 if fact["statisticId"] == "best_p_value"]
        assert len(facts) == 1
        assert facts[0]["value"] == tiny_p
        assert facts[0]["value"] != 0
        html_path = render_html_report(case.writer, run.root)

    html = html_path.read_text(encoding="utf-8")
    assert "1.25e-24" in html
    assert "significant at the 1% level" not in html


def test_report_formats_tiny_diagnostic_p_values(tmp_path: Path):
    project = create_project(tmp_path, "demo")
    run = create_run(project.root, mode="auto")
    report = {
        "title": "Demo Report",
        "facts": [],
        "claims": [],
        "warnings": [],
        "diagnostics": {
            "ols_1": {
                "breusch_pagan": {"lm": 25.0, "p_value": 0.0000002},
                "jarque_bera": {"statistic": 40.0, "p_value": 0.0000003},
            }
        },
    }

    summary = {"diagnostics": {"blockers": [], "warnings": [], "cautions": [], "info": []}}
    with _produced_report(run.root, model_results=_ols_source(label="Demo Report"),
                          diagnostic_summary=summary, report_fields=report) as case:
        html_path = render_html_report(case.writer, run.root)

    html = html_path.read_text(encoding="utf-8")
    assert "p &lt; 0.001" in html
    assert "p = 0.0000" not in html


@pytest.mark.parametrize("value", [True, False, 0])
def test_report_diagnostic_booleans_are_not_numbers(value):
    from jinja2 import Environment, FileSystemLoader
    import workbench.reporting as module

    # Exercise the renderer independently: upstream VIF admission already
    # rejects booleans, but rendering must never coerce them into statistics.
    variable = '<script>alert("vif")</script>'
    diagnostic = {
        "breusch_pagan": {"lm": value, "p_value": value},
        "durbin_watson": value,
        "vif": {variable: value},
    }
    template = Environment(
        loader=FileSystemLoader(Path(module.__file__).parent / "templates"), autoescape=True,
    ).get_template("report.html.j2")
    html = template.render(view_model={"typed_sections": {"model_diagnostics": {"ols_1": diagnostic}}})
    section = html.split('<section id="regression-diagnostics">', 1)[1].split("</section>", 1)[0]
    assert variable not in section
    assert "&lt;script&gt;" in section
    if type(value) is bool:
        assert section.count("unavailable (invalid numeric value: boolean)") == 4
        assert "LM = unavailable (invalid numeric value: boolean), p unavailable" in section
        assert "Durbin-Watson: unavailable (invalid numeric value: boolean)" in section
        assert "p &lt;" not in section and "p =" not in section
        assert "high" not in section and "moderate" not in section
    else:
        assert "LM = 0.00, p &lt; 0.001" in section
        assert "Durbin-Watson: 0.000" in section
        assert "0.0" in section.split("<tbody>", 1)[1].split("</tbody>", 1)[0]
        assert "unavailable" not in section


@pytest.mark.parametrize("missing_kind", ["missing", "null"])
def test_report_renders_missing_diagnostic_p_value(tmp_path: Path, missing_kind: str):
    project = create_project(tmp_path, "missing-diagnostic-p")
    run = create_run(project.root, mode="auto")
    bp = {"lm": 25.0}
    if missing_kind == "null":
        bp["p_value"] = None
    summary = {"diagnostics": {"blockers": [], "warnings": [], "cautions": [], "info": []}}
    with _produced_report(
        run.root, model_results=_ols_source(), diagnostic_summary=summary,
        report_fields={"diagnostics": {"ols_1": {"breusch_pagan": bp}}},
    ) as case:
        admitted = case.view["typed_sections"]["model_diagnostics"]["ols_1"]["breusch_pagan"]
        assert admitted["lm"] == 25.0
        if missing_kind == "missing":
            assert "p_value" not in admitted
        else:
            assert admitted["p_value"] == {
                "status": "unavailable", "reason": "PRODUCER_FIELD_UNDECLARED",
            }
        html = render_html_report(case.writer, run.root).read_text(encoding="utf-8")
    diagnostics = html.split('<section id="regression-diagnostics">', 1)[1].split("</section>", 1)[0]
    assert "Breusch-Pagan (heteroskedasticity): LM = 25.00, p unavailable (not declared)" in diagnostics
    assert "p &lt;" not in diagnostics
    assert "p =" not in diagnostics


@pytest.mark.parametrize("field,kind", [
    (field, kind) for field in ("lm", "durbin_watson", "vif")
    for kind in ("missing", "null", "numeric")
] + [("vif", "null_container")])
def test_report_renders_missing_diagnostic_numbers(tmp_path: Path, field: str, kind: str):
    project = create_project(tmp_path, "missing-diagnostic-number")
    run = create_run(project.root, mode="auto")
    diagnostic = {
        "breusch_pagan": {"lm": 25.0, "p_value": 0.04},
        "durbin_watson": 1.875,
        "vif": {"x": 12.5},
    }
    target, key = (diagnostic["breusch_pagan"], "lm") if field == "lm" else (
        (diagnostic["vif"], "x") if field == "vif" else (diagnostic, field))
    if kind == "missing":
        target.pop(key)
    elif kind == "null":
        target[key] = None
    elif kind == "null_container":
        diagnostic["vif"] = None
    summary = {"diagnostics": {"blockers": [], "warnings": [], "cautions": [], "info": []}}
    with _produced_report(
        run.root, model_results=_ols_source(), diagnostic_summary=summary,
        report_fields={"diagnostics": {"ols_1": diagnostic}},
    ) as case:
        admitted = case.view["typed_sections"]["model_diagnostics"]["ols_1"]
        admitted_target = admitted["breusch_pagan"] if field == "lm" else (
            admitted["vif"] if field == "vif" and kind != "null_container" else admitted)
        if kind == "missing":
            assert key not in admitted_target
        elif kind in ("null", "null_container"):
            assert admitted_target["vif" if kind == "null_container" else key] == {
                "status": "unavailable", "reason": "PRODUCER_FIELD_UNDECLARED",
            }
        else:
            assert admitted_target[key] == target[key]
        html = render_html_report(case.writer, run.root).read_text(encoding="utf-8")
    section = html.split('<section id="regression-diagnostics">', 1)[1].split("</section>", 1)[0]
    assert "p = 0.0400" in section
    expected = {
        "lm": "LM = 25.00" if kind == "numeric" else "LM = unavailable (not declared)",
        "durbin_watson": "Durbin-Watson: 1.875" if kind == "numeric" else "Durbin-Watson: unavailable (not declared)",
        "vif": "12.5" if kind == "numeric" else "unavailable (not declared)",
    }
    assert expected[field] in section
    if field != "lm":
        assert "LM = 25.00" in section
    if field != "durbin_watson":
        assert "Durbin-Watson: 1.875" in section
    if field != "vif" or kind == "numeric":
        assert "12.5" in section and "high" in section
    else:
        assert "high" not in section and "moderate" not in section
        if kind == "null":
            assert "<td>x</td>" in section


def test_report_renders_descriptive_statistics_section(tmp_path: Path):
    project = create_project(tmp_path, "demo")
    run = create_run(project.root, mode="auto")
    report = {
        "title": "Demo Report",
        "facts": [],
        "claims": [],
        "descriptive_stats": [
            {
                "column": "y",
                "dtype": "float64",
                "count": 35,
                "missing": 0,
                "missing_rate": 0.0,
                "unique_count": 35,
                "mean": 10.5,
                "std": 3.2,
                "min": 1.0,
                "max": 20.0,
            },
            {
                "column": "x",
                "dtype": "float64",
                "count": 35,
                "missing": 0,
                "missing_rate": 0.0,
                "unique_count": 35,
                "mean": 5.25,
                "std": 2.1,
                "min": 0.0,
                "max": 10.0,
            },
        ],
        "statistical_tests": [],
        "warnings": [],
    }

    with _produced_report(run.root, model_results=_ols_source(label="Demo Report", nobs=35),
                          diagnostic_summary={}, report_fields=report) as case:
        html_path = render_html_report(case.writer, run.root)
        pdf_path = export_pdf(case.writer, run.root)

    html = html_path.read_text(encoding="utf-8")
    assert "<h2>Descriptive Statistics</h2>" in html
    assert "y" in html
    assert "float64" in html
    assert "10.5" in html
    assert "3.2" in html
    assert b"Descriptive statistics" in pdf_path.read_bytes()
