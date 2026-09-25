from __future__ import annotations

from collections.abc import Mapping, Sequence
import io
import json
import os
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase.pdfmetrics import getFont, stringWidth
from reportlab.pdfgen import canvas

from .report_export import ReportEvidenceWriterV1, ReportExportError
from .reporting import (
    ReportDocumentWriterV1,
    WorkflowReportWriterV1,
    _write_report_output,
)


def _require_pdf_glyphs(text: str, font_name: str) -> None:
    """Reject missing glyphs before ReportLab can silently substitute squares.

    The renderer currently uses Type 1 fonts. This is a coverage guard, not
    Unicode/font-fallback support; no system fonts or new resources are loaded.
    """
    font = getFont(font_name)
    try:
        encoded = text.encode(font.encName, errors="strict")
    except UnicodeEncodeError as exc:
        raise ReportExportError(
            f"PDF_FONT_GLYPH_UNSUPPORTED: {font_name} lacks U+{ord(text[exc.start]):04X}"
        ) from exc
    if (encoded.decode(font.encName, errors="strict") != text
            or any(font.encoding.vector[code] not in font.face.glyphWidths for code in encoded)):
        raise ReportExportError(
            f"PDF_FONT_GLYPH_UNSUPPORTED: {font_name} cannot render text without substitution"
        )


def _draw_pdf_text(pdf: canvas.Canvas, x: float, y: float, text: str) -> None:
    _require_pdf_glyphs(text, pdf._fontname)
    pdf.drawString(x, y, text)


def _render_pdf_bytes(
    report: Mapping[str, Any],
    *,
    deterministic: bool = False,
) -> bytes:
    output = io.BytesIO()
    canvas_options = {"invariant": 1} if deterministic else {}
    pdf = canvas.Canvas(output, pagesize=letter, pageCompression=0, **canvas_options)
    _, height = letter
    y = height - 72
    if report.get("resolvedSections"):
        y = _draw_section(pdf, str(report.get("title", "Report")), [], y)
        for section in report["resolvedSections"]:
            lines = _structured_pdf_lines(section["content"]) if section["status"] == "available" else list(section["reasonCodes"])
            label = str(section.get("sectionType", "result")).replace("_", " ").capitalize()
            y = _draw_section(pdf, f"{label}: {section['sectionId']} ({section['status']})", lines, y)
    elif report.get("report_kind") == "arma_garch":
        _draw_arma_garch_pdf(pdf, report, y)
    else:
        pdf.setFont("Helvetica-Bold", 14)
        _draw_pdf_text(pdf, 72, y, str(report.get("title", "Report")))
        y -= 32
        y = _draw_section(pdf, "Facts", report.get("presentation_facts", report.get("facts", [])), y)
        y = _draw_section(pdf, "Descriptive statistics", report.get("descriptive_stats", []), y)
        regression_table = report.get("regression_table")
        if isinstance(regression_table, Mapping):
            y = _draw_section(pdf, "Regression table", _regression_pdf_items(regression_table), y)
        y = _draw_section(pdf, "Interpretation", report.get("claims", []), y)
        st = report.get("statistical_tests")
        if isinstance(st, Mapping):
            y = _draw_section(pdf, "Statistical tests (outcome-related)", st.get("y_related", []), y)
            other = st.get("other", [])
            if other:
                y = _draw_section(pdf, f"Other tests (+{st.get('other_truncated', 0)} truncated)", other, y)
        else:
            y = _draw_section(pdf, "Statistical tests", st or [], y)
        _draw_section(pdf, "Warnings", report.get("warnings", []), y)
    pdf.save()
    return output.getvalue()


def _structured_pdf_lines(value: Any, prefix: str = "") -> list[str]:
    """Walk all admitted fields without computing or truncating statistics."""
    if isinstance(value, Mapping) and value:
        return [line for key, item in value.items()
                for line in _structured_pdf_lines(item, f"{prefix}.{key}" if prefix else str(key))]
    if isinstance(value, (list, tuple)) and value:
        return [line for index, item in enumerate(value)
                for line in _structured_pdf_lines(item, f"{prefix}[{index}]")]
    return [f"{prefix}: {value}"]


def export_pdf(
    writer: ReportDocumentWriterV1 | WorkflowReportWriterV1,
    run_root: Path,
    *,
    filename: str = "report.pdf",
    artifact_id: str = "report_pdf",
    inputs: Sequence[str] = (),
    evidence_port: Any = None,
    admission_receipt: Any = None,
) -> Path:
    if type(writer) is WorkflowReportWriterV1:
        if writer.canonical_run_root != run_root.expanduser().resolve(strict=False):
            from .agent.workflow_runtime import WorkflowExecutionError

            raise WorkflowExecutionError(
                "WORKFLOW_ARTIFACT_PATH_CHANGED: workflow report run root changed"
            )
        return writer._write_formal_output(
            "pdf",
            directory_name="reports",
            filename=filename,
            artifact_id=artifact_id,
            artifact_type="report",
            content=_render_pdf_bytes(writer.view_model, deterministic=True),
        )
    if type(writer) is not ReportDocumentWriterV1:
        raise TypeError("export_pdf requires ReportDocumentWriterV1")
    writer._assert_server_owned()
    if writer.canonical_run_root != run_root.expanduser().resolve(strict=False):
        raise ValueError("REPORT_WRITER_RUN_ROOT_MISMATCH")
    report = writer.view_model
    output = _render_pdf_bytes(report)
    return _write_report_output(
        writer, run_root, directory="reports", filename=filename,
        artifact_id=artifact_id, artifact_type="report", step="export", inputs=inputs,
        content=output,
    )


def export_xlsx(
    writer: ReportEvidenceWriterV1 | WorkflowReportWriterV1,
    run_root: Path,
    *,
    filename: str = "tables.xlsx",
    artifact_id: str = "tables_xlsx",
    inputs: Sequence[str] = (),
    evidence_port: Any = None,
    admission_receipt: Any = None,
) -> Path:
    if type(writer) is WorkflowReportWriterV1:
        if writer.canonical_run_root != run_root.expanduser().resolve(strict=False):
            from .agent.workflow_runtime import WorkflowExecutionError

            raise WorkflowExecutionError(
                "WORKFLOW_ARTIFACT_PATH_CHANGED: workflow report run root changed"
            )
        return writer._write_formal_output(
            "xlsx",
            directory_name="exports",
            filename=filename,
            artifact_id=artifact_id,
            artifact_type="table_export",
            content=_render_workflow_xlsx(writer),
        )
    xlsx_bytes = render_typed_xlsx(writer)
    return _write_report_output(
        writer, run_root, directory="exports", filename=filename,
        artifact_id=artifact_id, artifact_type="table_export", step="export", inputs=inputs,
        content=xlsx_bytes,
    )


def render_typed_xlsx(
    writer: ReportEvidenceWriterV1,
) -> bytes:
    """Render a workbook only from the server-owned typed writer capability."""

    if type(writer) is not ReportEvidenceWriterV1:
        raise TypeError("render_typed_xlsx requires ReportEvidenceWriterV1")
    writer._assert_server_owned()
    tables = writer.tables()

    return _render_tables_xlsx(tables)


def _render_workflow_xlsx(writer: WorkflowReportWriterV1) -> bytes:
    writer._assert_server_owned()
    exploration = writer.view_model.get("exploration", {})
    steps = exploration.get("steps", {}) if isinstance(exploration, Mapping) else {}
    rows = [
        {
            "step_id": step_id,
            "status": value.get("status", ""),
            "artifact_ids": ", ".join(value.get("artifact_ids", [])),
        }
        for step_id, value in steps.items()
        if isinstance(value, Mapping)
    ]
    return _render_tables_xlsx({"workflow_steps": rows})


def _render_tables_xlsx(tables: Mapping[str, Sequence[Mapping[str, Any]]]) -> bytes:
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    for name, rows in tables.items():
        worksheet = workbook.create_sheet(title=_sheet_title(name))
        normalized_rows = list(rows)
        headers = _headers(normalized_rows)
        worksheet.append([_excel_value(header) for header in headers])
        worksheet.freeze_panes = "A2"
        for cell in worksheet[1]:
            cell.data_type = "s"
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for row in normalized_rows:
            values = [_excel_value(row.get(header, "")) for header in headers]
            worksheet.append(values)
            for index, value in enumerate(values, start=1):
                if isinstance(value, str):
                    # Labels are data, never spreadsheet formulas. Preserve the
                    # exact text rather than prefixing an extra apostrophe.
                    worksheet.cell(worksheet.max_row, index).data_type = "s"
        worksheet.auto_filter.ref = worksheet.dimensions
        for index, header in enumerate(headers, start=1):
            values = [str(header)] + [str(row.get(header, "")) for row in normalized_rows[:200]]
            worksheet.column_dimensions[get_column_letter(index)].width = min(48, max(12, max(map(len, values)) + 2))
    if not workbook.worksheets:
        workbook.create_sheet(title="tables")
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _wrap_pdf_text(text: str, font: str, size: float, width: float) -> list[str]:
    """Fit actual glyph advances; preserve spaces and hard line breaks."""
    # Newlines are layout, not glyphs; every other character must be renderable
    # before measuring widths (which can otherwise measure a substituted glyph).
    _require_pdf_glyphs(text.replace("\n", ""), font)
    lines = []
    for paragraph in text.split("\n"):
        while stringWidth(paragraph, font, size) > width:
            low, high = 0, len(paragraph)
            while low < high:
                middle = (low + high + 1) // 2
                if stringWidth(paragraph[:middle], font, size) <= width:
                    low = middle
                else:
                    high = middle - 1
            if low == 0:
                raise ReportExportError("PDF_GLYPH_EXCEEDS_LINE_WIDTH")
            # Prefer a word boundary, but split long unbroken labels as needed.
            boundary = max((i + 1 for i, char in enumerate(paragraph[:low]) if char.isspace()), default=0)
            cut = boundary or low
            lines.append(paragraph[:cut])
            paragraph = paragraph[cut:]
        lines.append(paragraph)
    return lines


def _draw_section(pdf: canvas.Canvas, title: str, items: Any, y: float) -> float:
    def draw(text: str, x: float, font: str, size: float, leading: float) -> None:
        nonlocal y
        if y < 72:
            pdf.showPage()
            y = letter[1] - 72
        pdf.setFont(font, size)
        _draw_pdf_text(pdf, x, y, text)
        y -= leading

    for line in _wrap_pdf_text(title, "Helvetica-Bold", 12, letter[0] - 144):
        draw(line, 72, "Helvetica-Bold", 12, 18)
    body_width = letter[0] - 72 - 90 - stringWidth("- ", "Helvetica", 10)
    for item in items:
        for index, line in enumerate(_wrap_pdf_text(_report_item_text(item), "Helvetica", 10, body_width)):
            draw(("- " if index == 0 else "  ") + line, 90, "Helvetica", 10, 14)
    return y - 10


def _draw_arma_garch_pdf(pdf: canvas.Canvas, report: Mapping[str, Any], y: float) -> None:
    """Render the native time-series evidence in a paginated, searchable PDF."""
    _, height = letter

    def heading(text: str, position: float, level: int = 1) -> float:
        if position < 86:
            pdf.showPage()
            position = height - 72
        pdf.setFont("Helvetica-Bold", 14 if level == 1 else 11)
        _draw_pdf_text(pdf, 54, position, text)
        return position - (24 if level == 1 else 18)

    def line(label: str, value: object, position: float) -> float:
        if position < 72:
            pdf.showPage()
            position = height - 72
        pdf.setFont("Helvetica-Bold", 9)
        _draw_pdf_text(pdf, 64, position, _pdf_text(label, 27, font_name="Helvetica-Bold"))
        pdf.setFont("Helvetica", 9)
        _draw_pdf_text(pdf, 210, position, _pdf_text(value, 70))
        return position - 13

    y = heading(str(report.get("title", "ARMA-GARCH Volatility Report")), y)
    pdf.setFont("Helvetica", 10)
    _draw_pdf_text(pdf, 54, y, _pdf_text(report.get("model_label", ""), 100))
    y -= 24
    y = heading("Run and model contract", y, level=2)
    for row in report.get("overview", []):
        if isinstance(row, Mapping):
            y = line(str(row.get("field", "")), row.get("value", ""), y)
    y = heading("Rolling forecast evaluation", y, level=2)
    for key in ("validation_n", "successful_forecast_n", "mae", "rmse", "interval_coverage", "average_interval_width"):
        y = line(key.replace("_", " ").title(), report.get("metrics", {}).get(key, "—") if isinstance(report.get("metrics"), Mapping) else "—", y)
    y = heading("Selected parameters", y, level=2)
    for row in report.get("parameters", []):
        if isinstance(row, Mapping):
            y = line(f"{row.get('component', '')}: {row.get('parameter', '')}", row.get("estimate", ""), y)
    y = heading("Production next-observation forecast", y, level=2)
    for row in report.get("next_forecast", []):
        if isinstance(row, Mapping):
            y = line(str(row.get("field", "")), row.get("value", ""), y)
    y = heading("Residual diagnostics", y, level=2)
    for row in report.get("diagnostics", []):
        if not isinstance(row, Mapping):
            continue
        statistic, p_value = row.get("statistic", "—"), row.get("p_value", "—")
        detail = row.get("detail", "")
        summary = f"stat={statistic} p={p_value} [{row.get('status', '—')}]"
        y = line(str(row.get("test", "")), f"{summary} {detail}".strip(), y)
    y = heading("Interpretation limits", y, level=2)
    for item in report.get("limitations", []):
        y = line("Limit", item, y)
    y = heading("Chart evidence", y, level=2)
    chart_artifacts = report.get("chart_artifacts", [])
    y = line("Registered chart artifacts", len(chart_artifacts), y)
    # Every registered chart is listed: a fixed cap here silently disagreed with
    # the count printed directly above it as soon as the pack grew a chart.
    for artifact_id in chart_artifacts:
        y = line("artifact", artifact_id, y)


def _pdf_text(value: object, length: int, *, font_name: str = "Helvetica") -> str:
    text = str(value).replace("\n", " ")
    # Check before the existing legacy length formatting, so an unsupported
    # character beyond its cutoff cannot silently disappear into an ellipsis.
    _require_pdf_glyphs(text, font_name)
    return text if len(text) <= length else f"{text[: length - 1]}…"


def _report_item_text(item: Any) -> str:
    if not isinstance(item, Mapping):
        return str(item)
    if "column" in item and "dtype" in item:
        raw_column = str(item.get("column", ""))
        display_column = str(item.get("label") or raw_column)
        column_text = (
            display_column
            if display_column == raw_column
            else f"{display_column} ({raw_column})"
        )
        value_labels = _value_labels_text(item.get("value_labels"))
        dtype = item.get("dtype", "")
        missing = item.get("missing", 0)
        unique = item.get("unique_count", 0)
        if item.get("mean") is not None:
            return (
                f"{column_text} ({dtype}): "
                f"mean={item['mean']:.4f} std={item['std']:.4f} "
                f"min={item['min']:.4f} max={item['max']:.4f} "
                f"missing={missing} unique={unique}{value_labels}"
            )
        return (
            f"{column_text} ({dtype}): "
            f"missing={missing} unique={unique}{value_labels}"
        )
    if item.get("label") and item.get("interpretation"):
        text = f"{item['label']}: {item['interpretation']}"
    else:
        text = str(item.get("claim", item.get("message", "")))
    source_id = item.get("source_id")
    if source_id:
        text = f"{text} [source: {source_id}]"
    confidence = item.get("confidence")
    if confidence is not None:
        text = f"{text} [confidence: {confidence}]"
    return text


def _value_labels_text(value_labels: object) -> str:
    if not isinstance(value_labels, Mapping) or not value_labels:
        return ""
    rendered = ", ".join(
        f"{value}={label}" for value, label in value_labels.items()
    )
    return f" values: {rendered}"


def _excel_value(value: Any) -> Any:
    """Convert structured metadata into a lossless, Excel-safe scalar."""
    if isinstance(value, (Mapping, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    # openpyxl silently slices strings during cell assignment. Fail before that
    # boundary, including for serialized nested metadata and column headers.
    if isinstance(value, str) and len(value) > 32767:
        raise ReportExportError("XLSX_CELL_TEXT_LIMIT_EXCEEDED: cell text exceeds 32767 characters")
    return value


def _regression_pdf_items(table: Mapping[str, Any]) -> list[str]:
    models = [model for model in table.get("models", []) if isinstance(model, Mapping)]
    items: list[str] = []
    for row in table.get("rows", []):
        if not isinstance(row, Mapping):
            continue
        parts = [f"{row.get('label', row.get('term', ''))} ({row.get('term', '')})"]
        cells = row.get("models", {})
        for model in models:
            model_id = str(model.get("id", ""))
            cell = cells.get(model_id) if isinstance(cells, Mapping) else None
            if not isinstance(cell, Mapping):
                parts.append(f"{model.get('label', model_id)}: unavailable")
                continue
            value = (
                f"{model.get('label', model_id)}: estimate={cell.get('estimate', '—')}"
                f" se={cell.get('std_error', '—')} p={cell.get('p_value', '—')}"
                f" {cell.get('significance', '')}".rstrip()
            )
            source_id = cell.get("source_id")
            if source_id:
                value += f" [source: {source_id}]"
            parts.append(value)
        items.append(" | ".join(parts))
    return items


def _headers(rows: Sequence[Mapping[str, Any]]) -> list[str]:
    headers: list[str] = []
    for row in rows:
        for key in row:
            if key not in headers:
                headers.append(str(key))
    return headers


def _sheet_title(name: str) -> str:
    invalid_characters = set("[]:*?/\\")
    title = "".join(
        character if character not in invalid_characters else "_"
        for character in str(name)
    )
    return (title or "table")[:31]
